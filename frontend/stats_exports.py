"""Exports of the exact dashboard snapshot and its rendered board."""

from itertools import zip_longest
from pathlib import Path

from openpyxl import Workbook
from PySide6.QtCore import QPoint, QRect
from PySide6.QtGui import QPixmap, QRegion
from PySide6.QtWidgets import QWidget

from statistics_model import BOARDS, summarize_buckets
from stats_text import group_label, label


def _number(value):
    return float(value) if value is not None else None


def _series_rows(metric_name, result, companies, layer):
    keys = sorted(result.series.keys() | result.comparison.keys())
    for company, group in keys:
        current = result.series.get((company, group), [])
        compared = result.comparison.get((company, group), [])
        base = [label(metric_name), company,
                companies.get(company, company) if company else '',
                group_label(group), layer, result.metric.unit]
        yield base + [label('summary'), result.current_window[0], result.current_window[1],
                      _number(summarize_buckets(current, result.metric)) if current else None,
                      all(bucket.complete for bucket in current) if current else False,
                      result.comparison_window[0] if result.comparison_window else None,
                      result.comparison_window[1] if result.comparison_window else None,
                      _number(summarize_buckets(compared, result.metric)) if compared else None,
                      all(bucket.complete for bucket in compared) if compared else False]
        for a, b in zip_longest(current, compared):
            yield base + [label('day') if result.query.grain == 'day' else label(result.query.grain),
                          a.start if a else None, a.end if a else None,
                          _number(a.value) if a else None, a.complete if a else False,
                          b.start if b else None, b.end if b else None,
                          _number(b.value) if b else None, b.complete if b else False]


def export_xlsx(snapshot, path, companies, network_snapshot=None, *, company_mode=None,
                satisfaction='satisfaction-speed'):
    """Write filters and all current/comparison buckets without source hours."""
    workbook = Workbook()
    info = workbook.active
    info.title = label('filters')
    filters = snapshot.filters
    info.append([label('companies'), ', '.join(companies.get(c, c) for c in filters.companies)])
    info.append([label('grain'), label(filters.grain)])
    info.append([label('current'), filters.start, filters.end])
    info.append([label('comparison-value'), *(filters.comparison or (None, None))])
    headers = [label('metrics'), label('company-id'), label('companies'), label('group'), label('layer'),
               label('unit'), label('grain'), label('start'), label('end'),
               label('current'), label('complete'), label('comparison-start'),
               label('comparison-end'), label('comparison-value'), label('comparison-complete')]
    for board_key, metric_names in BOARDS.items():
        sheet = workbook.create_sheet(label({'公司数据': 'company', '服务规模': 'service',
                                             '客流数据': 'passenger', '城市数据': 'city'}[board_key]))
        sheet.append(headers)
        for metric_name in metric_names:
            result = snapshot.results.get(metric_name)
            if result is not None:
                for row in _series_rows(metric_name, result, companies, label('summary')):
                    sheet.append(row)
            breakdown = snapshot.breakdowns.get(metric_name)
            if breakdown is not None and breakdown is not result:
                for row in _series_rows(metric_name, breakdown, companies, label('group')):
                    sheet.append(row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions

    if network_snapshot is not None:
        network = workbook.create_sheet('网络摘要')
        network.append(['分析模式', '范围开始', '范围结束', '对比开始', '对比结束',
                        '公司标识', '公司', '指标', '值', '单位',
                        '完整', '不可用原因', '分类', '分类值', '比较摘要', '比较详情'])
        filters = network_snapshot.filters
        comparison = filters.comparison or (None, None)
        for summary in network_snapshot.summaries:
            company_id = summary.company_id
            for value in summary.values:
                base = [{'overall': '总体', 'companies': '多公司对比',
                         'period': '单公司同期'}[network_snapshot.options.mode], filters.start, filters.end,
                        *comparison, company_id,
                        network_snapshot.companies.get(company_id, summary.title)
                        if company_id else summary.title,
                        value.title, _number(value.value), value.unit,
                        value.complete, value.reason]
                if value.details:
                    for category, amount in value.details:
                        network.append(base + [category, _number(amount), value.comparison.text, value.comparison.tooltip])
                else:
                    network.append(base + [None, None, value.comparison.text, value.comparison.tooltip])
        network.freeze_panes = 'A2'
        network.auto_filter.ref = network.dimensions
    if company_mode is not None:
        from card_comparisons import company_comparison
        comparisons = workbook.create_sheet('公司比较摘要')
        comparisons.append(['公司标识', '公司', '指标', '模式', '比较摘要', '比较详情', '变化幅度', '幅度单位', '排名'])
        for owner in snapshot.filters.companies:
            for key in ('cashflow', 'company-value', 'monthly-ticket', satisfaction, 'reputation', 'popularity'):
                comparison = company_comparison(snapshot, company_mode, key, owner, companies)
                comparisons.append([owner, companies.get(owner, owner), label(key), company_mode,
                    comparison.text, comparison.tooltip, _number(comparison.amount), comparison.unit, comparison.rank])
        comparisons.freeze_panes = 'A2'
    workbook.save(Path(path))


def export_city_xlsx(snapshot, path, state):
    """City-only values and actual source rows; presentation conversion occurs once."""
    import json
    workbook = Workbook()
    info = workbook.active
    info.title = '筛选'
    info.append(['范围', '全市'])
    info.append(['开始', snapshot.filters.start])
    info.append(['结束（不含）', snapshot.filters.end])
    info.append(['粒度', snapshot.filters.grain])
    info.append(['来源', snapshot.source])
    info.append(['能源换算', '游戏显示值 = 原始值 / 100；物理单位未确认'])
    display = workbook.create_sheet('显示状态')
    display.append(['方式图形', state.get('mode')])
    display.append(['显示曲线', json.dumps(state.get('curves', []), ensure_ascii=False)])
    display.append(['隐藏图例', json.dumps(state.get('hidden', {}), ensure_ascii=False)])
    display.append(['饼图有效', snapshot.pie_valid, snapshot.pie_reason])
    values = workbook.create_sheet('城市指标')
    values.append(['指标', '分项', '值', '单位', '观测时间', '完整', '口径', '比较摘要', '比较详情'])
    for model in snapshot.kpis.values():
        for detail in model.details or (model,):
            values.append([model.title, detail.title if model.details else '', _number(detail.value),
                detail.unit, detail.observed, detail.complete, detail.reason or model.reason,
                detail.comparison.text, detail.comparison.tooltip])
    charts = workbook.create_sheet('城市图表')
    charts.append(['图表', '系列', '开始', '结束', '观测时间', '值', '单位', '完整', '显示'])
    raw_sheet = workbook.create_sheet('原始观测')
    raw_sheet.append(['指标', '分组', '时间', '原值', '原分母', '当前槽位', '来源'])
    seen = set()
    for key, result in snapshot.charts.items():
        for (_, group), buckets in result.series.items():
            shown = (key != 'trip-time' or group in state.get('curves', [])) and group not in state.get('hidden', {}).get(key, [])
            for bucket in buckets:
                charts.append([result.metric.label, group, bucket.start, bucket.end, bucket.observed,
                    _number(bucket.value), result.metric.unit, bucket.complete, shown])
                for row in bucket.raw:
                    identity = (row.metric, row.group, row.time, row.value, row.divider)
                    if identity not in seen:
                        seen.add(identity)
                        raw_sheet.append([row.metric, row.group, row.time, row.value, row.divider, row.current,
                                          row.raw.get('分组来源', '原始城市分类')])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(Path(path))


def export_png(board_host: QWidget, path):
    """Render the entire board container, including content outside its viewport."""
    old_size = board_host.size()
    layout = board_host.layout()
    if layout is not None:
        layout.activate()
    target = old_size.expandedTo(board_host.sizeHint()).expandedTo(board_host.minimumSizeHint())
    try:
        board_host.resize(target)
        if layout is not None:
            layout.activate()
        image = QPixmap(target)
        image.fill()
        board_host.render(image, QPoint(), QRegion(QRect(QPoint(0, 0), target)),
                          QWidget.RenderFlag.DrawChildren)
        if not image.save(str(path), 'PNG'):
            raise OSError(path)
    finally:
        board_host.resize(old_size)
