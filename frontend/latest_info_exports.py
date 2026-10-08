"""Export the confirmed latest-information snapshot without querying or loading Qt."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from display_rules import store_text_literally, workbook_number, format_number, number_places, truncated_number
from statistics_model import METRICS
from stats_text import group_label

if TYPE_CHECKING:
    from dashboard_model import DashboardResult
    from latest_info_model import LatestInfoSnapshot


def _number(value, metric=None):
    return workbook_number(value, number_places(metric))


def _display(value):
    return format_number(value, grouped=False)


def _company(snapshot, owner):
    if not owner:
        return '全部公司'
    name = dict(snapshot.companies).get(owner, owner)
    return f'{name} [{owner}]'


def build_share_summary(snapshot: LatestInfoSnapshot) -> str:
    """Describe the exact displayed scope, including per-metric exceptions."""
    simulation = snapshot.simulation_time
    lines = ['最新信息', f'城市：{snapshot.city_name}', f'存档：{snapshot.save_name}',
             f'模拟时刻：{simulation:%Y-%m-%d %H:%M:%S}' if simulation else '模拟时刻：—',
             f'人口：{_display(snapshot.population)}',
             f'范围：{_company(snapshot, snapshot.company_id)} · {group_label(snapshot.mode)}']
    for metric in snapshot.metrics:
        unit = f' {metric.unit}' if metric.unit else ''
        scope_note = ('全市' if metric.key == 'public-transport-share' or
                      (metric.key == 'transfer-coefficient' and not snapshot.company_id) else '')
        mode_note = '全部制式' if metric.key == 'transfer-coefficient' and snapshot.mode != '综合' else ''
        notes = [part for part in (scope_note, mode_note,
                 '部分观测' if not metric.complete and _number(metric.value) is not None else '') if part]
        lines.append(f'{metric.title}：{format_number(metric.value, number_places(metric.key), grouped=False)}{unit}' +
                     (f'（{"；".join(notes)}）' if notes else ''))
    for highlight in snapshot.highlights:
        line = highlight.line
        detail = (f'{line.name}；{group_label(line.mode)}；{_company(snapshot, line.company_id)}；'
                  f'客流 {_display(line.passengers)} 人次；班次 {_display(line.departures)} 次') if line else '—'
        lines.append(f'{highlight.title}：{detail}')
    lines.append(f'总班次：{_display(snapshot.total_departures)} 次')
    return '\n'.join(lines)


def _append(sheet, values):
    sheet.append(values)
    # Source names are literal game/save text, including a leading equals sign.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def _line_values(line, snapshot):
    if line is None:
        return [None] * 9
    return [line.key, line.name, group_label(line.mode), line.company_id,
            dict(snapshot.companies).get(line.company_id, line.company_id),
            _number(line.passengers), _number(line.departures),
            _number(line.passengers_per_departure), _number(line.passengers_per_vehicle_km)]


def export_latest_info_xlsx(snapshot: LatestInfoSnapshot, alerts: DashboardResult | None, path) -> None:
    """Write numeric snapshot values; missing observations remain blank cells."""
    workbook = Workbook()
    info = workbook.active
    info.title = '范围与说明'
    _append(info, ['项目', '值', '说明'])
    for row in (
        ['城市', snapshot.city_name], ['存档', snapshot.save_name],
        ['模拟时刻', snapshot.simulation_time], ['人口', _number(snapshot.population)],
        ['公司', _company(snapshot, snapshot.company_id)], ['制式', group_label(snapshot.mode)],
        ['指标范围', snapshot.scope_text], ['缺测说明', '缺测为空值；真实零值保留为0；不补齐未知观测。'],
        ['提醒状态', '提醒未计算，不能判断是否有达标提醒。' if alerts is None else
         ('存在真实达标提醒。' if alerts.alerts else '无达标提醒；无可比数据时不生成提醒。')],
    ):
        _append(info, row)
    if alerts is not None:
        filters = alerts.filters
        _append(info, ['提醒公司', '、'.join(_company(snapshot, owner) for owner in filters.companies)
                      or '无公司选择（仅全市指标）'])
        _append(info, ['提醒粒度', {'hour': '小时', 'day': '日', 'week': '周', 'month': '月'}[filters.grain]])
        _append(info, ['提醒本期开始', filters.start])
        _append(info, ['提醒本期结束（不含）', filters.end])
        comparison = filters.comparison or (None, None)
        _append(info, ['提醒对比开始', comparison[0]])
        _append(info, ['提醒对比结束（不含）', comparison[1]])

    metrics = workbook.create_sheet('首页指标')
    _append(metrics, ['指标', '值', '单位', '指标范围', '观测状态', '缺测或口径说明'])
    for metric in snapshot.metrics:
        value = _number(metric.value, metric.key)
        _append(metrics, [metric.title, value, metric.unit, metric.scope,
                          '缺测' if value is None else ('完整' if metric.complete else '部分观测'), metric.reason])

    line_headers = ['线路标识', '线路', '制式', '公司标识', '公司', '当日客流（人次）', '当日班次（次）',
                    '每班次客流（人次）', '每车公里客流（人次每公里）']
    highlights = workbook.create_sheet('线路极值')
    _append(highlights, ['极值项目', *line_headers])
    for highlight in snapshot.highlights:
        _append(highlights, [highlight.title, *_line_values(highlight.line, snapshot)])
    top = workbook.create_sheet('客流前十')
    _append(top, ['排名', *line_headers])
    for rank, line in enumerate(snapshot.passenger_top10, 1):
        _append(top, [rank, *_line_values(line, snapshot)])

    departures = workbook.create_sheet('班次分类')
    _append(departures, ['制式', '班次（次）', '占总班次比例', '说明'])
    total = _number(snapshot.total_departures)
    for category in snapshot.departure_modes:
        value = _number(category.value)
        # Excel stores fractions; truncate the displayed percentage points.
        share = (float(truncated_number(Decimal(str(category.value)) * 100 /
                                       Decimal(str(snapshot.total_departures))) / 100)
                 if value is not None and total is not None and total > 0 else None)
        _append(departures, [group_label(category.mode), value, share,
                            '班次缺测' if value is None else ('总班次缺测或不大于零，无法计算占比' if share is None else '')])
    _append(departures, ['总计', total, 1 if total is not None and total > 0 else None,
                        '总班次缺测' if total is None else ''])
    for row in range(2, departures.max_row + 1):
        departures.cell(row, 3).number_format = '0.##%'

    reminder = workbook.create_sheet('关键提醒')
    _append(reminder, ['指标', '公司标识', '公司', '分组', '本期值', '对比值', '单位',
                      '本期开始', '本期结束（不含）', '对比开始', '对比结束（不含）', '原提醒原因'])
    if alerts is not None:
        names = dict(snapshot.companies)
        for alert in alerts.alerts:
            owner, group = alert.key
            metric = METRICS[alert.metric]
            _append(reminder, [metric.label, owner, names.get(owner, owner) if owner else '全市',
                               group_label(group), _number(alert.after, alert.metric), _number(alert.before, alert.metric), metric.unit,
                               alert.start, alert.end, alert.comparison_start, alert.comparison_end, alert.reason])

    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Microsoft YaHei UI', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if hasattr(cell.value, 'year'):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        for cell in sheet[1]:
            cell.font = Font(name='Microsoft YaHei UI', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='156E71')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                48, max(16, max(len(str(cell.value or '')) for cell in column) + 3))
    store_text_literally(workbook)
    workbook.save(Path(path))


def export_latest_info_png(target, path):
    """Qt is loaded only for an explicitly requested rendered-page export."""
    from stats_exports import export_png
    export_png(target, path)
