"""Opt-in local GUI acceptance using the actual application and real sessions.

Normal startup never imports this module. Outputs and settings are isolated;
the delivered executable exercises its own startup and parser dispatch.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
import time
import traceback


def _resize_for_workarea(window, width, height):
    from PySide6.QtCore import QSize
    from window_workarea import fit_window_to_workarea
    window._smoke_requested_size = [width, height]
    fit_window_to_workarea(window, requested_size=QSize(width, height))


def _window_workarea_evidence(widget):
    from PySide6.QtCore import QPoint, QRect
    window = widget.window()
    screen = window.screen()
    available, full, frame = screen.availableGeometry(), screen.geometry(), window.frameGeometry()
    def rectangle(rect):
        return [rect.x(), rect.y(), rect.width(), rect.height()]
    visible = QRect(window.mapToGlobal(QPoint(0, 0)), window.size()) if window.isMaximized() else frame
    inside = available.contains(visible)
    assert inside, ('window exceeds available workarea', rectangle(visible), rectangle(available))
    return dict(requested_size=getattr(window, '_smoke_requested_size', None),
                actual_window_size=[window.width(), window.height()],
                available_geometry=rectangle(available), screen_geometry=rectangle(full),
                frame_geometry=rectangle(frame), visible_geometry=rectangle(visible),
                within_available_geometry=inside,
                dpr=window.devicePixelRatioF())


def main(launcher: Path, argv):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--save', type=Path)
    source.add_argument('--job', type=Path)
    parser.add_argument('--tag', default='真实存档验收')
    parser.add_argument('--expected-dpr', type=float, required=True)
    parser.add_argument('--keep-open', action='store_true')
    parser.add_argument('--loading-profile-only', action='store_true')
    args = parser.parse_args(argv)
    if args.loading_profile_only and not args.save:
        parser.error('--loading-profile-only requires --save')
    original = args.save or args.job
    if not original.exists():
        parser.error('Save/session does not exist')
    if args.output.exists():
        parser.error('Use a new output directory')
    args.output.mkdir(parents=True)
    from startup_bootstrap import main as bootstrap_main

    def factory():
        from PySide6.QtCore import QSettings, QTimer
        import desktop_app
        settings = QSettings(str(args.output / 'prefs.ini'), QSettings.Format.IniFormat)
        desktop_app.QSettings = lambda *_args: settings
        desktop_app.JOBS = args.output / 'parsed-jobs'
        window = desktop_app.MainWindow()
        _resize_for_workarea(window, 960, 680)

        def begin():
            if window.property('startupHandoffPending'):
                QTimer.singleShot(30, begin)
                return
            try:
                if args.loading_profile_only:
                    from loading_profile import profile_loading
                    profile_loading(window, args)
                else:
                    _exercise(window, desktop_app, args)
                if args.keep_open:
                    _resize_for_workarea(window, 960, 680)
                    window.navigate(0)
                    window.show()
                    return
                code = 0
            except Exception:
                (args.output / 'failure.txt').write_text(traceback.format_exc(), encoding='utf-8')
                code = 1
            window.statistics_page.stop_workers()
            window.close()
            from PySide6.QtWidgets import QApplication
            QApplication.instance().exit(code)

        QTimer.singleShot(0, begin)
        return window

    result = bootstrap_main(launcher, window_factory=factory, trace_path=args.output / 'startup.json')
    if result:
        return result
    try:
        trace = json.loads((args.output / 'startup.json').read_text(encoding='utf-8'))
        names = [event['event'] for event in trace]
        assert 'helper_spawn_started' not in names
        assert names.index('logo_first_paint') < names.index('desktop_import_started') < names.index('window_first_paint')
        assert len({event['window_id'] for event in trace if 'window_id' in event}) == 1
    except Exception:
        (args.output / 'failure.txt').write_text(traceback.format_exc(), encoding='utf-8')
        return 1
    return 0



def _map_click(widget, settle):
    """Use the live widget's input path, including its enclosing scroll area."""
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    parent = widget.parentWidget()
    while parent is not None:
        if callable(getattr(parent, 'ensureWidgetVisible', None)):
            parent.ensureWidgetVisible(widget); settle(50); break
        parent = parent.parentWidget()
    assert widget.isVisible() and widget.isEnabled(), widget.accessibleName()
    QTest.mouseClick(widget, Qt.MouseButton.LeftButton)
    settle(50)


def _map_type(edit, text, settle):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    _map_click(edit, settle)
    QTest.keyClick(edit, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    QTest.keyClicks(edit, text)
    QTest.keyClick(edit, Qt.Key.Key_Return)
    settle(50)


def _map_choose(combo, value, settle):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    if combo.currentData() == value:
        return
    index = combo.findData(value); assert index >= 0
    _map_click(combo, settle); settle(300)
    menu = combo.dropMenu; assert menu is not None and menu.isVisible()
    view = menu.view
    item = next(view.item(i) for i in range(view.count())
                if view.item(i).text().strip() == combo.itemText(index))
    view.scrollToItem(item); settle(50)
    QTest.mouseClick(view.viewport(), Qt.MouseButton.LeftButton, pos=view.visualItemRect(item).center())
    settle(50)
    assert combo.currentData() == value


def _map_stable(page, settle, wait_until):
    def ready():
        motion = getattr(page, 'motion', getattr(page, '_navigation_motion', None))
        return page.canvas.can_export() and not (motion is not None and motion.running)
    page.canvas.prepare_frame()
    wait_until(ready, 120)
    settle(300)
    assert ready(), 'map changed during stable capture delay'


def _map_passenger_input_case(page, settle, wait_until, capture):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    _map_click(page.preset_pivot.items['network'], settle)
    _map_click(page.dock.tabs['filters'], settle)
    panel = page.panel_set
    _map_click(panel.reset_filters, settle)
    _map_stable(page, settle, wait_until)
    section = panel.filter_sections['passengers']
    if not section.is_expanded(): _map_click(section.toggle, settle)
    before = page.result; revision = page.canvas._revision
    editor = panel.passenger_editor
    _map_type(editor.min_edit, '5000', settle)
    _map_type(editor.max_edit, '15000', settle)
    typed = [editor.min_edit.text(), editor.max_edit.text()]
    positions = list(editor.axis.values())
    assert typed == ['5000', '15000'] and positions == [316, 548], (typed, positions)
    assert page.result is before and page.canvas._revision == revision
    assert panel.state()['passenger_min'] is None and panel.state()['passenger_max'] is None
    _map_stable(page, settle, wait_until); capture('map-native-passenger-typed')
    QTest.keyClick(editor.axis.handles[0], Qt.Key.Key_Right)
    settle(50)
    moved = editor.axis.values()[0]
    assert moved == 317 and editor.max_edit.text() == '15000'
    assert abs(float(editor.min_edit.text()) - 50000*(moved/1000)**2) < .01
    assert page.result is before and page.canvas._revision == revision
    slider_inputs = [editor.min_edit.text(), editor.max_edit.text()]
    _map_stable(page, settle, wait_until); capture('map-native-passenger-slider')
    _map_click(section.apply_button, settle)
    assert not section.is_expanded() and page.result is not before
    assert panel.state()['passenger_min'] == float(slider_inputs[0])
    assert panel.state()['passenger_max'] == 15000
    expected = {route.id for route in page.query.snapshot.routes
                if (value := page.query.stats.get(route.id)) is not None
                and value.passengers is not None and float(slider_inputs[0]) <= value.passengers <= 15000}
    assert {route.id for route in page.result.routes} == expected
    _map_stable(page, settle, wait_until); capture('map-native-passenger-confirmed')
    return dict(typed_inputs=typed, typed_positions=positions, slider_inputs=slider_inputs,
                draft_preserved_map=True, confirmed_collapsed=True,
                confirmed_route_ids=sorted(expected), summary=section.summary_label.text())


def _exercise_map_inputs(win, page, settle, wait_until, capture):
    """Frozen save/job smoke: actual Qt mouse/keyboard paths, no fake models."""
    from PySide6.QtCore import QPoint, QPointF, Qt
    from PySide6.QtGui import QWheelEvent
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication
    from unicodedata import normalize
    win.raise_(); win.activateWindow(); settle(300)
    canvas = page.canvas
    _map_stable(page, settle, wait_until)
    center, zoom = canvas.center, canvas.zoom
    origin = QPoint(canvas.width()//2, canvas.height()//2)
    target = origin+QPoint(55,35)
    QTest.mousePress(canvas, Qt.MouseButton.LeftButton, pos=origin)
    QTest.mouseMove(canvas, target, delay=30)
    QTest.mouseRelease(canvas, Qt.MouseButton.LeftButton, pos=target)
    assert canvas.center != center
    QApplication.sendEvent(canvas, QWheelEvent(QPointF(origin), QPointF(canvas.mapToGlobal(origin)),
        QPoint(), QPoint(0,120), Qt.MouseButton.NoButton, Qt.KeyboardModifier.NoModifier,
        Qt.ScrollPhase.ScrollUpdate, False))
    assert canvas.zoom > zoom
    _map_stable(page, settle, wait_until); capture('map-native-drag-zoom')
    navigation = dict(dragged=True, zoomed=True, center=list(canvas.center), zoom=canvas.zoom,
                      canvas_size=[canvas.width(),canvas.height()], window_size=[win.width(),win.height()])

    _map_click(page.preset_pivot.items['single'], settle)
    _map_click(page.dock.tabs['lines'], settle)
    listing = page.single_list_panel; view = listing.route_list
    item = next(view.item(i) for i in range(view.count()) if view.item(i).flags() & Qt.ItemFlag.ItemIsEnabled)
    identity = item.data(Qt.ItemDataRole.UserRole)
    view.scrollToItem(item); settle(50)
    QTest.mouseClick(view.viewport(), Qt.MouseButton.LeftButton, pos=view.visualItemRect(item).center())
    _map_stable(page, settle, wait_until)
    assert page.single_panel.state()['selected_route_id'] == identity
    assert page.surface.search_button.isVisible() and not page.surface.search.isVisible()
    native_window=win.window().windowHandle()
    assert native_window is not None
    QTest.mouseMove(native_window, canvas.mapTo(win.window(),QPoint(canvas.width()-30,canvas.height()-50)))
    settle(50)
    QTest.mouseMove(native_window,page.surface.search_button.mapTo(win.window(),page.surface.search_button.rect().center()))
    try:
        wait_until(lambda: page.surface.search.isVisible(), 5)
    except (RuntimeError,AssertionError):
        from PySide6.QtGui import QCursor
        point=page.surface.search_button.mapToGlobal(page.surface.search_button.rect().center())
        under=QApplication.widgetAt(QCursor.pos())
        raise AssertionError(dict(hover_failed=True,active=win.isActiveWindow(),
                                  cursor=[QCursor.pos().x(),QCursor.pos().y()],target=[point.x(),point.y()],
                                  under=type(under).__name__,under_is_target=under is page.surface.search_button,
                                  button_under_mouse=page.surface.search_button.underMouse())) from None
    _map_type(page.surface.search, '7', settle)
    expected = {row['id'] for row in page._presentation_catalog()
                if any('7' in normalize('NFKC',str(value)) for value in (row.get('search_name') or row['name'], row.get('number')) if value is not None)}
    found = {result.id for result in page.search('7')}
    assert found == expected
    if page.surface.results.isVisible():
        actual = {page.surface.results.item(i).data(Qt.ItemDataRole.UserRole).id
                  for i in range(page.surface.results.count())}
        assert actual == expected
    else:
        assert (not expected and page.single_panel.state()['selected_route_id'] == identity
                or len(expected) == 1 and page.single_panel.state()['selected_route_id'] in expected)
    _map_stable(page, settle, wait_until); capture('map-native-search-7')
    search = dict(query='7', matches=sorted(found), labels=[result.label for result in page.search('7')], hover_expanded=True)
    current=page.surface.results.currentItem()
    if page.surface.results.isVisible() and current is not None and current.flags() & Qt.ItemFlag.ItemIsEnabled:
        QTest.keyClick(page.surface.results, Qt.Key.Key_Return)
        assert page.single_panel.state()['selected_route_id'] in expected
        assert {route.id for route in page.result.routes} == {page.single_panel.state()['selected_route_id']}
        _map_stable(page,settle,wait_until);capture('map-native-search-selected')
    search['selected_id']=page.single_panel.state()['selected_route_id']
    # Dismiss results through the map and return through the actual list tab.
    QTest.mouseClick(canvas, Qt.MouseButton.LeftButton, pos=QPoint(canvas.width()-60,60))
    _map_click(page.dock.tabs['lines'], settle)
    identity = page.single_panel.state()['selected_route_id']
    sort_view = (*canvas.center,canvas.zoom)
    original_ids = {view.item(i).data(Qt.ItemDataRole.UserRole) for i in range(view.count())}
    _map_choose(listing.sort_combo, 'passengers', settle)
    if not listing.state()['sort_desc']: _map_click(listing.sort_order_button, settle)
    order = [view.item(i).data(Qt.ItemDataRole.UserRole) for i in range(view.count())]
    values = [getattr(page.query.stats.get(key),'passengers',None) for key in order]
    known = [value for value in values if value is not None]
    assert known == sorted(known, reverse=True)
    assert values == known+[None]*(len(values)-len(known))
    assert set(order) == original_ids and view.currentItem().data(Qt.ItemDataRole.UserRole) == identity
    assert page.single_panel.state()['selected_route_id'] == identity
    assert (*canvas.center,canvas.zoom) == sort_view
    _map_stable(page, settle, wait_until); capture('map-native-list-sorted')
    sorting = dict(by='passengers', descending=True, order=order, values=values, selected_id=identity, selection_preserved=True)

    _map_click(page.preset_pivot.items['planning'], settle)
    _map_stable(page, settle, wait_until)
    catalog = {row['id']:row for row in page._presentation_catalog()}
    for building in canvas.snapshot.buildings:
        source = building.service_lines
        if not (source.known and source.complete and source.route_ids and not source.unresolved_refs
                and all(key in catalog for key in source.route_ids) and len(building.polygon)>=3): continue
        if not any(catalog.get(key,{}).get('selectable') for key in source.route_ids): continue
        point = canvas.world_to_screen((sum(p[0] for p in building.polygon)/len(building.polygon),
                                        sum(p[2] for p in building.polygon)/len(building.polygon))).toPoint()
        hit = canvas.building_at(point)
        if canvas.rect().contains(point) and hit is not None and hit.id == building.id: break
    else: raise AssertionError('No visible native building has a selectable service-line menu')
    QTest.mouseClick(canvas, Qt.MouseButton.LeftButton, pos=point); settle(50)
    menu = page.planning_panel.building_menu
    assert menu is not None and menu.isVisible() and menu.source_complete
    candidates = [menu.line_list.item(i).data(Qt.ItemDataRole.UserRole) for i in range(menu.line_list.count())]
    assert candidates == [key for key in dict.fromkeys(source.route_ids) if key in catalog]
    enabled = {key for key in candidates if catalog[key]['selectable']}
    outside = set(page.planning_panel.state()['selected_ids'])-enabled
    _map_click(menu.clear_selection, settle)
    chosen = next(menu.line_list.item(i) for i in range(menu.line_list.count())
                  if menu.line_list.item(i).data(Qt.ItemDataRole.UserRole) in enabled)
    menu.line_list.scrollToItem(chosen); settle(50)
    rect = menu.line_list.visualItemRect(chosen)
    QTest.mouseClick(menu.line_list.viewport(), Qt.MouseButton.LeftButton, pos=QPoint(rect.right()-8,rect.center().y()))
    settle(50)
    assert chosen.checkState() == Qt.CheckState.Checked
    assert set(page.planning_panel.state()['selected_ids']) == outside | {chosen.data(Qt.ItemDataRole.UserRole)}
    selected_one=sorted(page.planning_panel.state()['selected_ids'])
    _map_stable(page, settle, wait_until); capture('map-native-building-selected')
    _map_click(menu.select_all, settle)
    assert set(page.planning_panel.state()['selected_ids']) == outside | enabled
    _map_stable(page, settle, wait_until); capture('map-native-building-all')
    building_evidence = dict(building_id=building.id,building_name=building.name,source=source.source,
                             candidate_ids=candidates,candidate_labels=[menu.line_list.item(i).text() for i in range(menu.line_list.count())],
                             selectable_ids=sorted(enabled),selected_one=selected_one,outside_ids=sorted(outside),
                             selected_ids=sorted(page.planning_panel.state()['selected_ids']),native_complete=True,
                             clicked_canvas_position=[point.x(),point.y()])
    _map_click(menu.close_button, settle)
    passenger = _map_passenger_input_case(page,settle,wait_until,capture)
    _map_click(page.panel_set.reset_filters, settle)
    _map_click(page.dock.tabs['display'], settle)
    panel=page.panel_set; metrics=[]
    _map_click(panel.controls['color_by'].buttons['interval'], settle)
    for mode in ('daytime','peak','all_day','passengers'):
        if mode == 'passengers': _map_click(panel.controls['color_by'].buttons['passengers'], settle)
        else: _map_choose(panel.controls['interval_mode'], mode, settle)
        _map_stable(page,settle,wait_until); capture('map-native-legend-'+mode)
        legend=canvas.options['legend_items']
        expected_metric='passengers' if mode=='passengers' else 'interval_peak' if mode=='peak' else 'interval'
        assert getattr(legend,'metric',None) == expected_metric
        metrics.append(dict(mode=mode,title=legend.title,date=legend.date,source=legend.source,unit=legend.unit,
                            labels=list(legend.labels),gradient_stops=list(legend.gradient_stops),
                            missing_colour=legend.missing_colour,route_colors=dict(page.result.route_colors)))
    return dict(navigation=navigation,search=search,sorting=sorting,building=building_evidence,passengers=passenger,metrics=metrics)


def _map_current_export(page, path):
    from PySide6.QtGui import QImage
    assert page.canvas.can_export(), 'map export requires the current complete frame'
    assert page.export_image(path)
    image=QImage(str(path))
    assert not image.isNull() and image.size()==page.canvas.export_size()
    legend_height=page.canvas.legend_strip_height() if page.canvas.options['legend_external'] else 0
    assert image.width()==page.canvas.width() and image.height()==page.canvas.height()+legend_height
    return dict(path=str(path),size=[image.width(),image.height()],
                canvas_size=[page.canvas.width(),page.canvas.height()],legend_height=legend_height)


def _exercise(win, desktop_app, args):
    from dataclasses import replace
    from PySide6.QtCore import QEventLoop, QPoint, QTimer, Qt
    from PySide6.QtWidgets import QApplication, QToolTip
    from shiboken6 import isValid
    from report_model import load_session
    from stats_charts import ChartPanel
    app = QApplication.instance()
    records = []

    def settle(ms=300):
        loop = QEventLoop()
        QTimer.singleShot(ms, loop.quit)
        loop.exec()

    def wait_until(predicate, timeout=120):
        end = time.monotonic() + timeout
        while not predicate():
            if time.monotonic() >= end:
                raise RuntimeError('Application readiness timed out')
            settle(25)
        settle()

    def capture(name, widget=win, *, delay=300, activate=True):
        if delay:
            settle(delay)
        path = args.output / (name + '.png')
        assert widget.grab().save(str(path))
        composited = None
        if app.platformName() == 'windows' and activate:
            widget.window().raise_()
            widget.window().activateWindow()
            settle(80)
            import ctypes
            foreground = ctypes.windll.user32.GetForegroundWindow
            foreground.restype = ctypes.c_void_p
            if widget.window().isActiveWindow() and foreground() == int(widget.window().winId()):
                position = widget.mapToGlobal(QPoint(0, 0))
                pixels = widget.screen().grabWindow(0, position.x(), position.y(),
                                                    widget.width(), widget.height())
                if foreground() == int(widget.window().winId()):
                    composited = args.output / (name + '-composited.png')
                    assert pixels.save(str(composited))
        records.append({'case':name, 'size':[widget.width(), widget.height()],
                        'screenshot':str(path), 'composited':str(composited) if composited else None})
        records[-1]['workarea'] = _window_workarea_evidence(widget)
        if name.startswith('map-'):
            canvas=win.map_page.canvas
            assert canvas.can_export(), 'map capture must use the current complete frame'
            motion=getattr(win.map_page,'motion',getattr(win.map_page,'_navigation_motion',None))
            assert motion is None or not motion.running, 'map capture must wait for navigation animation'
            records[-1].update(canvas_size=[canvas.width(),canvas.height()],dpr=canvas.devicePixelRatioF(),
                               center=list(canvas.center),zoom=canvas.zoom,complete_frame=True,animation_stable=True)
            strip=win.map_page.surface.legend_strip
            assert canvas.options['legend_external']
            if canvas.legend_strip_height():
                assert strip.isVisible() and canvas.geometry().bottom()<strip.geometry().top()
                assert strip.height()==canvas.legend_strip_height()
            else:
                assert strip.isHidden()
            def geometry(widget):
                rect=widget.geometry()
                return [rect.x(),rect.y(),rect.width(),rect.height()]
            records[-1].update(plot_rect=geometry(canvas),legend_rect=geometry(strip),legend_covers_map=False)

    assert abs(win.devicePixelRatioF() - args.expected_dpr) < .001, (win.devicePixelRatioF(), args.expected_dpr)
    from app_metadata import APP_NAME
    assert win.windowTitle() == APP_NAME
    assert win.windowFlags() & Qt.WindowType.FramelessWindowHint
    assert win.titleBar.height() == 32 and win.titleBar.iconLabel.isVisible()
    capture('startup-welcome')
    digest = None
    if args.save:
        digest = hashlib.sha256(args.save.read_bytes()).hexdigest()
        errors = []
        win.start_parse(args.save)
        assert win.worker is not None
        win.worker.failed.connect(errors.append)
        capture('actual-save-parsing')
        wait_until(lambda: bool(errors) or (win.worker is None and bool(win.data) and not win._awaiting_dashboards), 360)
        assert not errors, errors
        assert hashlib.sha256(args.save.read_bytes()).hexdigest() == digest
        data = win.data
        job = desktop_app.JOBS / data['manifest']['job_id']
    else:
        job = args.job
        data = load_session(job, args.tag)
        manifest = json.loads((job / 'manifest.json').read_text(encoding='utf-8'))
        data['manifest'] = manifest
        data['save_path'] = manifest['save_path']
        data['outputs'] = manifest['output_files']
        data['validation_status'] = manifest['validation_status']
        win.on_completed(data)
        wait_until(lambda: not win._awaiting_dashboards)
    page = win.statistics_page
    assert not page.analysis_mode_control._buttons['companies'].isEnabled()
    assert not page.network_mode_control._buttons['companies'].isEnabled()
    for width, height in ((960, 680), (1600, 900)):
        _resize_for_workarea(win, width, height)
        for index, label in ((0, 'home'), (2, 'lines'), (3, 'stats')):
            win.navigate(index)
            settle()
            if index == 2 and data.get('lines'):
                win.line_clicked(0, 0)
            if index != 3:
                capture(f'{label}-{width}x{height}')
            else:
                for family in ('company', 'network', 'city'):
                    page.tab_bar.setCurrentItem(family)
                    page._tab_changed(family)
                    capture(f'stats-{family}-{width}x{height}')
        win.navigate(0)
        for _ in range(50):
            win._refresh_exports()
        assert win.header.export_menu.view.count() == 6
        assert len(win.header.export_menu.actions()) == 5
        win.header._show_export_menu()
        capture(f'export-{width}x{height}', win.header.export_menu)
        win.header.export_menu.close()
    map_page = win.map_page
    map_errors = []
    map_page.failed.connect(map_errors.append)
    win.navigate(1)
    wait_until(lambda: bool(map_errors) or map_page.result is not None, 360)
    assert not map_errors, map_errors
    route = next(route for route in map_page.query.snapshot.routes
                 if route.mode != 'waterbus' and any(len(path) > 1 for path in route.paths))
    map_facts = {}
    for width, height in ((1440, 960), (960, 680)):
        _resize_for_workarea(win, width, height)
        settle()
        for preset in ('single', 'network', 'planning'):
            map_page.set_preset(preset)
            if preset == 'single':
                map_page.show_route(route.id, map_page.save_token)
            elif preset == 'planning':
                map_page.planning_panel.selectionChanged.emit({route.id})
                assert map_page.planning_panel.state()['building_emphasis'] is True
            map_page.canvas.prepare_frame()
            wait_until(lambda: bool(map_errors) or (
                map_page.canvas._frame_key == map_page.canvas._render_key()
                and map_page.canvas._frame_view == (*map_page.canvas.center, map_page.canvas.zoom)), 120)
            assert not map_errors, map_errors
            capture(f'map-{preset}-{width}x{height}')
            if preset == 'single':
                map_facts = {key: label.text() for key, label in map_page.single_panel.fact_labels.items()}
            elif preset == 'network':
                for panel in ('filters', 'display'):
                    map_page.dock.activate_panel(panel)
                    capture(f'map-network-{panel}-{width}x{height}')
                map_page.dock.activate_panel('layers')
    _resize_for_workarea(win, 1440, 960)
    settle()
    map_native_inputs = _exercise_map_inputs(win,map_page,settle,wait_until,capture)
    assert not map_errors, map_errors
    map_current_export = _map_current_export(map_page,args.output/'map-current.png')
    _resize_for_workarea(win, 960, 680)
    win.navigate(3)
    page.tab_bar.setCurrentItem('network')
    page._tab_changed('network')
    wait_until(lambda: page.network_snapshot is not None)
    average_chart = next(c for c in page.network_snapshot.charts if c.key == 'vehicles-running')
    page._network_options_changed(replace(page.network_options, vehicle='maximum'))
    wait_until(lambda: page.network_snapshot is not None)
    demand = next(v for v in page.network_snapshot.summaries[0].values if v.metric_id == 'vehicles-running')
    assert demand.complete and not demand.context and demand.title == '车辆'
    assert next(c for c in page.network_snapshot.charts if c.key == 'vehicles-running') == average_chart
    capture('demand-with-average-trend')
    page.network_mode_control.setCurrentKey('period')
    wait_until(lambda: page.network_snapshot is not None
               and page.network_snapshot.options.mode == 'period'
               and not page.network_workers and not page.workers)
    from PySide6.QtCore import QEvent, QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    comparison = next(p for p in page.network_dashboard.findChildren(ChartPanel)
                      if p.chart_views and p.mode == 'trend-bar'
                      and len({s.stack or s.key for s in p.chart_views[0].data.series}) >= 2)
    ancestor = comparison.parentWidget()
    while ancestor is not None:
        if callable(getattr(ancestor, 'ensureWidgetVisible', None)):
            ancestor.ensureWidgetVisible(comparison)
            break
        ancestor = ancestor.parentWidget()
    settle()
    bar_evidence = []

    def inspect_bars(panel, name):
        canvas = panel.chart_views[0]
        canvas.grab()
        hits = {}
        for index, stack, path in canvas._bar_hits:
            hits.setdefault(stack, (index, path.boundingRect().center()))
        assert len(hits) >= 2, name
        for position, (stack, (index, point)) in enumerate(list(hits.items())[:2]):
            QApplication.sendEvent(canvas, QMouseEvent(QEvent.Type.MouseMove, point,
                canvas.mapToGlobal(point.toPoint()), Qt.MouseButton.NoButton,
                Qt.MouseButton.NoButton, Qt.KeyboardModifier.NoModifier))
            settle()
            canvas.grab()
            assert canvas._hover == index and canvas._hover_stack == stack
            rows = canvas._tooltip_rows(index, stack)
            members = [s for s in canvas.visible_series() if (s.stack or s.key) == stack
                       and index < len(s.values) and s.values[index] is not None]
            assert rows and len(rows) == len(members)
            tip = canvas._overflow_tip
            if tip is not None and tip.isVisible():
                assert tip.height() >= tip.content_layout[1]
                assert tip.screen().availableGeometry().contains(tip.geometry())
                capture(name + '-' + str(position) + '-tooltip', tip, delay=0, activate=False)
            capture(name + '-' + str(position), panel, activate=False)
            bar_evidence.append({'case':name, 'stack':stack,
                                 'title':members[0].titles[index], 'rows':len(rows),
                                 'overflow':bool(tip is not None and tip.isVisible())})
        QApplication.sendEvent(canvas, QMouseEvent(QEvent.Type.MouseMove, QPointF(1, 1),
            canvas.mapToGlobal(QPoint(1, 1)), Qt.MouseButton.NoButton,
            Qt.MouseButton.NoButton, Qt.KeyboardModifier.NoModifier))
        settle()
        assert canvas._hover is None
        assert canvas._overflow_tip is None or not canvas._overflow_tip.isVisible()

    inspect_bars(comparison, 'comparison-bars')
    clone = comparison._open_fullscreen()
    settle()
    inspect_bars(clone, 'comparison-bars-expanded')
    clone.fullscreen_button.click()
    settle()
    compact_comparison = comparison._create_clone(win)
    compact_comparison.set_compact_height(190)
    compact_comparison.setGeometry(80, 130, 400, 190)
    compact_comparison.show()
    compact_comparison.raise_()
    settle()
    inspect_bars(compact_comparison, 'comparison-bars-400x190')
    compact_comparison.close()
    compact_comparison.deleteLater()
    settle()
    company = next(p for p in page.company_dashboard.findChildren(ChartPanel) if p.result is not None and p.chart_views)
    compact = ChartPanel(company.title_label.text(), default_mode=company.mode,
                         allowed_modes=company._allowed_modes if company._restricted_modes else None)
    compact.set_company_palette({key:color.name() for key, color in company._company_palette.items()})
    compact.set_category_palette(company._category_palette)
    compact.set_result(company.result, company.companies)
    compact.set_compact_height(190)
    _resize_for_workarea(compact, 400, 190)
    compact.show()
    capture('compact-400x190', compact)
    compact.close()
    for family in ('company', 'network', 'city', 'home'):
        win.navigate(0 if family == 'home' else 3)
        if family != 'home':
            page.tab_bar.setCurrentItem(family)
            page._tab_changed(family)
        _resize_for_workarea(win, 960, 680)
        settle()
        if family == 'home':
            source = win.latest_info_page.trend
        elif family == 'city':
            source = page.city_dashboard.panels['population']
        else:
            dashboard = page.company_dashboard if family == 'company' else page.network_dashboard
            wait_until(lambda: any(p.result is not None and p.chart_views
                                   for p in dashboard.findChildren(ChartPanel)))
            source = next(p for p in dashboard.findChildren(ChartPanel) if p.result is not None and p.chart_views)
        ancestor = source.parentWidget()
        while ancestor is not None:
            if callable(getattr(ancestor, 'ensureWidgetVisible', None)):
                ancestor.ensureWidgetVisible(source)
                settle()
                break
            ancestor = ancestor.parentWidget()
        clone = source._open_fullscreen()
        detail = source._fullscreen_dialog
        assert clone is not None and detail is not None and not detail.isWindow()
        settle()
        assert clone.fullscreen_button.isVisible() and clone.fullscreen_button.isEnabled()
        assert clone.geometry().bottom() <= detail.height()
        picture = detail.grab().toImage()
        for point in [QPoint(3, 3), *(c.mapTo(detail, QPoint(2, 2)) for c in clone.chart_views)]:
            assert picture.pixelColor(point).name() == '#ffffff'
        capture('expanded-' + family, detail)
        hidden = set(clone._hidden_groups)
        clone.fullscreen_button.click()
        wait_until(lambda: not isValid(detail) or not detail.isVisible(), 5)
        assert source._hidden_groups == hidden
    win.navigate(0)
    win.showNormal()
    _resize_for_workarea(win, 960, 680)
    settle()
    # A previous detail/activation can leave native geometry pending.
    _resize_for_workarea(win, 960, 680)
    settle()
    _window_workarea_evidence(win)
    home_structure = {}
    for name, panel in (('passengers', win.latest_info_page.passengers),
                        ('departures', win.latest_info_page.departures)):
        panel.show_structure()
        win.latest_info_page.scroll.ensureWidgetVisible(panel)
        settle()
        rows = panel.visible_mode_rows()
        assert rows and all(row.number.isVisible() and row.share.isVisible() for row in rows)
        assert panel.ring.center_total.isVisible() or panel.total_label.isVisible()
        home_structure[name] = {'structure': [{'count':row.number.text(), 'share':row.share.text()} for row in rows]}
        capture('home-structure-' + name, panel)
        panel.show_ranking()
        settle()
        ranking = panel.visible_ranking_rows()
        assert ranking and all(row.number.isVisible() for row in ranking)
        assert all(row.number.contentsRect().width() >= row.number.fontMetrics().horizontalAdvance(row.number.text())
                   for row in ranking)
        home_structure[name]['ranking'] = [{'key':row.line.key, 'count':row.number.text()} for row in ranking]
        capture('home-ranking-' + name, panel)
        panel.show_line_share()
        settle()
        shares = panel.visible_share_rows()
        assert shares and all(row.number.isVisible() and row.share.isVisible() for row in shares)
        assert panel.share_ring.center_total.isVisible() != panel.total_label.isVisible()
        assert all(row.name.width() <= 112 for row in shares)
        assert len({row.number.x() for row in shares}) == len({row.share.x() for row in shares}) == 1
        ring_left = panel.share_ring.mapTo(panel.line_share, QPoint(0, 0)).x()
        data_right = max(row.share.mapTo(panel.line_share, row.share.rect().bottomRight()).x() for row in shares)
        assert abs(ring_left - (panel.line_share.width() - 1 - data_right)) <= 8
        for row in shares:
            assert row.number.x() - row.name.geometry().right() - 1 >= 4
            assert row.share.x() - row.number.geometry().right() - 1 == 8
            assert all(field.contentsRect().width() >= field.fontMetrics().horizontalAdvance(field.text())
                       for field in (row.number, row.share))
            assert panel.rect().contains(row.mapTo(panel, row.rect().bottomRight()))
        home_structure[name]['line_share'] = [{'key':row.entry.key, 'count':row.number.text(), 'share':row.share.text()}
                                              for row in shares]
        home_structure[name]['line_share_total'] = panel.share_ring.center_total.text()
        home_structure[name]['line_share_columns'] = {'name_width':shares[0].name.width(),
            'count_x':shares[0].number.x(), 'share_x':shares[0].share.x(), 'row_width':shares[0].width()}
        capture('home-line-share-' + name, panel)
        panel.show_structure()
    # Compare both real model-backed cards at identical narrow/medium/wide
    # widths, rather than inferring the narrow case from a wide screenshot.
    from PySide6.QtWidgets import QWidget, QVBoxLayout
    from stats_controls import StatisticsScrollArea
    from latest_info_charts import PassengerRanking, DepartureStructure
    home_structure_responsive = []
    snapshot = win.latest_info_controller.snapshot
    for width in (400, 520, 864):
        host = QWidget(win, Qt.WindowType.Tool | Qt.WindowType.FramelessWindowHint)
        outer = QVBoxLayout(host)
        outer.setContentsMargins(0, 0, 0, 0)
        scroll = StatisticsScrollArea(host)
        scroll.setWidgetResizable(True)
        body = QWidget()
        scroll.setWidget(body)
        outer.addWidget(scroll)
        layout = QVBoxLayout(body)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)
        passengers, departures = PassengerRanking(), DepartureStructure()
        passengers.set_data(snapshot.lines, snapshot.passenger_top10, snapshot.passenger_modes)
        departures.set_data(snapshot.lines, snapshot.departure_modes, snapshot.total_departures)
        for panel in (passengers, departures):
            panel.show_line_share()
            layout.addWidget(panel)
        content_height = passengers.height() + departures.height() + 12
        scrollbar_width = scroll.verticalScrollBar().sizeHint().width()
        _resize_for_workarea(host, width + scrollbar_width, content_height)
        host.show()
        settle()
        geometries = []
        for name, panel in (('passengers', passengers), ('departures', departures)):
            scroll.ensureWidgetVisible(panel)
            settle()
            rows = panel.visible_share_rows()
            assert rows and all(row.number.isVisible() and row.share.isVisible() for row in rows)
            assert all(field.width() >= field.fontMetrics().horizontalAdvance(field.text())
                       for row in rows for field in (row.number, row.share))
            left = panel.share_ring.mapTo(panel.line_share, QPoint(0, 0)).x()
            right = panel.line_share.width() - 1 - max(
                row.share.mapTo(panel.line_share, row.share.rect().bottomRight()).x() for row in rows)
            assert abs(left - right) <= 8
            dot_left = rows[0].dot.mapTo(panel.line_share, QPoint(0, 0)).x()
            ring_right = panel.share_ring.mapTo(panel.line_share, panel.share_ring.rect().bottomRight()).x()
            ring_gap = dot_left - ring_right - 1
            name_gap = rows[0].number.x() - rows[0].name.geometry().right() - 1
            if panel.width() >= 520:
                assert max(left, right, ring_gap, name_gap) - min(left, right, ring_gap, name_gap) <= 2
            else:
                assert ring_gap >= 16 and name_gap >= 4
            assert all(row.share.x() - row.number.geometry().right() - 1 == 8 for row in rows)
            assert len({row.number.x() for row in rows}) == len({row.share.x() for row in rows}) == 1
            assert all(panel.rect().contains(row.mapTo(panel, row.rect().bottomRight())) for row in rows)
            geometries.append({'left':left, 'right':right, 'ring_gap':ring_gap, 'name_gap':name_gap,
                               'name_width':rows[0].name.width(), 'count_width':rows[0].number.width(),
                               'share_width':rows[0].share.width(), 'actual_card_width':panel.width()})
            capture('home-line-share-' + name + '-workarea-' + str(width), host)
        assert geometries[0] == geometries[1]
        capture('home-line-share-responsive-' + str(width), host)
        home_structure_responsive.append({'width':width, 'passengers':geometries[0], 'departures':geometries[1]})
        host.close()
        host.deleteLater()
        settle()
    # Use actual native button contents and painted arrow bounds, after theme
    # polish. Check all captions in both real home cards at the current DPR.
    from PySide6.QtCore import QRectF
    from PySide6.QtWidgets import QStyle, QStyleOptionButton
    for name, panel in (('passengers', win.latest_info_page.passengers),
                        ('departures', win.latest_info_page.departures)):
        panel.show_ranking()
        win.latest_info_page.scroll.ensureWidgetVisible(panel)
        menu = panel.mode_combo
        for index in range(menu.count()):
            menu.setCurrentIndex(index)
            settle()
            arrows = []
            original = menu._drawDropDownIcon
            menu._drawDropDownIcon = lambda painter, rect: (arrows.append(QRectF(rect)), original(painter, rect))[-1]
            try:
                menu.grab()
            finally:
                menu._drawDropDownIcon = original
            option = QStyleOptionButton()
            menu.initStyleOption(option)
            contents = menu.style().subElementRect(QStyle.SubElement.SE_PushButtonContents, option, menu)
            ink = menu.fontMetrics().boundingRect(contents, Qt.AlignmentFlag.AlignCenter | Qt.TextFlag.TextSingleLine, menu.text())
            assert arrows and ink.right() + 4 < arrows[-1].left()
            assert contents.contains(ink)
            capture('home-mode-menu-' + name + '-' + str(index), panel, delay=0, activate=False)
        panel.show_structure()
    for index, text in enumerate(('查看完整名称', '长标签验证：真实公司的完整名称    前一完整日' * 4, '短说明')):
        QToolTip.showText(win.mapToGlobal(QPoint(260, 180)), text, win)
        settle()
        tooltip = next((w for w in app.topLevelWidgets() if w.objectName() == 'qtooltip_label' and w.isVisible()), None)
        assert tooltip is not None
        assert tooltip.font().pixelSize() == 12 and not tooltip.font().bold()
        capture('native-tooltip-' + str(index), tooltip, delay=0, activate=False)
    QToolTip.hideText()
    settle()
    # Exercise the same export functions shipped in the GUI, using this real
    # snapshot. File picking is verified separately; no source file is edited.
    from latest_info_exports import export_latest_info_png, export_latest_info_xlsx
    from PySide6.QtGui import QImage
    from openpyxl import load_workbook
    exports = args.output / 'exports'
    exports.mkdir()
    export_latest_info_png(win.latest_info_page.export_target(), exports / 'home.png')
    export_latest_info_xlsx(win.latest_info_controller.snapshot,
                            win.latest_info_controller.alerts_snapshot, exports / 'home.xlsx')
    assert not QImage(str(exports / 'home.png')).isNull()
    exported = [exports / 'home.xlsx']
    if args.save:
        from display_rules import export_precision_workbook
        for kind, source in data.get('outputs', {}).items():
            if Path(source).suffix.lower() != '.xlsx':
                continue
            before = hashlib.sha256(Path(source).read_bytes()).hexdigest()
            destination = exports / (kind + '.xlsx')
            export_precision_workbook(source, destination)
            assert hashlib.sha256(Path(source).read_bytes()).hexdigest() == before
            exported.append(destination)
        assert len(exported) >= 3
    workbooks = []
    for path in exported:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            assert workbook.worksheets and all(sheet.max_row > 0 for sheet in workbook.worksheets)
            workbooks.append({'path':str(path), 'sheets':workbook.sheetnames})
        finally:
            workbook.close()
    from datetime import datetime, timedelta
    from PySide6.QtSvgWidgets import QSvgWidget
    from about_dialog import AboutDialog
    from stats_range_picker import RangePicker
    from stats_dialogs import _MessageSurface
    from window_chrome import FluentFileDialog
    now = datetime(2013, 5, 21)
    dialogs = [('about', AboutDialog(win)),
               ('range', RangePicker(now - timedelta(days=1), now, parent=win)),
               ('message', _MessageSurface('验证窗口', '窗口行为检查', win)),
               ('export-file-picker', FluentFileDialog(win, '保存 XLSX', str(job), 'Excel 工作簿 (*.xlsx)'))]
    dialogs[-1][1].setAcceptMode(FluentFileDialog.AcceptMode.AcceptSave)
    dialogs[-1][1].setFileMode(FluentFileDialog.FileMode.AnyFile)
    for name, dialog in dialogs:
        dialog.show()
        settle()
        if dialog.isWindow():
            assert dialog.windowFlags() & Qt.WindowType.FramelessWindowHint
        else:
            assert dialog.window() is win
        icon = dialog.windowIcon().pixmap(16, 16).toImage()
        assert not icon.isNull()
        assert all(icon.pixelColor(x, y).alpha() == 0
                   for x in range(icon.width()) for y in range(icon.height()))
        if name == 'about':
            assert not dialog.findChildren(QSvgWidget)
        if hasattr(dialog, 'titleBar'):
            assert not dialog.titleBar.iconLabel.isVisible()
        capture('secondary-' + name, dialog)
        dialog.close()
        settle()
        dialog.deleteLater()
        settle()
    report = {'frozen':bool(getattr(sys, 'frozen', False)), 'executable':sys.executable,
              'scale':os.environ.get('QT_SCALE_FACTOR', '1'), 'dpr':win.devicePixelRatioF(),
              'real_job':str(job), 'real_history_rows':len(data.get('history', [])),
              'current_demand':str(demand.value), 'raw_save':str(args.save) if args.save else None,
              'raw_save_sha256_unchanged':digest, 'synthetic_tooltip_style_probe':True,
              'comparison_bars':bar_evidence, 'exported_workbooks':workbooks, 'cases':records}
    report['home_structure'] = home_structure
    report['map'] = {'schema': map_page.query.snapshot.schema_version,
                     'route_id': route.id, 'single_facts': map_facts,
                     'routes': len(map_page.query.snapshot.routes),
                     'buildings': len(map_page.query.snapshot.buildings)}
    report['map']['native_inputs'] = map_native_inputs
    report['map']['current_export'] = map_current_export
    report['home_structure_responsive'] = home_structure_responsive
    (args.output / 'evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
