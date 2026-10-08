from __future__ import annotations

import os
import json
import subprocess
import sys
import uuid
import time
from zipfile import BadZipFile
from pathlib import Path
from shiboken6 import isValid

from PySide6.QtCore import QEvent, Qt, QSettings, QThread, Signal, QTimer
from PySide6.QtGui import QColor, QDragEnterEvent, QDropEvent, QFont, QIcon
from PySide6.QtWidgets import (QApplication, QHBoxLayout, QStackedWidget,
                               QTableWidgetItem, QVBoxLayout, QWidget)

PROJECT = Path(__file__).resolve().parents[1]
SRC = PROJECT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
# Namespace-qualified frontend/src imports must also work from a foreign cwd.
if str(PROJECT) not in sys.path:
    sys.path.insert(sys.path.index(str(SRC)) + 1, str(PROJECT))
from stats_style import initialize_theme
from stats_tokens import FONT_FAMILY, FONT_SIZE_BODY, NAV_WIDTH_EXPANDED, PAGE_BG, TEXT_PRIMARY, ACCENT
from window_chrome import FluentMainWindow, FluentFileDialog as QFileDialog
from startup_surface import (MINIMUM_WINDOW_SIZE, StartupSurface, center_startup_window,
                             initial_window_size)
from app_paths import jobs_directory
from app_metadata import APP_NAME, application_version, icon_svg_path
from stats_identity import save_fingerprint
APP_ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else PROJECT
JOBS = jobs_directory(APP_ROOT)
BUNDLE_ROOT = Path(getattr(sys, "_MEIPASS", "")) if getattr(sys, "frozen", False) else PROJECT
RUNTIME_DATA = BUNDLE_ROOT / "data"
CATALOG_DIR = BUNDLE_ROOT / "exports"
APP_VERSION = application_version(BUNDLE_ROOT)
DEFAULT_MANAGED = Path(r"D:\Program Files (x86)\Steam\steamapps\common\Cities in Motion 2\CIM2_Data\Managed")

_UI_LOADED = False
_UI_NAMES = {'psutil', 'load_workbook', 'ERROR_CODES', 'MODES', 'load_session',
             'export_precision_workbook', 'StatisticsPage', 'FluentIcon', 'InfoBar',
             'InfoBarPosition', 'NavigationDisplayMode', 'NavigationInterface',
             'NavigationItemPosition', 'PAGE_GUTTER', 'AppHeader', 'StartupWelcome',
             'LoadingOverlay', 'ParseProgressEstimator', 'parse_time', 'QMessageBox', 'PAGES'}


def _load_ui_dependencies():
    """Paint the same main window's logo before importing the business UI."""
    global _UI_LOADED, psutil, load_workbook, ERROR_CODES, MODES, load_session
    global export_precision_workbook, StatisticsPage, FluentIcon, InfoBar, InfoBarPosition
    global NavigationDisplayMode, NavigationInterface, NavigationItemPosition
    global PAGE_GUTTER, AppHeader, StartupWelcome, LoadingOverlay, ParseProgressEstimator
    global parse_time, QMessageBox, PAGES
    if _UI_LOADED:
        return
    import psutil
    from openpyxl import load_workbook
    from openpyxl.cell.cell import ERROR_CODES
    from report_model import MODES, load_session
    from display_rules import export_precision_workbook
    from statistics_page import StatisticsPage
    from qfluentwidgets import (FluentIcon, InfoBar, InfoBarPosition, NavigationDisplayMode,
                                NavigationInterface, NavigationItemPosition)
    from app_shell import PAGE_GUTTER, AppHeader
    from startup_welcome import StartupWelcome
    from loading_overlay import LoadingOverlay
    from parse_progress import ParseProgressEstimator
    from statistics_model import parse_time
    from stats_dialogs import FluentMessageBox as QMessageBox
    from map_icon import MAP_ICON
    PAGES = (('overview', '最新信息', FluentIcon.HOME),
             ('map', '地图显示', MAP_ICON),
             ('lines', '线路查询', FluentIcon.SEARCH),
             ('statistics', '统计数据', FluentIcon.PIE_SINGLE))
    _UI_LOADED = True


def __getattr__(name):
    # Existing diagnostic/test callers retain their module-level integration
    # seams while normal startup loads them only after the first logo paint.
    if name in _UI_NAMES:
        _load_ui_dependencies()
        return globals()[name]
    raise AttributeError(name)


def bundled_managed_root() -> Path:
    """Return the Managed directory shipped beside the portable executable."""
    extracted = Path(getattr(sys, "_MEIPASS", "")) / "game_runtime" / "Managed"
    if extracted.exists():
        return extracted
    return APP_ROOT / "game_runtime" / "Managed"


def find_managed_roots() -> list[Path]:
    # The portable runtime is authoritative.  Steam locations are fallbacks
    # for development builds or packages built without the local game files.
    roots = [bundled_managed_root(), APP_ROOT / "game_runtime" / "Managed"]
    configured = os.environ.get("CIM2_MANAGED_ROOT", "")
    if configured:
        roots.append(Path(configured))
    roots.append(DEFAULT_MANAGED)
    for base in (Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")), Path(os.environ.get("PROGRAMFILES", r"C:\Program Files"))):
        roots.extend([
            base / "Steam" / "steamapps" / "common" / "Cities in Motion 2" / "CIM2_Data" / "Managed",
            base / "SteamLibrary" / "steamapps" / "common" / "Cities in Motion 2" / "CIM2_Data" / "Managed",
        ])
    seen = set()
    return [p for p in roots if not (str(p).lower() in seen or seen.add(str(p).lower()))]


def locate_managed(saved: str = "") -> Path | None:
    # A frozen portable build must be self-contained.  Do not silently fall
    # back to Steam or a stale user setting, otherwise packaging omissions are
    # hidden on the development machine and only fail after distribution.
    if getattr(sys, "frozen", False):
        embedded = bundled_managed_root()
        return embedded if (embedded / "Assembly-CSharp.dll").exists() else None
    candidates = [bundled_managed_root()]
    if saved:
        candidates.append(Path(saved))
    candidates.extend(find_managed_roots())
    seen = set()
    for path in candidates:
        path = path / "Managed" if path.name.lower() == "cim2_data" else path
        key = str(path).lower()
        if key in seen:
            continue
        seen.add(key)
        if (path / "Assembly-CSharp.dll").exists():
            return path
    return None


def fmt(value, digits=2):
    if value is None or value == "":
        return "-"
    if isinstance(value, float):
        return f"{value:,.{digits}f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


class SortItem(QTableWidgetItem):
    def __lt__(self, other):
        left = self.data(Qt.ItemDataRole.UserRole)
        right = other.data(Qt.ItemDataRole.UserRole) if other is not None else None
        if isinstance(left, (int, float)) and isinstance(right, (int, float)):
            return left < right
        return self.text().casefold() < (other.text().casefold() if other is not None else "")


def item(value) -> QTableWidgetItem:
    cell = SortItem(fmt(value) if isinstance(value, (int, float)) else str(value or ""))
    if isinstance(value, (int, float)):
        cell.setData(Qt.ItemDataRole.UserRole, value)
    cell.setFlags(cell.flags() & ~Qt.ItemFlag.ItemIsEditable)
    return cell


class ParseWorker(QThread):
    progress = Signal(int, str)
    log = Signal(str)
    # A dict signal converts the entire session to QVariantMap and back while
    # holding Python execution, duplicating large histories. The worker stops
    # mutating data after emit; pass that session as a queued Python object.
    completed = Signal(object)
    failed = Signal(str)

    def __init__(self, save_path: Path, managed: Path, parent=None):
        super().__init__(parent)
        self.save_path = save_path
        self.managed = managed
        self.job_dir = JOBS / uuid.uuid4().hex
        self.cancel_requested = False
        self.process: subprocess.Popen | None = None

    def cancel(self):
        self.cancel_requested = True
        process = self.process
        if process and process.poll() is None:
            try:
                process.terminate()
            except ProcessLookupError:
                pass

    def run_command(self, command: list[str], env: dict[str, str], stage: int, label: str):
        if self.cancel_requested:
            raise RuntimeError("已取消解析")
        self.progress.emit(stage, label)
        log_path = self.job_dir / "backend.log"
        output_lines: list[str] = []
        self.process = subprocess.Popen(command, cwd=PROJECT, env=env, text=True,
                                        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                        encoding="utf-8", errors="replace", bufsize=1)
        # Cancellation may arrive while Popen is starting, before self.process exists.
        if self.cancel_requested:
            self.process.terminate()
            self.process.wait()
            self.process = None
            raise RuntimeError("已取消解析")
        assert self.process.stdout is not None
        with log_path.open("a", encoding="utf-8") as log_file:
            for line in self.process.stdout:
                if self.cancel_requested:
                    self.process.terminate()
                    self.process.wait()
                    self.process = None
                    raise RuntimeError("已取消解析")
                line = line.rstrip()
                if line:
                    output_lines.append(line)
                    log_file.write(line + "\n")
                    log_file.flush()
                    self.log.emit(line)
        code = self.process.wait()
        self.process = None
        if code != 0:
            tail = "\n".join(output_lines[-12:])
            detail = f"\n\n{tail}" if tail else ""
            raise RuntimeError(f"{label}失败（退出码 {code}）{detail}\n\n完整日志：{log_path}")

    def run(self):
        try:
            if self.save_path.suffix.lower() != ".save":
                raise ValueError("请选择 .save 存档文件")
            if not self.save_path.exists():
                raise FileNotFoundError(f"存档不存在：{self.save_path}")
            if not (self.managed / "Assembly-CSharp.dll").exists():
                raise FileNotFoundError("安装包缺少内置的 CIM2 v1.6.3 程序集，请重新下载完整 EXE")
            self.job_dir.mkdir(parents=True, exist_ok=True)
            env = os.environ.copy()
            env.update({
                "CIM2_EXPORT_DIR": str(self.job_dir),
                "CIM2_PAYLOAD_DIR": str(self.job_dir),
                "CIM2_RUNTIME_DATA_DIR": str(RUNTIME_DATA),
                "CIM2_MANAGED_ROOT": str(self.managed),
                "PYTHONIOENCODING": "utf-8",
            })
            tag = "运行时" if self.save_path.stem == "望春市6" else f"{self.save_path.stem}_运行时"
            if getattr(sys, "frozen", False):
                # The parser is embedded in this same one-file executable.
                command = [str(sys.executable), "--backend", str(self.save_path), tag]
                self.run_command(command, env, 10, "读取存档对象并生成报表")
            else:
                python = sys.executable
                self.run_command([python, str(SRC / "extract_runtime_data.py"), str(self.save_path)], env, 10, "读取存档对象")
                self.run_command([python, str(SRC / "build_line_workbook.py"), tag], env, 70, "生成线路工作簿")
                self.run_command([python, str(SRC / "build_company_workbook.py"), tag], env, 82, "生成公司工作簿")
            self.progress.emit(92, "校验并建立查询索引")
            from report_model import load_session
            from openpyxl import load_workbook
            from openpyxl.cell.cell import ERROR_CODES
            data = load_session(self.job_dir, tag, CATALOG_DIR)
            line_xlsx = self.job_dir / f"CIM2_线路发班整理_{tag}.xlsx"
            company_xlsx = self.job_dir / f"CIM2_公司信息整理_{tag}.xlsx"
            expected_sheets = ((line_xlsx, {"线路信息", "公司概览"}), (company_xlsx, {"表1_公司信息", "表2_人员表", "表3_票价表", "表4_周收支表"}))
            self.log.emit('CIM2_PROGRESS ' + json.dumps(dict(event='progress', phase='validation', done=0, total=len(expected_sheets))))
            for workbook_index, (workbook_path, required) in enumerate(expected_sheets, 1):
                if not workbook_path.exists():
                    raise RuntimeError(f"未生成工作簿：{workbook_path.name}")
                workbook = load_workbook(workbook_path, read_only=True, data_only=True)
                try:
                    missing = required - set(workbook.sheetnames)
                    if missing:
                        raise RuntimeError(f"工作簿缺少工作表：{', '.join(sorted(missing))}")
                    for sheet in workbook.worksheets:
                        for row in sheet.iter_rows():
                            for cell in row:
                                # Only Excel error codes are errors; user text such
                                # as a line named "#3" is valid data.
                                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in ERROR_CODES):
                                    raise RuntimeError(f"工作簿存在错误值：{sheet.title}!{cell.coordinate}")
                finally:
                    # read_only workbooks keep the file open (and locked on Windows).
                    workbook.close()
                self.log.emit('CIM2_PROGRESS ' + json.dumps(dict(event='progress', phase='validation', done=workbook_index, total=len(expected_sheets))))
            manifest = {
                "job_id": self.job_dir.name,
                "save_path": str(self.save_path),
                "save_type": data["save_type"],
                "simulation_date": data["metadata"].get("当前日期", ""),
                "simulation_time": data["metadata"].get("当前时间", ""),
                "company_count": data["counts"].get("companies", 0),
                "line_count": data["counts"].get("lines", 0),
                "vehicle_count": data["counts"].get("vehicles", 0),
                "timetable_count": data["counts"].get("timetables", 0),
                "departure_count": data["counts"].get("departures", 0),
                "output_files": {
                    "line_workbook": str(self.job_dir / f"CIM2_线路发班整理_{tag}.xlsx"),
                    "company_workbook": str(self.job_dir / f"CIM2_公司信息整理_{tag}.xlsx"),
                },
                "validation_status": "通过",
                "warnings": [],
            }
            (self.job_dir / "manifest.json").write_text(__import__("json").dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            data["manifest"] = manifest
            data["outputs"] = {
                "line_workbook": self.job_dir / f"CIM2_线路发班整理_{tag}.xlsx",
                "company_workbook": self.job_dir / f"CIM2_公司信息整理_{tag}.xlsx",
            }
            data["save_path"] = str(self.save_path)
            data["save_key"] = save_fingerprint(self.save_path)
            data["managed_root"] = str(self.managed)
            data["validation_status"] = "通过"
            if self.cancel_requested:
                raise RuntimeError("已取消解析")
            self.progress.emit(100, "完成")
            self.completed.emit(data)
        except Exception as exc:
            if self.cancel_requested:
                self.failed.emit("已取消解析")
            else:
                self.failed.emit(str(exc))


class StartupImports(QThread):
    loaded = Signal(object)

    def run(self):
        try:
            import importlib
            # These modules contain no widgets. File access and Python imports
            # need not freeze the already painted logo or its system window.
            for name in ('openpyxl', 'report_model', 'statistics_model',
                         'dashboard_model', 'network_model', 'city_model', 'latest_info_model'):
                if self.isInterruptionRequested():
                    return
                importlib.import_module(name)
            # This defines/imports classes and enum metadata only. QApplication
            # theme, fonts and every widget constructor remain on the UI thread.
            if not self.isInterruptionRequested():
                _load_ui_dependencies()
            self.loaded.emit(None)
        except Exception as exc:
            self.loaded.emit(exc)



class MainWindow(FluentMainWindow):
    """Application shell: navigation, one shared header, page stack and parsing."""

    def __init__(self, *, defer_startup=None):
        super().__init__()
        self.setWindowTitle(APP_NAME)
        self.setWindowIcon(QIcon(str(icon_svg_path(BUNDLE_ROOT))))
        self.setStyleSheet(f'QMainWindow {{ background: {PAGE_BG}; }}')
        self.resize(initial_window_size())
        self.setMinimumSize(MINIMUM_WINDOW_SIZE)
        center_startup_window(self)
        self._content_ready = False
        self._content_initializing = False
        self._welcome_handoff_complete = False
        self._welcome_handoff_gate = None
        self._welcome_handoff_callbacks = []
        self.setAcceptDrops(False)
        QApplication.instance().installEventFilter(self)
        self._startup_surface = StartupSurface(icon_svg_path(BUNDLE_ROOT), self)
        self.setCentralWidget(self._startup_surface)
        if defer_startup is None:
            defer_startup = bool(QApplication.instance().property('deferWindowStartup'))
        if not defer_startup:
            self.initialize_content()

    def initialize_content(self):
        if self._content_ready or self._content_initializing:
            return
        self._content_initializing = True
        self._prepare_content()
        self.build_ui()
        self._finish_content()

    def _prepare_content(self):
        _load_ui_dependencies()
        initialize_theme(QApplication.instance())
        # Run before the app-wide touch synthesizer while construction is partial.
        QApplication.instance().removeEventFilter(self)
        QApplication.instance().installEventFilter(self)
        self.settings = QSettings("CIM2SaveStats", "Desktop")
        self.data: dict = {}
        self.worker: ParseWorker | None = None
        self.selected_key = ""
        self._selected_line = None
        self._awaiting_dashboards = False

    def _finish_content(self):
        self._enable_mica()
        self.check_install()
        self._content_ready = True
        self._content_initializing = False
        QApplication.instance().removeEventFilter(self)
        self.setAcceptDrops(True)
        if self._startup_surface is not None:
            # The welcome paints the same opaque logo frame above the retained
            # cover. Keeping the cover above it would suppress backing-store
            # paints forever; deleting it here would leave a construction gap.
            self._startup_surface.stackUnder(self.empty_state)
        if not self.property('startupHandoffPending'):
            self.prepare_welcome_handoff(self._begin_welcome_if_allowed)

    def _begin_welcome_if_allowed(self):
        if not self.property('startupHandoffPending'):
            self.begin_welcome_transition()

    def prepare_welcome_handoff(self, ready):
        from startup_readiness import FirstFrameGate
        if self._welcome_handoff_complete:
            ready()
            return
        self._welcome_handoff_callbacks.append(ready)
        if self._welcome_handoff_gate is None:
            self._welcome_handoff_gate = FirstFrameGate(
                self, self._complete_welcome_handoff, surface=self.empty_state)
        self.empty_state.update()

    def _complete_welcome_handoff(self):
        if getattr(self, '_closing_app', False):
            return
        cover, self._startup_surface = self._startup_surface, None
        if cover is not None:
            cover.hide()
            cover.deleteLater()
        self._welcome_handoff_complete = True
        callbacks, self._welcome_handoff_callbacks = self._welcome_handoff_callbacks, []
        for callback in callbacks:
            callback()

    def initialize_content_async(self, ready, failed):
        if self._content_ready:
            ready()
            return
        if self._content_initializing:
            return
        self._content_initializing = True
        self._startup_ready, self._startup_failed = ready, failed
        self._startup_imports = StartupImports(self)
        self._startup_imports.loaded.connect(self._startup_imports_ready)
        self._startup_imports.start()

    def _startup_imports_ready(self, error):
        if getattr(self, '_closing_app', False):
            return
        if error is not None:
            self._startup_failed(error)
            return
        self._startup_steps = iter((self._prepare_content, self._build_shell,
                                   self.build_overview, self.build_map, self.build_lines,
                                   self._build_statistics, self._finish_ui, self._finish_content))
        self._startup_next()

    def _startup_next(self):
        if getattr(self, '_closing_app', False):
            return
        try:
            step = next(self._startup_steps, None)
            if step is None:
                self._startup_ready()
                return
            step()
            cover = self._startup_surface
            if cover is not None:
                cover.setGeometry(self.centralWidget().rect())
                if self._content_ready:
                    cover.stackUnder(self.empty_state)
                else:
                    cover.raise_()
            QTimer.singleShot(0, self._startup_next)
        except Exception as exc:
            self._startup_failed(exc)

    # ------------------------------------------------------------ window
    def _enable_mica(self):
        self.setProperty('nativeMicaEnabled', False)
        if (sys.platform != 'win32' or sys.getwindowsversion().build < 22000 or
                QApplication.platformName() == 'offscreen'):
            return
        from qframelesswindow import WindowEffect
        self.setObjectName('micaMainWindow')
        self.setStyleSheet(f'QMainWindow#micaMainWindow {{ background: {PAGE_BG}; }}')
        if not hasattr(self, '_mica_effect'):
            self._mica_effect = WindowEffect(self)
        self._mica_effect.setMicaEffect(self.winId(), isDarkMode=False, isAlt=False)
        self.setProperty('nativeMicaEnabled', True)

    def showEvent(self, event):
        super().showEvent(event)
        self._enable_mica()
        if self._content_ready and not self.property('startupHandoffPending'):
            if self._welcome_handoff_complete:
                self.begin_welcome_transition()
            elif self._welcome_handoff_gate is None:
                self.prepare_welcome_handoff(self._begin_welcome_if_allowed)

    def begin_welcome_transition(self):
        self.empty_state.begin_transition()
        if self.empty_state.isVisible():
            self.empty_state.open_button.setFocus()

    def focusNextPrevChild(self, forward):
        if not self._content_ready:
            controls = [button for button in (self.titleBar.minBtn, self.titleBar.maxBtn, self.titleBar.closeBtn)
                        if button.isVisible() and button.isEnabled()]
            if controls:
                current = QApplication.focusWidget()
                index = controls.index(current) if current in controls else (-1 if forward else 0)
                controls[(index + (1 if forward else -1)) % len(controls)].setFocus(Qt.FocusReason.TabFocusReason)
            return True
        # Keep keyboard traversal on the active covering surface without
        # disabling hidden page controls used by asynchronous preparation.
        if hasattr(self, 'empty_state') and self.empty_state.isVisible():
            target = (self.loading_overlay.cancel_button if self.loading_overlay.isVisible()
                      else self.empty_state.open_button)
            controls = [button for button in (self.titleBar.minBtn, self.titleBar.maxBtn,
                        self.titleBar.closeBtn, target) if button.isVisible() and button.isEnabled()]
            current = QApplication.focusWidget()
            index = controls.index(current) if current in controls else (-1 if forward else 0)
            controls[(index + (1 if forward else -1)) % len(controls)].setFocus(
                Qt.FocusReason.TabFocusReason if forward else Qt.FocusReason.BacktabFocusReason)
            return True
        return super().focusNextPrevChild(forward)

    def eventFilter(self, watched, event):
        if not getattr(self, '_content_ready', False) and isinstance(watched, QWidget):
            central = self.centralWidget()
            if central is not None and (watched is central or central.isAncestorOf(watched)):
                if event.type() in (QEvent.Type.MouseButtonPress, QEvent.Type.MouseButtonRelease,
                                    QEvent.Type.MouseButtonDblClick, QEvent.Type.TouchBegin,
                                    QEvent.Type.TouchUpdate, QEvent.Type.TouchEnd, QEvent.Type.TouchCancel,
                                    QEvent.Type.KeyPress, QEvent.Type.KeyRelease, QEvent.Type.ContextMenu,
                                    QEvent.Type.Wheel):
                    event.accept()
                    return True
        return super().eventFilter(watched, event)

    def build_ui(self):
        self._build_shell()
        self.build_overview()
        self.build_map()
        self.build_lines()
        self._build_statistics()
        self._finish_ui()

    def _build_shell(self):
        shell = QWidget()
        cover = self.takeCentralWidget()
        self.setCentralWidget(shell)
        if cover is not None:
            cover.setParent(shell)
            cover.setGeometry(shell.rect())
            cover.show()
        # setCentralWidget defers showing a replacement central widget. Show
        # it now so reparenting the only painted cover cannot hide that cover
        # for the next asynchronous construction slice.
        shell.show()
        outer = QHBoxLayout(shell)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)
        self.sidebar = NavigationInterface(shell, showReturnButton=False)
        # Fluent rewrites this action tip during both animated and immediate
        # navigation changes. Translate at that boundary without changing them.
        from PySide6.QtCore import QObject, QEvent

        class NavigationActionTip(QObject):
            def eventFilter(self, watched, event):
                if event.type() == QEvent.Type.ToolTipChange:
                    translated = {'Open Navigation': '展开导航',
                                  'Close Navigation': '收起导航'}.get(watched.toolTip())
                    if translated:
                        watched.setToolTip(translated)
                return False

        menu_button = self.sidebar.panel.menuButton
        self._navigation_action_tip = NavigationActionTip(menu_button)
        menu_button.installEventFilter(self._navigation_action_tip)
        menu_button.setToolTip('收起导航' if self.sidebar.panel.displayMode in (
            NavigationDisplayMode.EXPAND, NavigationDisplayMode.MENU) else '展开导航')
        self.sidebar.setExpandWidth(NAV_WIDTH_EXPANDED)
        self.sidebar.setMinimumExpandWidth(1000)
        self.sidebar.displayModeChanged.connect(self._navigation_mode_changed)
        self.nav_buttons = []
        for index, (route, title, icon) in enumerate(PAGES):
            self.nav_buttons.append(self.sidebar.addItem(
                route, icon, title, onClick=lambda checked=False, i=index: self.navigate(i), tooltip=title))
        from about_dialog import show_about_dialog
        self.about_button = self.sidebar.addItem(
            'about', FluentIcon.INFO, '关于', position=NavigationItemPosition.BOTTOM,
            onClick=lambda checked=False: show_about_dialog(self, root=BUNDLE_ROOT), selectable=False, tooltip='关于')
        outer.addWidget(self.sidebar)

        self.content_host = QWidget(shell)
        self.content_host.setObjectName('contentHost')
        self.content_host.setStyleSheet(f'QWidget#contentHost {{ background: {PAGE_BG}; }}')
        content = QVBoxLayout(self.content_host)
        content.setContentsMargins(0, 0, 0, 0)
        content.setSpacing(0)
        self.header = AppHeader(self.content_host)
        self.header.open_requested.connect(self.open_dialog)
        self.header.export_requested.connect(self._export_action)
        self.header.about_to_show_exports = self._refresh_exports
        content.addWidget(self.header)
        self.body = QStackedWidget(self.content_host)
        page_host = QWidget(self.body)
        page_layout = QVBoxLayout(page_host)
        page_layout.setContentsMargins(PAGE_GUTTER, 8, PAGE_GUTTER, 0)
        page_layout.setSpacing(0)
        self.pages = QStackedWidget(page_host)
        page_layout.addWidget(self.pages)
        self.body.addWidget(page_host)
        content.addWidget(self.body, 1)
        outer.addWidget(self.content_host, 1)
        self.empty_state = StartupWelcome(shell)
        self.empty_state.open_requested.connect(self.open_dialog)
        if cover is not None:
            cover.raise_()

    def _build_statistics(self):
        self.statistics_page = StatisticsPage(self.settings, parent=self.pages)
        self.pages.addWidget(self.statistics_page)
        # The statistics sub-tabs live in the shared header next to the title.
        self.stats_tabs = self.statistics_page.tab_bar
        self._mount_header_tabs(self.statistics_page,self.stats_tabs)
        self.statistics_page.export_button.hide()
        self.statistics_page.stats_integration.export_failed.connect(
            lambda error: self._notify('导出失败', str(error), error=True))
        self.statistics_page.query_failed.connect(self._dashboard_failed)

    def _mount_header_tabs(self,page,tabs):
        """Share the existing statistics tab sizing and AppHeader reflow."""
        from stats_controls import configure_navigation_pivot
        page.layout().removeWidget(tabs)
        tabs.setParent(self.header)
        tabs.hide()
        configure_navigation_pivot(tabs,in_header=True)

    def _finish_ui(self):
        self.loading_overlay = LoadingOverlay(self.centralWidget(), self.cancel_parse,
                                              prepare_destination=self._refresh_body)
        self._progress_predictor = None
        self._parse_stage = 0
        self._parse_stage_text = ''
        self._parse_started = 0.0
        self._progress_timer = QTimer(self)
        self._progress_timer.setInterval(250)
        self._progress_timer.timeout.connect(self._update_estimated_progress)
        self._ready_timer = QTimer(self)
        self._ready_timer.setInterval(30)
        self._ready_timer.timeout.connect(self._check_dashboard_ready)
        self.status_text = ''
        self._connect_latest_info()
        collapsed = str(self.settings.value('shell/sidebar_collapsed', 'false')).lower() == 'true'
        self.set_sidebar_collapsed(collapsed)
        self.navigate(0)
        self._refresh_body()

    # -------------------------------------------------------- navigation
    def _navigation_mode_changed(self, mode):
        collapsed = mode not in (NavigationDisplayMode.EXPAND, NavigationDisplayMode.MENU)
        self._sidebar_collapsed = collapsed
        if self.width() >= 1000:
            self.settings.setValue('shell/sidebar_collapsed', collapsed)

    def set_sidebar_collapsed(self, collapsed):
        self._sidebar_collapsed = bool(collapsed)
        if collapsed:
            panel = self.sidebar.panel
            panel.collapse()
            panel.expandAni.stop()
            panel.resize(48, panel.height())
            panel._onExpandAniFinished()
            self.sidebar.setFixedWidth(48)
        else:
            self.sidebar.expand(False)

    def toggle_sidebar(self):
        self.set_sidebar_collapsed(not self._sidebar_collapsed)
        self.settings.setValue('shell/sidebar_collapsed', self._sidebar_collapsed)

    def navigate(self, index):
        from stats_motion import SurfaceMotion
        previous = self.pages.currentWidget()
        changed = self.pages.currentIndex() != index
        if changed and previous is not None:
            motion = getattr(previous, 'motion', getattr(previous, '_navigation_motion', None))
            if motion is not None:
                motion.finish()
        self.pages.setCurrentIndex(index)
        route, title, icon = PAGES[index]
        tabs=self.stats_tabs if index==3 else self.map_tabs if index==1 else None
        self.header.set_page(title, icon, tabs)
        self.sidebar.setCurrentItem(route)
        for position, button in enumerate(self.nav_buttons):
            glyph = PAGES[position][2]
            button.setIcon(glyph.colored(ACCENT, ACCENT) if position == index else glyph)
        self._refresh_exports()
        if changed and self.data:
            page = self.pages.currentWidget()
            motion = getattr(page, 'motion', getattr(page, '_navigation_motion', None))
            if motion is None:
                motion = SurfaceMotion(page)
                page._navigation_motion = motion
            motion.reveal(float_in=True)

    @property
    def page_title(self):
        return self.header.title

    def _refresh_body(self):
        ready = bool(self.data) and not self._awaiting_dashboards
        self.body.setCurrentIndex(0)
        self.empty_state.setVisible(not ready)
        if not ready:
            self.empty_state.raise_()
            if self.loading_overlay.isVisible():
                self.loading_overlay.raise_()
        self.loading_overlay.raise_transition()
        self.stats_tabs.setEnabled(bool(self.data))

    # ----------------------------------------------------------- exports
    def _export_items(self):
        has_line = bool(self.data) and self._workbook('line_workbook') is not None
        has_company = bool(self.data) and self._workbook('company_workbook') is not None
        items = []
        index = self.pages.currentIndex()
        if index == 0:
            ready = self.latest_info_controller.snapshot is not None
            items += [('latest-xlsx', '首页报告 XLSX', FluentIcon.DOCUMENT, ready),
                      ('latest-png', '首页图片 PNG', FluentIcon.PHOTO, ready),
                      ('latest-copy', '复制当前摘要', FluentIcon.COPY, ready), (None, '', None, False)]
        elif index == 1:
            ready = self.map_page.result is not None and self.map_page.canvas.can_export()
            items += [('map-png', '地图图片 PNG', FluentIcon.PHOTO, ready), (None, '', None, False)]
        elif index == 3:
            items += [('stats-report', '统计报表（PNG + XLSX）', FluentIcon.DOCUMENT,
                       self.statistics_page.export_button.isEnabled()), (None, '', None, False)]
        items += [('line_workbook', '线路工作簿 XLSX', FluentIcon.SAVE, has_line),
                  ('company_workbook', '公司工作簿 XLSX', FluentIcon.SAVE, has_company)]
        return items

    def _refresh_exports(self):
        self.header.set_exports(self._export_items())
        self.header.export_button.setEnabled(bool(self.data))

    def _export_action(self, key):
        controller = self.latest_info_controller
        if key == 'map-png':
            page = self.map_page
            if page.result is None or not page.canvas.can_export():return
            result, token = page.result, page.save_token
            path, _ = QFileDialog.getSaveFileName(self, '导出地图', '地图.png', 'PNG 图片 (*.png)')
            if page.result is not result or page.save_token != token or not page.canvas.can_export():return
            if path and not page.export_image(path):
                self._notify('导出失败', '无法写入地图图片', error=True)
        elif key == 'latest-xlsx':
            controller._export('xlsx')
        elif key == 'latest-png':
            controller._export('png')
        elif key == 'latest-copy':
            controller.copy_summary()
            self._notify('已复制', '当前摘要已复制到剪贴板')
        elif key == 'stats-report':
            self.statistics_page.export_button.click()
        else:
            self.export_file(key)

    def _workbook(self, kind):
        source = self.data.get('outputs', {}).get(kind)
        return source if source and Path(source).is_file() else None

    def _notify(self, title, text, *, error=False):
        self.status_text = text
        if QApplication.platformName() == 'offscreen':
            return
        (InfoBar.error if error else InfoBar.success)(
            title, text, parent=self.content_host, position=InfoBarPosition.TOP_RIGHT,
            duration=6000 if error else 3000)

    # ------------------------------------------------------------- close
    def closeEvent(self, event):
        self._closing_app = True
        if not self._content_ready:
            preloader = getattr(self, '_startup_imports', None)
            if preloader is not None and preloader.isRunning():
                preloader.requestInterruption()
                preloader.finished.connect(self.close)
                event.ignore()
                return
            super().closeEvent(event)
            return
        self._ready_timer.stop()
        self._progress_timer.stop()
        self.loading_overlay.finish()
        self.statistics_page.stop_workers()
        if not self.map_page.stop_workers():
            event.ignore()
            for task in self.map_page.workers:
                if task.isRunning(): task.finished.connect(self.close)
            return
        if not self.latest_info_controller.stop_workers():
            event.ignore()
            for task in self.latest_info_controller.workers:
                if task.isRunning():
                    task.finished.connect(self.close)
            return
        worker = self.worker
        if worker and worker.isRunning():
            worker.cancel()
            if not worker.wait(3000):
                event.ignore()
                if not getattr(self, '_close_waiting', False):
                    self._close_waiting = True
                    worker.finished.connect(self.close)
                return
        super().closeEvent(event)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, 'sidebar') and self.width() < 1000 and not self._sidebar_collapsed:
            self.set_sidebar_collapsed(True)

    # ------------------------------------------------------------- pages
    def build_map(self):
        from map_page import MapPage
        self.map_page = MapPage(self.settings, self.pages, cache_dir=JOBS / 'map-cache')
        self.map_tabs=self.map_page.preset_pivot
        self._mount_header_tabs(self.map_page,self.map_tabs)
        self.map_page.failed.connect(lambda message: self._notify('地图读取失败', message, error=True))
        self.map_page.ready.connect(self._refresh_exports)
        self.map_page.canvas.export_availability_changed.connect(lambda _:self._refresh_exports())
        self.pages.addWidget(self.map_page)

    def build_overview(self):
        from latest_info_page import LatestInfoPage
        from latest_info_controller import LatestInfoController
        self.latest_info_page = LatestInfoPage(self.settings, parent=self.pages)
        self.overview_tab = self.latest_info_page
        self.company_combo = self.latest_info_page.company_combo
        self.mode_combo = self.latest_info_page.mode_combo
        self.latest_info_controller = LatestInfoController(self.latest_info_page, self.settings, self)
        self.pages.addWidget(self.latest_info_page)

    def _connect_latest_info(self):
        page, controller = self.latest_info_page, self.latest_info_controller
        page.open_save_requested.connect(self.open_dialog)
        page.line_export_requested.connect(lambda: self.export_file('line_workbook'))
        page.company_export_requested.connect(lambda: self.export_file('company_workbook'))
        page.line_requested.connect(self._open_latest_line)
        controller.thresholds_changed.connect(self.statistics_page.set_thresholds)
        self.statistics_page.set_thresholds(controller.thresholds)
        controller.snapshot_changed.connect(self._latest_info_ready)
        controller.query_failed.connect(lambda error: self._dashboard_failed(str(error)))
        controller.export_failed.connect(lambda error: self._notify('导出失败', str(error), error=True))

    def _latest_info_ready(self, snapshot):
        if snapshot is None or self.data is not self.latest_info_controller.data:
            return
        if snapshot.simulation_time is not None and self.statistics_page.store is None:
            data = dict(self.data, simulation_time=snapshot.simulation_time.isoformat(sep=' '))
            self.statistics_page.set_session(data)
        if self._awaiting_dashboards:
            self._check_dashboard_ready()
        self._refresh_exports()

    def _open_latest_line(self, key):
        line = next((row for row in self.data.get('lines', []) if row.get('key') == key), None)
        if line is None:
            return
        self.navigate(2)
        self.query.clear()
        self.line_company.setCurrentIndex(max(0, self.line_company.findData(str(line.get('公司标识') or ''))))
        self.line_mode.setCurrentIndex(max(0, self.line_mode.findData(line.get('运输制式') or '')))
        self.selected_key = key
        self.refresh_lines()

    def build_lines(self):
        from line_query_page import LinesPage
        from line_schedule_view import SchedulePanel
        self.lines_page = LinesPage(self, SchedulePanel)
        self.lines_page.mapRequested.connect(self._show_line_on_map)
        self.lines_tab = self.lines_page
        for name in ('query', 'line_company', 'line_mode', 'line_columns_button',
                     'line_columns_menu', 'line_count', 'line_table', 'detail_title',
                     'fact_menu_button', 'fact_menu', 'facts_host', 'facts_grid', 'schedule_panel'):
            setattr(self, name, getattr(self.lines_page, name))
        self.fact_cards = []
        self.pages.addWidget(self.lines_page)

    def _show_line_on_map(self, route_id, save_token):
        # Identity and epoch are captured when the detail is rendered, not clicked.
        if tuple(save_token)!=self.map_page.save_token:
            return False
        if not self.map_page.show_route(route_id, save_token):
            return False
        self.navigate(1)
        return True

    # ------------------------------------------------------------ import
    def check_install(self):
        managed = locate_managed(self.settings.value("managed_root", ""))
        self.status_text = "安装包缺少内置的 CIM2 v1.6.3 程序集" if managed is None else "等待导入 .save 存档"

    def choose_managed(self):
        selected = QFileDialog.getExistingDirectory(self, "选择 CIM2_Data 或 Managed 目录")
        if not selected:
            return
        managed = locate_managed(selected)
        if managed is None:
            QMessageBox.warning(self, "程序集不可用", "所选目录中没有 Assembly-CSharp.dll")
            return
        self.settings.setValue("managed_root", str(managed))
        self.status_text = f"已设置程序集目录：{managed}"

    def open_dialog(self):
        if not getattr(self, '_content_ready', True):
            return
        path, _ = QFileDialog.getOpenFileName(self, "打开 Cities in Motion 2 存档", "",
                                              "Cities in Motion 2 存档 (*.save)")
        if path:
            self.start_parse(Path(path))

    def dragEnterEvent(self, event: QDragEnterEvent):
        if not self._content_ready:
            event.ignore()
            return
        if event.mimeData().hasUrls() and any(url.toLocalFile().lower().endswith(".save")
                                              for url in event.mimeData().urls()):
            event.acceptProposedAction()
            self.empty_state.set_drag_active(True)
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self.empty_state.set_drag_active(False)
        super().dragLeaveEvent(event)

    def dropEvent(self, event: QDropEvent):
        if not self._content_ready:
            event.ignore()
            return
        self.empty_state.set_drag_active(False)
        for url in event.mimeData().urls():
            path = Path(url.toLocalFile())
            if path.suffix.lower() == ".save":
                self.start_parse(path)
                event.acceptProposedAction()
                return
        event.ignore()

    def _reset_line_page(self):
        self.selected_key = ''
        self._selected_line = None
        self.lines_page.set_map_target()
        self.clear_fact_cards()
        self.clear_schedule_tabs()
        self.detail_title.setText('选择线路查看详情')
        self.lines_page.list_footer.setText('从列表选择线路')

    def start_parse(self, path: Path):
        if not self._content_ready:
            return
        # Normalize Unicode paths before crossing the Qt/subprocess boundary.
        # Resolving/stat-ing an online placeholder can involve the cloud provider.
        # Only normalize text here; file I/O belongs to the cancellable worker.
        path = Path(os.path.abspath(os.path.expanduser(os.fsdecode(str(path)))))
        running = False
        if self.worker:
            try:
                running = self.worker.isRunning()
            except RuntimeError:
                self.worker = None
        if running:
            QMessageBox.information(self, "正在解析", "当前任务尚未完成，请先取消或等待完成")
            return
        managed = locate_managed(self.settings.value("managed_root", ""))
        if managed is None:
            if getattr(sys, "frozen", False):
                QMessageBox.critical(self, "安装包不完整", "当前 EXE 缺少内置的 CIM2 v1.6.3 程序集，请重新下载完整 EXE。")
            else:
                result = QMessageBox.warning(self, "缺少游戏程序集", "未找到 v1.6.3 程序集。现在选择安装目录？",
                                             QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
                if result == QMessageBox.StandardButton.Yes:
                    self.choose_managed()
                    managed = locate_managed(self.settings.value("managed_root", ""))
            if managed is None:
                return
        self._ready_timer.stop()
        self._awaiting_dashboards = False
        self.latest_info_controller.clear_session()
        self.data = {}
        self.statistics_page.clear_session()
        # Clear the previous session immediately so a replacement import
        # cannot leave stale rows, charts or export actions on screen.
        self.line_company.clear()
        self.line_mode.clear()
        self.line_table.setRowCount(0)
        self._reset_line_page()
        self.empty_state.set_note('')
        self.header.save_chip.set_context(path.name, '正在读取存档…')
        self.status_text = f"准备解析：{path.name}"
        self._progress_predictor = ParseProgressEstimator(0)
        self._parse_started = time.monotonic()
        self._parse_stage, self._parse_stage_text = 0, '准备读取存档'
        worker = ParseWorker(path, managed, self)
        self.worker = worker
        self._refresh_body()
        self._refresh_exports()
        self.loading_overlay.begin("准备读取存档")
        self._progress_timer.start()

        def current():
            return self.worker is worker and not getattr(self, '_closing_app', False)
        worker.progress.connect(lambda value, text: self.on_progress(value, text) if current() else None)
        worker.log.connect(lambda text: self.on_parse_log(text) if current() else None)
        worker.completed.connect(lambda data: self.on_completed(data)
                                 if current() and not worker.cancel_requested else None)
        worker.failed.connect(lambda message: self.on_failed(message) if current() else None)
        worker.finished.connect(lambda: self.worker_finished(worker))
        worker.start()
        # Warm the map cache alongside the statistics parser so navigating to
        # the map can reuse this exact worker or its completed snapshot.
        self.map_page.start_prefetch(path)

    def worker_finished(self, worker=None):
        # Drop the finished thread before the next import; a Python reference
        # to a deleted QThread makes ``isRunning`` raise on the next import.
        worker = worker if worker is not None else self.worker
        if self.worker is worker:
            self.worker = None
            if not self._awaiting_dashboards:
                self.loading_overlay.finish()
                self._progress_timer.stop()
            self._refresh_body()
        if worker:
            worker.deleteLater()

    def cancel_parse(self):
        if self._awaiting_dashboards:
            if self.worker:
                self.worker.cancel()
            self.on_failed('已取消解析')
        elif self.worker:
            self.worker.cancel()
            self.map_page.cancel_prefetch()
            self.status_text = "正在取消…"
            self.loading_overlay.cancel_button.setEnabled(False)
            self.loading_overlay.update_progress("正在取消…")

    def on_progress(self, value, label):
        self.status_text = label
        # Existing values are stage weights, not measured object percentages.
        self._parse_stage = min(value, 99)
        self._parse_stage_text = '准备图表' if value == 100 else label
        self._update_estimated_progress()

    def _update_estimated_progress(self):
        predictor = self._progress_predictor
        percent = None
        worker = self.worker
        cancelling = worker is not None and worker.cancel_requested
        if predictor is not None and not cancelling:
            rss = 0
            process = worker.process if worker is not None else None
            if process is not None and not getattr(predictor, 'actual_progress', False):
                try:
                    root = psutil.Process(process.pid)
                    rss = sum(p.memory_info().rss for p in [root, *root.children(recursive=True)])
                except (psutil.NoSuchProcess, psutil.AccessDenied):
                    pass
            percent = predictor.update(time.monotonic() - self._parse_started, rss, self._parse_stage)
        estimated = percent is not None and percent < 100 and not getattr(predictor, 'actual_progress', False)
        self.loading_overlay.update_progress(self._parse_stage_text, percent, estimated=estimated)

    def on_parse_log(self, text):
        if text.startswith('CIM2_PROGRESS '):
            try:
                details = json.loads(text.removeprefix('CIM2_PROGRESS '))
            except (ValueError, TypeError):
                return
            if isinstance(details, dict) and self._progress_predictor is not None:
                if details.get('event') == 'progress':
                    done, total = details.get('done'), details.get('total')
                    if type(done) is not int or type(total) is not int or total <= 0 or not 0 <= done <= total:
                        return
                    phase = {'lines': '提取线路', 'history': '读取历史记录',
                             'line_workbook': '生成线路报表', 'company_workbook': '生成公司报表',
                             'validation': '校验输出'}.get(details.get('phase'))
                    if phase:
                        self._parse_stage_text = f'{phase}：{done}/{total}'
                observer = getattr(self._progress_predictor, 'observe', None)
                if callable(observer):
                    observer(time.monotonic() - self._parse_started, details)
                self._update_estimated_progress()
            return
        self.status_text = text
        self._parse_stage_text = text
        self._update_estimated_progress()

    def on_failed(self, message):
        self.map_page.cancel_prefetch()
        self.latest_info_controller.clear_session()
        self.statistics_page.clear_session()
        self.data = {}
        self._awaiting_dashboards = False
        self._ready_timer.stop()
        self.loading_overlay.finish()
        self._progress_timer.stop()
        self.status_text = message
        cancelled = message == '已取消解析'
        self.header.save_chip.set_context('未载入存档', '已取消读取' if cancelled else '读取失败')
        self.empty_state.set_note('' if cancelled else f'读取失败：{message.splitlines()[0] if message else ""}')
        self._refresh_body()
        self._refresh_exports()
        if not cancelled:
            QMessageBox.critical(self, "解析失败", message)

    def on_completed(self, data):
        self._reset_line_page()
        self._awaiting_dashboards = True
        self._ready_started = time.monotonic()
        self._ready_data = data
        self._parse_stage, self._parse_stage_text = 99, '准备图表'
        self.loading_overlay.update_progress('准备图表')
        self._ready_timer.start()
        self.data = data
        self.map_page.set_session(data)
        self.statistics_page.clear_session()
        self.latest_info_controller.set_session(data)
        counts = data.get("counts", {})
        kind = "多人" if data.get("save_type") == "multiplayer" else "单人"
        save_name = Path(data["save_path"]).name
        try:
            simulation = parse_time(data.get('simulation_time'))
        except (TypeError, ValueError, AttributeError):
            simulation = None
        total_companies = counts.get('companies', len(data.get('companies', ())))
        clock = f'{simulation:%Y-%m-%d %H:%M}' if simulation else '模拟时间未提供'
        self.header.save_chip.set_context(save_name, f'{clock} · {kind} · {total_companies} 公司')
        self.status_text = ''
        self.line_company.blockSignals(True)
        self.line_company.clear()
        self.line_company.addItem("全部公司", userData="")
        names = [str(x.get('公司名称', '')) for x in data.get('companies', [])]
        for company in data.get('companies', []):
            name = str(company.get('公司名称', ''))
            company_id = str(company.get('公司标识') or name)
            self.line_company.addItem(f'{name} [{company_id}]' if names.count(name) > 1 else name, userData=company_id)
        self.line_company.blockSignals(False)
        self.line_mode.blockSignals(True)
        self.line_mode.clear()
        self.line_mode.addItem("全部制式", userData="")
        for mode in MODES:
            if any(x["运输制式"] == mode for x in data.get("lines", [])):
                self.line_mode.addItem(mode, userData=mode)
        self.line_mode.blockSignals(False)
        self.refresh_lines()
        self._refresh_body()
        self.navigate(0)

    def _dashboard_failed(self, message):
        if self._awaiting_dashboards:
            self.on_failed(f'准备图表失败：{message}')

    def _check_dashboard_ready(self):
        if not self._awaiting_dashboards or self.data is not self._ready_data:
            self._ready_timer.stop()
            return
        page, home = self.statistics_page, self.latest_info_controller
        snapshot = home.snapshot
        home_ready = (snapshot is not None and not home.query_timer.isActive()
                      and not any(worker.isRunning() for worker in home.workers)
                      and (not home.enabled or snapshot.simulation_time is None or home.alerts_snapshot is not None))
        stats_ready = (snapshot is not None and (snapshot.simulation_time is None or
                       (page.snapshot is not None and page.network_snapshot is not None
                        and not page.query_timer.isActive()
                        and not any(worker.isRunning() for worker in (*page.workers, *page.network_workers)))))
        if home_ready and stats_ready:
            self._awaiting_dashboards = False
            self._ready_timer.stop()
            self._progress_timer.stop()
            finish = getattr(self._progress_predictor, 'finish', None)
            if callable(finish):
                finish()
            self.loading_overlay.finish()
            self._refresh_body()
            self._refresh_exports()
        elif time.monotonic() - self._ready_started > 60:
            self.on_failed('准备图表超时，请重新导入存档')

    # -------------------------------------------------------------- lines
    def _company_matches(self, row, selected):
        if not selected:
            return True
        owner = str(row.get('公司标识') or '')
        if owner:
            return owner == selected
        selected_name = next((str(c.get('公司名称', '')) for c in self.data.get('companies', [])
                              if str(c.get('公司标识') or c.get('公司名称') or '') == selected), '')
        if not selected_name:
            return False
        same_name = [c for c in self.data.get('companies', []) if str(c.get('公司名称', '')) == selected_name]
        return len(same_name) == 1 and str(row.get('公司名称', '')) == selected_name

    def filtered_lines(self):
        from line_query_page import line_query_name
        query = self.query.text().strip().lower()
        company = self.line_company.currentData() or ""
        mode = self.line_mode.currentData() or ""
        return [x for x in self.data.get("lines", [])
                if self._company_matches(x, company) and (not mode or x["运输制式"] == mode)
                and (not query or query in " ".join(str(v) for v in x.values() if not isinstance(v, dict)).lower()
                     or query in line_query_name(x, self.data.get('lines', ())).lower())]

    def refresh_lines(self):
        if not hasattr(self, 'line_table'):
            return
        from line_query_page import line_display_value
        rows = self.filtered_lines()
        self.line_count.setText(f"{len(rows)} / {len(self.data.get('lines', []))} 条")
        sorting = self.line_table.isSortingEnabled()
        self.line_table.blockSignals(True)
        self.line_table.setSortingEnabled(False)
        self.line_table.setRowCount(len(rows))
        from stats_typography import ui_font
        body_font = ui_font(FONT_SIZE_BODY)
        for ri, row in enumerate(rows):
            values = [line_display_value(row, name, self.data.get('lines', ())) for name in (
                "公司名称", "运输制式", "线路名称", "地图里程", "折算里程", "单程时间", "核定速度", "今日客流",
                "当日发班数", "理论最大车辆需求数", "每周收入", "每周支出", "今日平均单班人次", "今日平均车公里人次")]
            for ci, value in enumerate(values):
                cell = item(value if value is not None else "—")
                cell.setToolTip(cell.text())
                cell.setFont(body_font)
                cell.setForeground(QColor(TEXT_PRIMARY))
                if ci == 2:
                    font = cell.font()
                    font.setWeight(QFont.Weight.Medium)
                    cell.setFont(font)
                self.line_table.setItem(ri, ci, cell)
            self.line_table.item(ri, 0).setData(Qt.ItemDataRole.UserRole, row["key"])
        self.line_table.setSortingEnabled(sorting)
        # Sorting changes row indices. Resolve the stable key from actual sorted items.
        selected = next((r for r in rows if r['key'] == self.selected_key), None)
        if selected:
            for ri in range(self.line_table.rowCount()):
                if self.line_table.item(ri, 0).data(Qt.ItemDataRole.UserRole) == self.selected_key:
                    self.line_table.selectRow(ri)
                    break
        else:
            self.line_table.clearSelection()
        self.line_table.blockSignals(False)
        if selected:
            self.show_line(selected)
        elif self.selected_key:
            self._reset_line_page()
            self.detail_title.setText('筛选结果中没有已选线路')

    def line_selection_changed(self):
        row = self.line_table.currentRow()
        if row >= 0 and self.line_table.selectedItems():
            self.line_clicked(row, 0)

    def line_clicked(self, row, _column):
        key_item = self.line_table.item(row, 0)
        key = key_item.data(Qt.ItemDataRole.UserRole) if key_item else ""
        line = next((x for x in self.data.get("lines", []) if x["key"] == key), None)
        if line:
            if self.selected_key == key and self._selected_line is line:
                return
            self.selected_key = key
            self.show_line(line)

    def show_line(self, line):
        from line_query_page import CompactFactCard, shown, line_display_value, line_query_name
        from PySide6.QtGui import QAction
        self._selected_line = line
        from display_rules import display_mode
        identity=line.get('对象ID',line.get('原始字段',{}).get('对象ID'))
        try:identity=int(identity)
        except (TypeError,ValueError):identity=None
        if display_mode(line.get('运输制式',''))=='水上巴士':identity=None
        self.lines_page.set_map_target(identity,self.map_page.save_token)
        display_name = line_query_name(line, self.data.get('lines', ()))
        self.detail_title.setText(f"{display_name} · {line['公司名称']}")
        facts = [
            (("线路车库", line.get("线路车库")), None),
            (("开线日期", line.get("开线日期")), ("最近改线日期", line.get("最近改线日期"))),
            (("地图里程", shown(line_display_value(line, "地图里程"), "km")), ("站点数", shown(line_display_value(line, "站点数"), "站"))),
            (("单程时间", shown(line.get("单程时间"), "min")), ("核定速度", shown(line_display_value(line, "核定速度"), "km/h"))),
            (("当日发班数", shown(line_display_value(line, "当日发班数"), "班")), ("平均间隔", line.get("平均间隔"))),
            (("理论最大车辆需求数", shown(line.get("理论最大车辆需求数"), "辆")), ("平均车辆需求数", shown(line.get("平均车辆需求数"), "辆"))),
            (("今日客流", shown(line.get("今日客流"), "人次")), ("平均客流", shown(line.get("平均客流"), "人次"))),
            (("今日平均单班人次", line_display_value(line, "今日平均单班人次")), ("今日平均车公里人次", line_display_value(line, "今日平均车公里人次"))),
            (("每周收入", line.get("每周收入")), ("每周支出", line.get("每周支出"))),
        ]
        self.fact_specs = facts
        self.clear_fact_cards()
        self.fact_menu.clear()
        for index, (primary, alternate) in enumerate(facts):
            tooltip = ''
            if index == 4:
                tooltip = '该线路该天运营时间内的平均间隔'
                if line.get('平均间隔') in (None, '', '—') and line.get('平均间隔Tooltip'):
                    tooltip += '\n' + line['平均间隔Tooltip']
            elif index == 5:
                tooltip = '游戏系统计算的线路平均车辆需求数'
                if line.get('平均车辆需求数') is None and line.get('平均车辆需求数Tooltip'):
                    tooltip += '\n' + line['平均车辆需求数Tooltip']
            elif index == 6:
                tooltip = '自开线以来的日均客流'
            card = CompactFactCard(primary, alternate, tooltip=tooltip)
            if index == 4:
                card.setToolTip('')
                card.label.setToolTip('')
                card.value.setToolTip('')
            self.fact_cards.append(card)
            self.facts_grid.addWidget(card, index // 3, index % 3)
            title = primary[0] + (' / ' + alternate[0] if alternate else '')
            action = QAction(title, self.fact_menu)
            action.setCheckable(True)
            action.setChecked(True)
            action.toggled.connect(card.setVisible)
            self.fact_menu.addAction(action)
        self.schedule_panel.set_line(line)
        self.lines_page.list_footer.setText(f"共 {len(self.data.get('lines', []))} 条线路    选中：{display_name}")

    def clear_schedule_tabs(self):
        self.schedule_panel.clear()

    def clear_fact_cards(self):
        for card in getattr(self, "fact_cards", []):
            self.facts_grid.removeWidget(card)
            card.hide()
            card.deleteLater()
        self.fact_cards = []

    def export_file(self, kind):
        data, controller = self.data, self.latest_info_controller
        token = controller.token
        session_key = str(data.get('save_key') or data.get('session_key') or '')
        source = data.get('outputs', {}).get(kind)
        if not source or not Path(source).is_file():
            QMessageBox.warning(self, '无法导出', '当前存档尚未生成该工作簿')
            return
        target, _ = QFileDialog.getSaveFileName(self, '保存 XLSX', Path(source).name, 'Excel 工作簿 (*.xlsx)')
        if (not target or self.data is not data or controller.token != token
                or str(self.data.get('save_key') or self.data.get('session_key') or '') != session_key
                or self.data.get('outputs', {}).get(kind) != source or not Path(source).is_file()):
            return
        try:
            export_precision_workbook(source, target)
            self._notify('已导出', Path(target).name)
        except (OSError, ValueError, BadZipFile) as exc:
            QMessageBox.critical(self, '导出失败', str(exc))


def main():
    from startup_bootstrap import main as bootstrap_main
    return bootstrap_main(PROJECT / 'CIM2_SaveStats.py')


if __name__ == "__main__":
    main()
