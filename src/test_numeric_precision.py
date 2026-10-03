"""User-facing precision is truncation; calculations retain the source values."""
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QPoint, QRect
from PySide6.QtWidgets import QLabel

from test_stats_charts import app, bucket, panel, result


@pytest.mark.parametrize('value,expected', [
    ('1.239', '1.23'), ('-1.239', '-1.23'), ('0.009', '0'),
    ('-0.009', '0'), ('1234.999', '1,234.99'), ('2.100', '2.1'),
])
def test_all_shared_ui_formatters_truncate(value, expected):
    from chart_canvas import format_value
    from chart_details import exact_number
    from city_dashboard import number
    from company_dashboard import _display
    from latest_info_charts import shown
    from card_comparisons import number as comparison_number
    from network_dashboard import _number
    from statistics_page import display
    for formatter in (format_value, exact_number, number, _display, shown,
                      comparison_number, _number, display):
        assert formatter(Decimal(value)) == expected


def test_vehicle_detail_and_chart_use_one_decimal_without_mutating_result():
    original = Decimal('121.92857724144345238')
    source = panel('平均运行车辆')
    data = result({('a', '总计'): [bucket(1, original)]}, metric='vehicles-running')
    source.set_result(data, {'a': '甲公司'})
    clone = source._open_fullscreen()
    try:
        dialog = source._fullscreen_dialog
        dialog.showNormal()
        dialog.resize(960, 680)
        app().processEvents()
        labels = clone.detail_summary.findChildren(QLabel, 'detailSummaryNumber')
        assert [label.text() for label in labels if label.isVisibleTo(clone.detail_summary)] == ['121.9']
        assert clone.chart_views[0].data.decimal_places == 1
        assert data.series[('a', '总计')][0].value == original
    finally:
        source._fullscreen_dialog.close()
        source.close()


def test_summary_card_reserves_large_number_and_unit_width():
    source = panel('现金流')
    source.set_result(result({('a', '总计'): [bucket(1, Decimal('123456789012345.239'))]},
                             metric='cashflow'), {'a': '甲公司'})
    clone = source._open_fullscreen()
    try:
        dialog = source._fullscreen_dialog
        dialog.showNormal()
        dialog.resize(960, 680)
        app().processEvents()
        for label in clone.detail_summary.findChildren(QLabel):
            if label.objectName() not in ('detailSummaryNumber', 'detailSummaryUnit'):
                continue
            card = label.parentWidget()
            assert card.rect().contains(QRect(label.mapTo(card, QPoint()), label.size()))
            assert label.fontMetrics().horizontalAdvance(label.text()) <= label.width()
    finally:
        source._fullscreen_dialog.close()
        source.close()


def test_summary_releases_preferred_width_when_host_shrinks():
    from PySide6.QtWidgets import QWidget
    from chart_details import DetailSummary
    app()
    host = QWidget()
    host.result = result({('a', '总计'): [bucket(1, 123), bucket(2, 111)]})
    host.companies = {'a': '甲公司'}
    host._comparison_name = lambda: '对比'
    host.resize(600, 300)
    summary = DetailSummary(host)
    summary.refresh()
    wide = summary._surfaces[0][0].width()
    host.resize(250, 300)
    summary._reflow_cards()
    assert summary._surfaces[0][0].width() < wide
    assert summary._surfaces[0][0].width() <= host.width() - 32
    host.close()


def test_statistics_workbook_truncates_average_vehicles_and_other_values(tmp_path):
    from test_stats_exports import snapshot
    from stats_exports import export_xlsx
    snap = snapshot()
    vehicles = result({('a', '总计'): [bucket(1, Decimal('121.928577'))]}, metric='vehicles-running')
    cash = result({('a', '总计'): [bucket(1, Decimal('-1.239'))]}, metric='cashflow')
    snap.results.update({'vehicles-running': vehicles, 'cashflow': cash})
    snap.breakdowns.update({'vehicles-running': vehicles, 'cashflow': cash})
    path = tmp_path / 'precision.xlsx'
    export_xlsx(snap, path, {'a': '甲公司'})
    book = load_workbook(path, data_only=True)
    numbers = [cell.value for sheet in book for row in sheet for cell in row if cell.data_type == 'n']
    assert 121.9 in numbers and -1.23 in numbers
    assert 121.928577 not in numbers and -1.239 not in numbers
    assert vehicles.series[('a', '总计')][0].value == Decimal('121.928577')


def test_latest_workbook_and_share_text_truncate(tmp_path):
    from test_latest_info_exports import sample_snapshot
    from latest_info_exports import export_latest_info_xlsx, build_share_summary
    snap = sample_snapshot()
    snap.metrics[4].value = Decimal('120.25999')
    path = tmp_path / 'latest.xlsx'
    export_latest_info_xlsx(snap, None, path)
    book = load_workbook(path, data_only=True)
    assert book['首页指标'].cell(6, 2).value == 120.25
    assert '周收入：120.25' in build_share_summary(snap)
    assert snap.metrics[4].value == Decimal('120.25999')


def test_percentage_export_truncates_percentage_points(tmp_path):
    from test_latest_info_exports import sample_snapshot
    from latest_info_exports import export_latest_info_xlsx
    snap = sample_snapshot(total_departures=7)
    snap.departure_modes[0].value = 3
    path = tmp_path / 'shares.xlsx'
    export_latest_info_xlsx(snap, None, path)
    book = load_workbook(path, data_only=True)
    assert book['班次分类']['C2'].value == .4285  # 42.85%, not rounded 42.86%


def test_network_comparison_keeps_one_decimal_for_vehicle_values():
    from card_comparisons import change
    comparison = change(Decimal('121.928577'), Decimal('110.9988'), '辆', '较上日', value_places=1)
    assert '当前 121.9 辆' in comparison.tooltip
    assert '对比 110.9 辆' in comparison.tooltip
    assert comparison.amount == (Decimal('121.928577') - Decimal('110.9988')) * 100 / Decimal('110.9988')


def test_cached_export_truncates_copy_and_preserves_source_and_identities(tmp_path):
    from openpyxl import Workbook
    from display_rules import export_precision_workbook
    source, target = tmp_path / 'cache.xlsx', tmp_path / 'export.xlsx'
    book = Workbook()
    book.active.append(['公司标识', '日期', '名称', '金额', '平均运行车辆数'])
    from datetime import datetime
    book.active.append(['001', datetime(2024, 1, 1), '=公司', -1.239, 121.928577])
    book.active['C2'].data_type = 's'
    book.save(source)
    original = source.read_bytes()
    export_precision_workbook(source, target)
    assert source.read_bytes() == original
    exported = load_workbook(target)
    assert list(exported.active.values)[1] == ('001', datetime(2024, 1, 1), '=公司', -1.23, 121.9)
    assert exported.active['C2'].data_type == 's'
    with pytest.raises(ValueError):
        export_precision_workbook(source, source)
    assert source.read_bytes() == original


def test_desktop_cached_export_action_uses_precision_copy(tmp_path, monkeypatch):
    from types import SimpleNamespace as NS
    from openpyxl import Workbook
    import desktop_app
    source, target = tmp_path / 'cache.xlsx', tmp_path / 'export.xlsx'
    book = Workbook()
    book.active.append(['金额'])
    book.active.append([1.239])
    book.save(source)
    original = source.read_bytes()
    data = {'save_key': 'stable-save', 'outputs': {'line_workbook': str(source)}}
    shell = NS(data=data, latest_info_controller=NS(token=3), _notify=lambda *_: None)
    monkeypatch.setattr(desktop_app.QFileDialog, 'getSaveFileName', lambda *_: (str(target), ''))
    desktop_app.MainWindow.export_file(shell, 'line_workbook')
    assert load_workbook(target).active['A2'].value == 1.23
    assert source.read_bytes() == original
