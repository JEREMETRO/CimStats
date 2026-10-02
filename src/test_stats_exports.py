import os
from datetime import datetime as D
from decimal import Decimal

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor
from PySide6.QtWidgets import QApplication, QScrollArea, QVBoxLayout, QWidget, QLabel

from dashboard_model import DashboardResult, FilterState
from statistics_model import Bucket, METRICS, Query, Result
from stats_exports import export_png, export_xlsx
from stats_charts import ChartPanel
import stats_tokens as tokens


def bucket(day, value, complete=True, numerator=None, denominator=None):
    return Bucket(D(2024, 1, day), D(2024, 1, day + 1), value, complete,
                  False, D(2024, 1, day, 23), numerator, denominator, 24)


def result(metric, group, values, comparisons):
    start, end = D(2024, 1, 3), D(2024, 1, 5)
    prior = (D(2024, 1, 1), D(2024, 1, 3))
    query = Query(metric, ('p1',), None, start, end, 'day', prior)
    return Result(query, METRICS[metric], {('p1', group): values},
                  {('p1', group): comparisons}, (start, end), prior)


def snapshot():
    filters = FilterState(('p1',), D(2024, 1, 3), D(2024, 1, 5),
                          'day', (D(2024, 1, 1), D(2024, 1, 3)))
    cash = result('cashflow', '总计', [bucket(3, Decimal(-2)), bucket(4, Decimal(5), False)],
                  [bucket(1, Decimal(1)), bucket(2, Decimal(3))])
    stops = result('stopcount', 'bus', [bucket(3, Decimal(4))],
                   [bucket(1, Decimal(2))])
    coefficient = result('transfer-coefficient', '总计',
                         [bucket(3, Decimal(3), numerator=30, denominator=10),
                          bucket(4, Decimal(2), numerator=20, denominator=10)],
                         [bucket(1, Decimal(1), numerator=10, denominator=10)])
    return DashboardResult(filters, {'cashflow': cash, 'stopcount': stops,
                                     'transfer-coefficient': coefficient},
                           {'cashflow': cash, 'stopcount': stops,
                            'transfer-coefficient': coefficient}, [])


def test_xlsx_has_filters_every_bucket_and_weighted_coefficient(tmp_path):
    path = tmp_path / 'board.xlsx'
    export_xlsx(snapshot(), path, {'p1': 'Transit'})
    wb = load_workbook(path, data_only=True)
    assert len(wb.sheetnames) == 5
    assert '原始' not in ' '.join(wb.sheetnames)
    rows = [tuple(row) for sheet in wb.worksheets[1:] for row in sheet.values]
    assert any('p1' in row and 'Transit' in row for row in rows)
    assert any('Transit' in row and '2024-01-04' in ' '.join(map(str, row)) and -2 in row for row in rows)
    assert any('Transit' in row and 5 in row and False in row for row in rows)
    assert any('Transit' in row and 3 in row and 1 in row for row in rows)
    assert any('Transit' in row and 2 in row for row in rows)
    assert any('Transit' in row and 2.5 in row for row in rows)
    assert any('bus' in row or '公交' in row for row in rows)


def test_png_captures_content_below_viewport_without_geometry_change(tmp_path):
    app = QApplication.instance() or QApplication([])
    host = QWidget()
    layout = QVBoxLayout(host)
    top = QLabel(); top.setFixedSize(80, 60); top.setStyleSheet('background:#ff0000')
    bottom = QLabel(); bottom.setFixedSize(80, 60); bottom.setStyleSheet('background:#0000ff')
    layout.addWidget(top); layout.addWidget(bottom)
    area = QScrollArea(); area.resize(120, 70); area.setWidget(host)
    area.show(); app.processEvents()
    old_size = host.size()
    path = tmp_path / 'board.png'
    export_png(host, path)
    from PySide6.QtGui import QImage
    image = QImage(str(path))
    assert image.height() > area.viewport().height()
    assert image.pixelColor(20, image.height() - 35).blue() > 200
    assert host.size() == old_size


def test_png_export_contains_the_fluent_chart_marks(tmp_path):
    app = QApplication.instance() or QApplication([])
    panel = ChartPanel('现金流', default_mode='trend-bar')
    panel.set_company_palette({'p1': tokens.COMPANY_COLORS[0]})
    panel.set_result(result('cashflow', '总计',
                            [bucket(3, Decimal(2)), bucket(4, Decimal(5))], []))
    panel.resize(700, 320)
    panel.show()
    app.processEvents()
    view = panel.chart_views[0]
    view.grab()
    region = next(region for region, hit in view._hit_regions
                  if hit['type'] == 'time-hit' and hit['bucket'].value == 5)
    pixel = view.viewport().mapTo(panel, view.mapFromScene(region.center()))
    path = tmp_path / 'fluent-chart.png'
    export_png(panel, path)
    from PySide6.QtGui import QImage
    image = QImage(str(path))
    color = image.pixelColor(pixel)
    assert not image.isNull()
    assert color.blue() > color.red() + 40
    assert panel.size().width() == 700
    panel.close()
