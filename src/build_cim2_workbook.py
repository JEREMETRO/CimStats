from pathlib import Path
import csv
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parents[1]
EXPORT = PROJECT / "exports"
TAG = sys.argv[1] if len(sys.argv) > 1 else "运行时"
SHEETS = [
    ("公司信息", f"CIM2_公司信息_{TAG}.csv"),
    ("公司车型", f"CIM2_公司车型数据_{TAG}.csv"),
    ("线路客流", f"CIM2_线路客流_完整导出_{TAG}.csv"),
    ("配车", f"CIM2_配车信息_完整导出_{TAG}.csv"),
    ("时刻表", f"CIM2_发班信息_完整导出_{TAG}.csv"),
    ("发班记录", f"CIM2_发班记录_完整导出_{TAG}.csv"),
    ("人工比对", f"CIM2_人工数据比对_完整_{TAG}.csv"),
    ("历史指标", f"CIM2_城市历史指标_完整_{TAG}.csv"),
    ("历史汇总", f"CIM2_城市历史指标汇总_{TAG}.csv"),
    ("客流分布", f"CIM2_城市客流分布_完整_{TAG}.csv"),
    ("站点客流", f"CIM2_线路站点客流分布_当前_{TAG}.csv"),
    ("线路汇总", f"CIM2_线路发班配车汇总_{TAG}.csv"),
    ("历史核心", f"CIM2_城市历史核心指标_{TAG}.csv"),
    ("历史元数据", f"CIM2_城市历史元数据_{TAG}.csv"),
    ("指标字典", f"CIM2_城市历史指标字典_{TAG}.csv"),
]

wb = Workbook()
ws = wb.active
ws.title = "说明"
ws.append(["Cities in Motion 2 存档数据导出"])
ws.append(["导出标签", TAG])
ws.append(["数据目录", str(EXPORT)])
ws.append(["程序集", "Cities in Motion 2 v1.6.3 Assembly-CSharp.dll"])
ws.append(["读取方式", "模拟 DataSerializer 对象表并屏蔽 Unity 后处理副作用；原始存档未修改"])
ws.append(["对象图结果", "详细数量见公司信息、线路客流、配车、时刻表和发班记录工作表"])
ws.append(["类型字段", "按存档 LineData.m_type 的 VehicleTypeObject.m_id 直接映射为游戏线路类型（有轨电车/公交等）；不使用能源或燃料推断"])
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 100

for name, filename in SHEETS:
    sheet = wb.create_sheet(name)
    with (EXPORT / filename).open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for col in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(r, col).value or "") for r in range(1, min(sheet.max_row, 200) + 1)]
        width = min(32, max(10, max((len(v) for v in values), default=10) + 2))
        sheet.column_dimensions[get_column_letter(col)].width = width

out = EXPORT / f"CIM2_{TAG}_线路配车发班导出.xlsx"
wb.save(out)
print(out)
