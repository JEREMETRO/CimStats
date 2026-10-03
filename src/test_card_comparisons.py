from datetime import datetime as D
from decimal import Decimal as V

from card_comparisons import change, peer_comparisons, previous_period, baseline_window
from dashboard_model import FilterState, build_dashboard
from statistics_model import HistoryStore
from test_statistics_model import row


def test_change_percent_points_negative_zero_and_missing():
    assert change(V(120), V(100), '人', '较上日').amount == 20
    assert change(V(68), V(70), '%', '较上日').amount == -2
    assert '-2 点' in change(V(68), V(70), '%', '较上日').text
    assert change(V(-80), V(-100), '货币', '较上日').amount == 20
    assert '增加 5 人' in change(V(5), V(0), '人', '较上日').text
    assert '持平' in change(V(0), V(0), '人', '较上日').text
    assert not change(V(5), None, '人', '较上日').available


def test_peer_two_companies_and_ranking_ties_and_missing():
    names = {'a': '甲', 'b': '乙', 'c': '丙', 'd': '丁'}
    pair = peer_comparisons({'a': V(100), 'b': V(80)}, names, '人')
    assert '领先' in pair['a'].text and pair['a'].amount == 25
    assert '落后' in pair['b'].text and pair['b'].amount == -20
    ranked = peer_comparisons({'a': V(100), 'b': V(100), 'c': V(80), 'd': None}, names, '人')
    assert ranked['a'].rank == ranked['b'].rank == 1
    assert '并列' in ranked['a'].text
    assert ranked['c'].rank == 3 and '20%' in ranked['c'].text
    assert not ranked['d'].available
    assert not peer_comparisons({'a': V(100)}, names, '人')['a'].available


def test_previous_period_handles_calendar_month_and_year():
    assert previous_period(D(2024, 3, 1), 'month') == (D(2024, 2, 1), D(2024, 3, 1))
    assert previous_period(D(2024, 1, 1), 'month') == (D(2023, 12, 1), D(2024, 1, 1))
    assert baseline_window(D(2024, 3, 1), D(2024, 4, 1)) == (D(2024, 2, 1), D(2024, 3, 1))
    assert baseline_window(D(2024, 1, 6), D(2024, 1, 8)) == (D(2023, 12, 30), D(2024, 1, 1))
    assert baseline_window(D(2024, 1, 6), D(2024, 1, 16)) == (D(2023, 12, 23), D(2024, 1, 2))


def test_dashboard_loads_previous_grain_outside_selected_range():
    rows = [row('cashflow', 'bus', D(2024, 1, day), amount, company='a')
            for day, amount in ((1, 10000), (2, 12000))]
    source = build_dashboard(HistoryStore(rows, D(2024, 1, 3)),
                             FilterState(('a',), D(2024, 1, 2), D(2024, 1, 3), 'day'))
    assert source.card_previous['cashflow'].series[('a', '总计')][0].value == 100


def test_network_period_summary_uses_selected_custom_baseline():
    from network_model import build_network_snapshot, NetworkOptions
    rows = [row('linecount', 'bus', D(2024, 1, day), amount, company='a')
            for day, amount in ((1, 4), (2, 5))]
    source = build_dashboard(HistoryStore(rows, D(2024, 1, 3)),
        FilterState(('a',), D(2024, 1, 2), D(2024, 1, 3), 'day', (D(2024, 1, 1), D(2024, 1, 2))))
    result = build_network_snapshot(source, NetworkOptions(mode='period'), {'a': '甲'})
    delta = result.summaries[0].values[0].comparison
    assert delta.amount == 25 and '同比' in delta.text


def test_city_weekend_comparison_is_city_only_and_keeps_density_peak():
    from city_model import build_city_snapshot
    rows = [row('population', 'A', D(2024, 1, day), amount)
            for day, amount in ((6, 100), (7, 120))]
    rows += [row('population', 'A', D(2023, 12, 31), 100)]
    filters = FilterState(('a',), D(2024, 1, 6), D(2024, 1, 8), 'day')
    snapshot = build_city_snapshot(HistoryStore(rows, D(2024, 1, 9)), filters)
    assert snapshot.kpis['population'].comparison.amount == 20
    assert '2023-12-30' in snapshot.kpis['population'].comparison.tooltip


def test_company_ui_and_export_share_peer_summary(tmp_path):
    import os
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from company_dashboard import CompanyDashboard
    from stats_exports import export_xlsx
    from openpyxl import load_workbook
    QApplication.instance() or QApplication([])
    rows = [row('cashflow', 'bus', D(2024, 1, 2), amount, company=company)
            for company, amount in [('a', 10000), ('b', 8000)]]
    source = build_dashboard(HistoryStore(rows, D(2024, 1, 3)),
        FilterState(('a', 'b'), D(2024, 1, 2), D(2024, 1, 3), 'day'))
    names = {'a': '甲', 'b': '乙'}
    widget = CompanyDashboard()
    from company_dashboard import DEFAULT_SLOTS
    widget.render(source, ('a', 'b'), names, {'a': '#1677FF', 'b': '#159A79'},
                  'companies', 'satisfaction-speed', DEFAULT_SLOTS)
    assert '25%' in widget.groups['a'].kpis['cashflow'].comparison_label.full_text
    path = tmp_path / 'comparisons.xlsx'
    export_xlsx(source, path, names, company_mode='companies')
    book = load_workbook(path, data_only=True)
    exported = list(book['公司比较摘要'].values)
    assert any(row[0] == 'a' and row[6] == 25 for row in exported[1:])
    book.close()
    widget.close()


def test_comparison_label_bolds_only_numbers_and_units():
    from PySide6.QtGui import QTextDocument
    from PySide6.QtWidgets import QApplication
    from PySide6.QtTest import QTest
    from card_comparison_label import ComparisonLabel
    QApplication.instance() or QApplication([])
    label = ComparisonLabel(change(V(125), V(100), '人', '较上日'))
    label.resize(360, 24)
    label.show()
    QTest.qWait(20)
    document = QTextDocument()
    document.setDefaultFont(label.font())
    document.setHtml(label.text())
    assert document.find('+25%').charFormat().fontWeight() == 600
    assert document.find('较上日').charFormat().fontWeight() == 400
    label.set_comparison(change(V(68), V(70), '%', '较上日'))
    document.setHtml(label.text())
    assert document.find('-2 点').charFormat().fontWeight() == 600
    assert label.accessibleName() == label.full_text
    label.close()
