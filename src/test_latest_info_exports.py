"""Pure export contract tests; importing this module must never start Qt."""
from dataclasses import replace
from datetime import datetime as D
from decimal import Decimal
import importlib
import importlib.util
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace as NS

from openpyxl import load_workbook
import pytest

from dashboard_model import DashboardResult, FilterState
from statistics_model import Alert


def exports():
    assert importlib.util.find_spec('latest_info_exports') is not None, '首页导出模块尚未实现'
    return importlib.import_module('latest_info_exports')


def sample_snapshot(**changes):
    # Complete frozen A contract, independent of its concurrent implementation.
    specs = (
        ('line-count', '线路数', 10, '条'), ('fleet', '车辆数', 20, '辆'),
        ('drive-minutes', '行车总时间', 80, '分钟'), ('turnover', '周转量', 120, '人公里'),
        ('weekly-income', '周收入', Decimal('120.25'), '货币'),
        ('weekly-expense', '周支出', 100, '货币'), ('profit', '利润', Decimal('20.25'), '货币'),
        ('interval', '平均间隔', None, '分钟'), ('speed', '平均速度', 30, '公里每小时'),
        ('passengers-per-run', '每班次客流', Decimal('2.5'), '人次'),
        ('passengers-per-km', '每车公里客流', Decimal('1.5'), '人次每公里'),
        ('public-transport-share', '公共交通分担率', Decimal('45.5'), '%'),
        ('transfer-coefficient', '换乘系数', Decimal('3.25'), ''),
    )
    metrics = tuple(NS(key=k, title=t, value=v, unit=u,
                       scope='全市当日' if k == 'public-transport-share' else '公司当日',
                       reason='没有运行日记录' if v is None else '', complete=v is not None)
                    for k, t, v, u in specs)
    lines = tuple(NS(key=f'线路-{i}', name=f'长名称线路{i}', mode='公交', company_id='p2',
                     passengers=100-i, departures=i, passengers_per_departure=Decimal('2.5'),
                     passengers_per_vehicle_km=None) for i in range(1, 11))
    snap = NS(session_key='save-a', company_id='p2', mode='公交', city_name='秋山市',
              simulation_time=D(2024, 3, 3, 12, 30), population=None, save_name='秋山市.save',
              companies=(('p1', '同名公司'), ('p2', '同名公司')),
              modes=('公交', '有轨电车', '无轨电车', '地铁', '水上巴士', '其他'),
              metrics=metrics, lines=lines, passenger_top10=lines,
              highlights=tuple(NS(title=t, line=l) for t, l in zip(
                  ('当日最大客流线路', '当日最小客流线路', '当日最多班次线路', '当日最少班次线路'),
                  (lines[0], lines[-1], lines[-1], lines[0]))),
              passenger_modes=(NS(mode='公交', value=955),),
              departure_modes=tuple(NS(mode=m, value=v) for m, v in zip(
                  ('公交', '有轨电车', '无轨电车', '地铁', '水上巴士', '其他'),
                  (30, 10, 5, 5, 0, None))), total_departures=50, trend=None,
              scope_text='公司同名公司 [p2]；公交；模拟当日截至12:30')
    snap.__dict__.update(changes)
    return snap


def sample_alerts(alerts=None):
    filters = FilterState(('p2',), D(2024, 3, 2), D(2024, 3, 3), 'hour',
                          (D(2024, 3, 1), D(2024, 3, 2)))
    alert = Alert('cashflow', ('p2', '总计'), '现金流转负', Decimal('10.25'), Decimal('-3.5'),
                  filters.start, filters.comparison[0], filters.end, filters.comparison[1])
    return DashboardResult(filters, {}, {}, [alert] if alerts is None else alerts)


def workbook(tmp_path, snap=None, alerts=None):
    path = tmp_path / '最新信息.xlsx'
    exports().export_latest_info_xlsx(snap or sample_snapshot(), alerts, path)
    return load_workbook(path, data_only=True)


def test_summary_reports_simulation_scope_and_missing_instead_of_zero():
    snap = sample_snapshot()
    summary = exports().build_share_summary(snap)
    for text in ('秋山市', '秋山市.save', '2024-03-03 12:30', '同名公司', 'p2', '公交',
                 '全市', '全部制式', '45.5', '3.25'):
        assert text in summary
    assert '平均间隔：—' in summary and '人口：—' in summary
    assert '周收入：120.25 货币' in summary and '利润：20.25 货币' in summary
    assert all(metric.title + '：' in summary for metric in snap.metrics)
    assert all(title in summary for title in ('当日最大客流线路', '当日最小客流线路',
                                               '当日最多班次线路', '当日最少班次线路'))
    assert summary.count('范围：') == 1
    assert '公司当日' not in summary and '没有运行日记录' not in summary
    assert '缺测显示为—，不按零值统计。' not in summary

    snap.population = 0
    snap.metrics[0].value = 0
    snap.metrics[4].value = Decimal('120.250001')
    snap.metrics[4].complete = False
    values = exports().build_share_summary(snap)
    assert '人口：0' in values and '线路数：0 条' in values
    assert '周收入：120.250001 货币（部分观测）' in values
    assert '平均间隔：—' in values

    snap.mode = '综合'
    comprehensive = exports().build_share_summary(snap)
    assert next(line for line in comprehensive.splitlines() if line.startswith('换乘系数：')) == '换乘系数：3.25'
    assert '公共交通分担率：45.5 %（全市）' in comprehensive


def test_xlsx_exports_13_metrics_as_numbers_with_scope_reason_and_completeness(tmp_path):
    wb = workbook(tmp_path)
    assert wb.sheetnames == ['范围与说明', '首页指标', '线路极值', '客流前十', '班次分类', '关键提醒']
    rows = list(wb['首页指标'].values)
    assert len(rows) == 14
    by_title = {row[0]: row for row in rows[1:]}
    assert by_title['公共交通分担率'][1:4] == (45.5, '%', '全市当日')
    assert by_title['换乘系数'][1] == 3.25
    assert by_title['平均间隔'][1] is None
    assert '没有运行日记录' in by_title['平均间隔']
    assert '缺测' in by_title['平均间隔']
    assert by_title['周收入'][1] == 120.25
    assert isinstance(by_title['线路数'][1], int)
    assert len(list(wb['线路极值'].values)) == 5
    assert len(list(wb['客流前十'].values)) == 11
    assert list(wb['客流前十'].values)[-1][1:3] == ('线路-10', '长名称线路10')


def test_xlsx_keeps_all_mode_rows_total_zero_and_missing_share_distinct(tmp_path):
    wb = workbook(tmp_path)
    rows = list(wb['班次分类'].values)
    assert [r[0] for r in rows[1:]] == ['公交', '有轨电车', '无轨电车', '地铁', '水上巴士', '其他', '总计']
    assert rows[1][1:3] == (30, 0.6)
    assert rows[5][1:3] == (0, 0)
    assert rows[6][1:3] == (None, None)
    assert rows[7][1] == 50
    assert wb['班次分类']['C2'].number_format == '0.0%'


def test_alert_export_preserves_company_values_both_windows_and_original_reason(tmp_path):
    wb = workbook(tmp_path, alerts=sample_alerts())
    rows = list(wb['关键提醒'].values)
    assert len(rows) == 2
    row = dict(zip(rows[0], rows[1]))
    assert row['指标'] == '现金流'
    assert row['公司标识'] == 'p2'
    assert row['公司'] == '同名公司'
    assert row['原提醒原因'] == '现金流转负'
    assert row['本期值'] == -3.5 and row['对比值'] == 10.25
    assert row['本期开始'] == D(2024, 3, 2) and row['本期结束（不含）'] == D(2024, 3, 3)
    assert row['对比开始'] == D(2024, 3, 1) and row['对比结束（不含）'] == D(2024, 3, 2)


@pytest.mark.parametrize('companies, expected', [
    ((), '无公司选择（仅全市指标）'),
    (('p2',), '同名公司 [p2]'),
    (('p2', 'p1'), '同名公司 [p2]、同名公司 [p1]'),
])
def test_alert_company_scope_preserves_empty_selection_and_stable_ids(tmp_path, companies, expected):
    # Empty DashboardResult selection excludes company metrics, unlike empty home scope.
    alerts = sample_alerts([])
    alerts = replace(alerts, filters=replace(alerts.filters, companies=companies))
    wb = workbook(tmp_path, alerts=alerts)
    context = {row[0]: row[1] for row in list(wb['范围与说明'].values)[1:]}
    assert context['提醒公司'] == expected
    assert context['公司'] == '同名公司 [p2]'
    assert len(list(wb['关键提醒'].values)) == 1


def test_unavailable_or_empty_alerts_have_no_fabricated_alert_rows(tmp_path):
    for alert_snapshot in (None, sample_alerts([])):
        wb = workbook(tmp_path, alerts=alert_snapshot)
        assert len(list(wb['关键提醒'].values)) == 1
        notes = '\n'.join(str(value) for row in wb['范围与说明'].values for value in row if value is not None)
        assert '缺测' in notes
        if alert_snapshot is None:
            assert '提醒未计算' in notes
        else:
            assert '无达标提醒' in notes
            assert '2024-03-02' in notes and '2024-03-01' in notes


def test_empty_line_scope_keeps_four_missing_highlights_without_zero_filling(tmp_path):
    snap = sample_snapshot(lines=(), passenger_top10=(), departure_modes=(), total_departures=None,
                           highlights=tuple(NS(title=t, line=None) for t in
                               ('当日最大客流线路', '当日最小客流线路', '当日最多班次线路', '当日最少班次线路')))
    wb = workbook(tmp_path, snap)
    assert len(list(wb['客流前十'].values)) == 1
    assert all(row[1] is None for row in list(wb['线路极值'].values)[1:])
    assert list(wb['班次分类'].values)[-1][1:3] == (None, None)


@pytest.mark.parametrize('value', [float('nan'), float('inf'), Decimal('NaN'), Decimal('Infinity')])
def test_nonfinite_values_export_blank_without_excel_errors(tmp_path, value):
    snap = sample_snapshot(population=value, total_departures=value)
    snap.metrics[0].value = value
    wb = workbook(tmp_path, snap)
    assert list(wb['首页指标'].values)[1][1] is None
    assert list(wb['班次分类'].values)[-1][1] is None
    assert not any(cell.data_type == 'e' for sheet in wb for row in sheet for cell in row)


def test_source_text_cannot_become_an_excel_formula(tmp_path):
    snap = sample_snapshot(city_name='=1+1', save_name='=HYPERLINK("bad")')
    snap.lines[0].name = '=2+2'
    wb = workbook(tmp_path, snap)
    assert not any(cell.data_type == 'f' for sheet in wb for row in sheet for cell in row)
    assert wb['范围与说明']['B2'].value == '=1+1'
    assert list(wb['客流前十'].values)[1][2] == '=2+2'


def test_exports_import_and_xlsx_do_not_import_qt(tmp_path):
    exports()
    root = Path(__file__).resolve().parents[1]
    script = '''
import importlib.abc, sys
class NoQt(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.startswith(('PySide6', 'qfluentwidgets', 'stats_exports')):
            raise AssertionError('纯导出不应加载Qt: ' + fullname)
sys.meta_path.insert(0, NoQt())
from test_latest_info_exports import sample_snapshot
from latest_info_exports import build_share_summary, export_latest_info_xlsx
assert '秋山市' in build_share_summary(sample_snapshot())
export_latest_info_xlsx(sample_snapshot(), None, sys.argv[1])
'''
    env = dict(os.environ, PYTHONPATH=os.pathsep.join((str(root/'src'), str(root/'frontend'))))
    result = subprocess.run([sys.executable, '-X', 'utf8', '-c', script, str(tmp_path/'纯导出.xlsx')],
                            env=env, capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
