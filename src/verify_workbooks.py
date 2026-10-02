from openpyxl import load_workbook
from pathlib import Path
p=Path('exports/CIM2_线路发班整理_quicksave.76561198362520556-76561198845688243_运行时.xlsx')
wb=load_workbook(p,read_only=True,data_only=True)
print('line workbook sheets',len(wb.sheetnames),wb.sheetnames[:5],wb.sheetnames[-3:])
ws=wb['线路信息']; rows=list(ws.iter_rows(values_only=True)); print('line info rows',len(rows),'row1',rows[0][:4]); hdr=next(r for r in rows if r and '线路号' in r); print('headers',hdr); ix={v:i for i,v in enumerate(hdr)}
length_key = '折算里程' if '折算里程' in ix else ('线路长度' if '线路长度' in ix else None)
for row in rows:
 if row and row[ix['线路号']] in [201,301,327,901,107]:
  length = row[ix[length_key]] if length_key else None
  print('line',row[ix['线路号']],length,row[ix['单程时间']],row[ix['理论最大车辆需求数']],row[ix['开线日期']],row[ix['最近改线日期']])
q=Path('exports/CIM2_公司信息整理_quicksave.76561198362520556-76561198845688243_运行时.xlsx'); cw=load_workbook(q,read_only=True,data_only=True); print('company sheets',cw.sheetnames)
for sn in cw.sheetnames:
 print(sn,list(cw[sn].iter_rows(min_row=1,max_row=4,values_only=True)))
