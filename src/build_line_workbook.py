from __future__ import annotations

from display_rules import format_line_name, store_text_literally
from parse_events import ProgressReporter

import csv
import os
import re
import sys
from datetime import datetime
from pathlib import Path

# Use the stdlib/et_xmlfile writer in both source and frozen runs.  The lxml
# incremental writer can raise IO_WRITE from a PyInstaller one-file temp
# directory when a large workbook is serialized.
import openpyxl
openpyxl.LXML = False
import openpyxl.xml.functions as _oxml
from et_xmlfile import xmlfile as _stdlib_xmlfile
_oxml.LXML = False
_oxml.xmlfile = _stdlib_xmlfile
from openpyxl import Workbook

PROJECT = Path(__file__).resolve().parents[1]
EXPORT = Path(os.environ.get("CIM2_EXPORT_DIR", str(PROJECT / "exports")))
TAG = sys.argv[1] if len(sys.argv) > 1 else "运行时"

LINE_MONEY_SCALE = 102400
LINE_LENGTH_SCALE = 1024000
MAP_LENGTH_MULTIPLIER = 2.0

COMPANY_NAMES = {
    "826272703's Company": "六进公交",
    "jeremylin2005's Company": "八连交通集团",
}


def company_name(value: str) -> str:
    return COMPANY_NAMES.get(value, value)


def garage_name(company: str, mode: str, value: str) -> str:
    """Render the serialized depot index as the game's readable depot name."""
    try:
        number = int(float(value or 0))
    except (TypeError, ValueError):
        number = 0
    if number <= 0:
        return ""
    prefix = company_name(company).removesuffix("集团")
    return f"{prefix}{mode}车库{number}"


def read_csv(name: str) -> list[dict[str, str]]:
    path = EXPORT / f"{name}_{TAG}.csv"
    if not path.exists():
        path = EXPORT / "reanalysis" / f"{name}_{TAG}.csv"
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    # Older exports may lack the newly decoded line metadata.  Prefer the
    # reanalysis copy when it contains the richer schema.
    if name == "CIM2_线路客流_完整导出" and "开线日期" not in (rows[0] if rows else {}):
        fallback = EXPORT / "reanalysis" / f"{name}_{TAG}.csv"
        if fallback.exists():
            with fallback.open(encoding="utf-8-sig", newline="") as f:
                rows = list(csv.DictReader(f))
    return rows


def parse_datetime(value: str) -> datetime | None:
    value = (value or "").strip()
    for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    return None


def number(value: str) -> float:
    try:
        return float(value or 0)
    except ValueError:
        return 0.0


def integer(value: str) -> int:
    try:
        return int(float(value or 0))
    except ValueError:
        return 0


def safe_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "线路"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{base[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def date_groups(mask: int) -> list[str]:
    """Expand a game's weekday bitmask into the requested reporting groups."""
    groups = []
    # CIM2 uses Sunday=bit 0 through Saturday=bit 6.
    if mask & 0b0011110:  # Monday-Thursday
        groups.append("周一至周四")
    if mask & 0b0100000:  # Friday
        groups.append("周五")
    if mask & 0b1000000:  # Saturday
        groups.append("周六")
    if mask & 0b0000001:  # Sunday
        groups.append("周日")
    if mask == 0:
        groups.append("未启用")
    return groups or ["其他"]


def mask_label(mask: int) -> str:
    labels = {0: "未启用", 30: "周一至周四", 32: "周五", 62: "周一至周五", 65: "周六/周日", 127: "每天"}
    return labels.get(mask, f"掩码{mask}")


def preferred_vehicle(value: str) -> str:
    return {"0": "任意", "1": "小型", "2": "中型", "3": "大型"}.get(str(value), str(value))


MODE_ORDER = ["公交", "单轨列车", "地铁", "无轨电车", "有轨电车", "水上巴士", "综合"]


def current_day_bit(day: datetime) -> int:
    return 1 << ((day.weekday() + 1) % 7)


def minutes_from_ticks(value: str) -> float:
    return integer(value) / 10_000_000 / 60


def clock_minutes(value: str) -> str:
    """Render a timetable clock as HH:MM, hiding serialized seconds."""
    value = (value or "").strip()
    if not value:
        return ""
    match = re.match(r"^(\d{1,2}:\d{2})", value)
    return match.group(1) if match else value


def line_metrics(
    line: dict[str, str], departures: list[dict[str, str]], day_bit: int,
    duration_minutes: float, length_km: float,
) -> dict[str, float | int | str]:
    day_departures = [r for r in departures if integer(r.get("时刻表_运行日掩码", "0")) & day_bit]
    count = len(day_departures)
    passenger = integer(line.get("客流_今日", "0"))
    total_minutes = round(count * duration_minutes, 2)
    total_vehicle_km = count * length_km
    return {
        "当日发班数": count,
        "行车总时间": total_minutes,
        "行车总里程": round(total_vehicle_km, 6),
        "今日平均单班人次": round(passenger / count, 2) if count else 0,
        "今日平均车公里人次": round(passenger / total_vehicle_km, 2) if total_vehicle_km else 0,
    }


def main() -> None:
    reporter = ProgressReporter()
    reporter.start('line_workbook')
    company_info = read_csv("CIM2_公司信息")
    fleet_by_company = {
        company_name(r.get("公司名称", "")): integer(r.get("车辆总数", "0"))
        for r in company_info
    }
    fleet_by_company_mode = {
        (company_name(r.get("公司名称", "")), mode): integer(r.get(f"车辆总数_{mode}", "0"))
        for r in company_info
        for mode in MODE_ORDER[:-1]
    }
    lines = read_csv("CIM2_线路客流_完整导出")
    timetables = read_csv("CIM2_发班信息_完整导出")
    departures = read_csv("CIM2_发班记录_完整导出")
    metadata = read_csv("CIM2_城市历史元数据")[0]
    current = parse_datetime(metadata.get("模拟当前时间", "")) or datetime.now()
    simulation_start = parse_datetime(metadata.get("模拟开始", "")) or current
    try:
        current_tick = int(float(metadata.get("模拟当前时间_tick", "0")))
    except (TypeError, ValueError):
        current_tick = int((current - datetime(1, 1, 1)).total_seconds() * 10_000_000)
    current_date = current.strftime("%Y-%m-%d")
    current_time = current.strftime("%H:%M:%S")
    day_bit = current_day_bit(current)

    timetable_by_key = {
        (r["公司序号"], r["线路类型"], r["线路号"], r["时刻表序号"]): r
        for r in timetables
    }
    first_timetable_by_line: dict[tuple[str, str, str], dict[str, str]] = {}
    for timetable in timetables:
        line_key = (timetable["公司序号"], timetable["线路类型"], timetable["线路号"])
        current_first = first_timetable_by_line.get(line_key)
        if current_first is None or integer(timetable.get("时刻表序号", "0")) < integer(current_first.get("时刻表序号", "0")):
            first_timetable_by_line[line_key] = timetable
    departures_by_line: dict[tuple[str, str, str], list[dict[str, str]]] = {}
    for departure in departures:
        line_key = (departure["公司序号"], departure["线路类型"], departure["线路号"])
        timetable_key = line_key + (departure["时刻表序号"],)
        row = dict(departure)
        row.update({f"时刻表_{k}": v for k, v in timetable_by_key.get(timetable_key, {}).items()})
        departures_by_line.setdefault(line_key, []).append(row)
    for rows in departures_by_line.values():
        rows.sort(key=lambda r: (integer(r.get("发班_tick", "0")), integer(r.get("时刻表序号", "0")), integer(r.get("发班序号", "0"))))

    wb = Workbook()
    used_names = set()
    line_info_headers = [
        "公司名称", "运输制式", "线路号", "线路名称", "开线日期", "线路车库", "最近改线日期", "单程时间", "地图里程", "折算里程", "核定速度", "站点数",
        "当日发班数", "行车总时间", "理论最大车辆需求数", "累计客流", "平均客流", "今日客流",
        "今日平均单班人次", "今日平均车公里人次", "当前配车数", "每周收入", "每周支出",
        "周一至周四发班数量", "周五发班数量", "周六发班数量", "周日发班数量",
    ]
    line_info_rows = []
    overview_rows = []
    sheet_map = {}
    for line in lines:
        key = (line["公司序号"], line["线路类型"], line["线路号"])
        first_timetable = first_timetable_by_line.get(key, {})
        # Do not substitute the simulation start when a save lacks the line's
        # own build timestamp: that would silently turn an unknown opening
        # date into 2013-04-01 08:00:00.  The extractor obtains this value from
        # LineData's inherited ObjectData.m_buildTime.
        opened = parse_datetime(line.get("开线日期", ""))
        edited = parse_datetime(line.get("最近改线日期", ""))
        try:
            opened_tick = int(float(line.get("开线日期_tick", "0")))
        except (TypeError, ValueError):
            opened_tick = 0
        elapsed_shift = max((current_tick - opened_tick) >> 10, 1) if opened_tick > 0 else None
        sheet_name = safe_sheet_name(f"{line['线路类型']}{format_line_name(line['线路号'], line.get('线路名称', ''))}", used_names)
        sheet_map[key] = sheet_name
        grouped_departures = {group: [] for group in ("周一至周四", "周五", "周六", "周日")}
        for departure in departures_by_line.get(key, []):
            mask = integer(departure.get("时刻表_运行日掩码", "0"))
            for group in date_groups(mask):
                if group in grouped_departures:
                    grouped_departures[group].append(departure)
        for rows in grouped_departures.values():
            rows.sort(key=lambda r: (integer(r.get("发班_tick", "0")), integer(r.get("时刻表序号", "0")), integer(r.get("发班序号", "0"))))
        cumulative = integer(line.get("客流_累计", "0"))
        income = integer(line.get("收入_累计", "0"))
        expenses = integer(line.get("支出_累计", "0"))
        duration_minutes = minutes_from_ticks(line.get("单程时间_tick", "0"))
        # Both values are fixed-point integers in the extraction CSV.  The
        # raw map mileage is the pre-conversion road length; the converted
        # value is what CIM2 displays after expressway/junction adjustments.
        map_length_raw = number(line.get("地图里程", "0"))
        converted_length_raw = number(line.get("折算里程", line.get("线路长度", "0")))
        # Reporting convention doubles the pre-conversion map mileage.  The
        # game's converted/display mileage remains unchanged.
        map_length_km = (map_length_raw / LINE_LENGTH_SCALE) * MAP_LENGTH_MULTIPLIER
        converted_length_km = converted_length_raw / LINE_LENGTH_SCALE
        metrics = line_metrics(line, departures_by_line.get(key, []), day_bit, duration_minutes, map_length_km)
        mode = line.get("线路类型", "").replace("线路", "")
        display_company = company_name(line.get("公司名称", ""))
        line_number = integer(line.get("线路号", "0"))
        info_row = [
            display_company, mode, line_number,
            format_line_name(line_number, line.get("线路名称", "")), opened.strftime("%Y-%m-%d %H:%M:%S") if opened else "",
            garage_name(line.get("公司名称", ""), mode, line.get("线路车库", "0")),
            edited.strftime("%Y-%m-%d %H:%M:%S") if edited else "",
            round(duration_minutes, 2) if duration_minutes else "", round(map_length_km, 6),
            round(converted_length_km, 6),
            round(map_length_km / duration_minutes * 60, 2) if duration_minutes else "",
            integer(line.get("站点数", "0")),
            metrics["当日发班数"], metrics["行车总时间"], integer(line.get("最大配车数", "0")),
            cumulative,
            # InfoTool.FormatObjectInfo performs integer division first,
            # converts to Int32, then FormatNumberFixed applies the 10-bit
            # fixed-point shift.  Keep that order so displayed averages do
            # not depend on Python floating-point rounding.
            round(((cumulative * 864000000000) // elapsed_shift) / 1024, 2) if elapsed_shift else "",
            # Keep the direct LineData.m_transportedToday value in the
            # "今日客流" column before the derived per-departure metrics.
            integer(line.get("客流_今日", "0")),
            metrics["今日平均单班人次"], metrics["今日平均车公里人次"], integer(line.get("配车数", "0")),
            round(income / LINE_MONEY_SCALE, 2), round(expenses / LINE_MONEY_SCALE, 2),
            len(grouped_departures["周一至周四"]), len(grouped_departures["周五"]), len(grouped_departures["周六"]), len(grouped_departures["周日"]),
        ]
        line_info_rows.append(info_row)
        overview_rows.append({
            "公司名称": display_company, "运输制式": mode, "线路号": line_number, "线路名称": format_line_name(line_number, line.get("线路名称", "")),
            "今日客流": integer(line.get("客流_今日", "0")), "当日发班数": metrics["当日发班数"],
            "行车总时间": metrics["行车总时间"], "理论最大车辆需求数": integer(line.get("最大配车数", "0")),
            "当前配车数": integer(line.get("配车数", "0")), "平均单班人次": metrics["今日平均单班人次"],
            "行车总里程": metrics["行车总里程"], "平均车公里人次": metrics["今日平均车公里人次"],
        })

    info = wb.active
    info.title = "线路信息"
    info.append(["当前日期", current_date, "当前时间", current_time])
    info.append([])
    info.append(line_info_headers)
    for row in line_info_rows:
        info.append(row)

    overview = wb.create_sheet("公司概览", 1)
    overview.append(["当前日期", current_date, "当前时间", current_time])
    overview.append([])
    overview_headers = [
        "运输制式", "当日总客流", "当日发班数", "行车总时间", "理论最大车辆需求数", "配车总数", "车辆周转率",
        "平均单班人次", "平均车公里人次", "当日最大客流线路", "当日最大客流", "线路平均单班人次", "线路平均车公里人次",
        "当日最多班次线路", "最多班次", "线路平均单班人次", "线路平均车公里人次",
        "当日最小客流线路", "当日最小客流", "线路平均单班人次", "线路平均车公里人次",
        "当日最少班次线路", "最少班次", "线路平均单班人次", "线路平均车公里人次",
    ]
    for company_index, company in enumerate(dict.fromkeys(r["公司名称"] for r in overview_rows)):
        overview.append([company])
        overview.append(overview_headers)
        company_rows = [r for r in overview_rows if r["公司名称"] == company]
        for mode in MODE_ORDER:
            mode_rows = company_rows if mode == "综合" else [r for r in company_rows if r["运输制式"] == mode]
            total_passenger = sum(int(r["今日客流"]) for r in mode_rows)
            total_departures = sum(int(r["当日发班数"]) for r in mode_rows)
            total_minutes = round(sum(float(r["行车总时间"]) for r in mode_rows), 2)
            total_vehicle_km = sum(float(r["行车总里程"]) for r in mode_rows)
            max_vehicles = sum(int(r["理论最大车辆需求数"]) for r in mode_rows)
            # Mode rows show vehicles currently assigned to that mode.  The
            # aggregate row uses every vehicle in the company's depots,
            # including parked/unassigned vehicles omitted from LineData.
            vehicles = fleet_by_company.get(company, 0) if mode == "综合" else fleet_by_company_mode.get(
                (company, mode), sum(int(r["当前配车数"]) for r in mode_rows)
            )
            max_passenger = max(mode_rows, key=lambda r: (int(r["今日客流"]), -int(r["线路号"])), default=None)
            max_departures = max(mode_rows, key=lambda r: (int(r["当日发班数"]), -int(r["线路号"])), default=None)
            min_passenger = min(mode_rows, key=lambda r: (int(r["今日客流"]), int(r["线路号"])), default=None)
            min_departures = min(mode_rows, key=lambda r: (int(r["当日发班数"]), int(r["线路号"])), default=None)
            values = [
                mode, total_passenger, total_departures, total_minutes, max_vehicles, vehicles,
                round(total_departures / vehicles, 2) if vehicles else 0,
                round(total_passenger / total_departures, 2) if total_departures else 0,
                round(total_passenger / total_vehicle_km, 2) if total_vehicle_km else 0,
            ]
            for selected, metric in ((max_passenger, "今日客流"), (max_departures, "当日发班数"), (min_passenger, "今日客流"), (min_departures, "当日发班数")):
                values.extend([
                    selected["线路名称"] if selected else "", selected[metric] if selected else 0,
                    selected["平均单班人次"] if selected else 0, selected["平均车公里人次"] if selected else 0,
                ])
            overview.append(values)
        if company_index < len(set(r["公司名称"] for r in overview_rows)) - 1:
            overview.append([])
    reporter.progress('line_workbook', 2, len(lines) + 3)
    for line_index, line in enumerate(lines):
        key = (line["公司序号"], line["线路类型"], line["线路号"])
        ws = wb.create_sheet(sheet_map[key])
        ws.append(["当前日期", current_date, "当前时间", current_time])
        ws.append(["公司名称", company_name(line.get("公司名称", "")), "运输制式", line.get("线路类型", ""), "线路名称", format_line_name(line.get("线路号", "0"), line.get("线路名称", ""))])
        grouped_departures = {group: [] for group in ("周一至周四", "周五", "周六", "周日")}
        for row in departures_by_line.get(key, []):
            for group in date_groups(integer(row.get("时刻表_运行日掩码", "0"))):
                if group in grouped_departures:
                    grouped_departures[group].append(row)
        for rows in grouped_departures.values():
            rows.sort(key=lambda r: (integer(r.get("发班_tick", "0")), integer(r.get("时刻表序号", "0")), integer(r.get("发班序号", "0"))))
        ws.append(["周一至周四发班数量", len(grouped_departures["周一至周四"]), "周五发班数量", len(grouped_departures["周五"]), "周六发班数量", len(grouped_departures["周六"]), "周日发班数量", len(grouped_departures["周日"])])
        headers = []
        for group in ("周一至周四", "周五", "周六", "周日"):
            headers.extend([f"{group}发班时间", f"{group}首选车型"])
        ws.append(headers)
        max_rows = max((len(rows) for rows in grouped_departures.values()), default=0)
        for idx in range(max_rows):
            values = []
            for group in ("周一至周四", "周五", "周六", "周日"):
                row = grouped_departures[group][idx] if idx < len(grouped_departures[group]) else None
                values.extend([
                    clock_minutes(row.get("发班时间", "")) if row else "",
                    preferred_vehicle(row.get("时刻表_首选车型", "0")) if row else "",
                ])
            ws.append(values)
        reporter.progress('line_workbook', line_index + 3, len(lines) + 3)

    out = EXPORT / f"CIM2_线路发班整理_{TAG}.xlsx"
    # Keep the workbook unstyled while sizing columns to their contents.
    for ws in wb.worksheets:
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                widths[cell.column] = max(widths.get(cell.column, 0), len(value))
        for column, width in widths.items():
            ws.column_dimensions[ws.cell(1, column).column_letter].width = min(max(width + 2, 10), 32)
    store_text_literally(wb)
    wb.save(out)
    reporter.progress('line_workbook', len(lines) + 3, len(lines) + 3)
    reporter.finish('line_workbook')
    print(out)
    print(f"line_sheets={len(lines)} departure_rows={len(departures)}")


if __name__ == "__main__":
    main()
