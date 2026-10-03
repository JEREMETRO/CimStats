"""Regression gate for the two existing tabs during the Fluent migration."""
import copy
import os
import time
from pathlib import Path
from decimal import Decimal

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import pytest
from PySide6.QtCore import QCoreApplication, QEvent, QObject, QSettings, Qt, Signal
from PySide6.QtWidgets import QApplication, QLabel
import desktop_app as desktop
from report_model import DAY_GROUPS, load_session


def wait_home(window):
    from test_latest_info_integration import wait_for
    wait_for(QApplication.instance(), lambda: window.latest_info_controller.snapshot is not None)


def metric(window, key):
    label = window.latest_info_page.findChild(QLabel, f'metric-{key}-value')
    assert label is not None
    text = label.text().replace(',', '')
    return None if text == '—' else Decimal(text)


@pytest.fixture
def window(tmp_path, monkeypatch):
    application = QApplication.instance() or QApplication([])
    assert application.platformName() == 'offscreen'
    settings = QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat)
    monkeypatch.setattr(desktop, 'QSettings', lambda *args: settings)
    monkeypatch.setattr(desktop.MainWindow, 'check_install', lambda self: None)
    widget = desktop.MainWindow()
    data = load_session(Path(__file__).resolve().parents[1] / 'exports', '望春市_test_运行时')
    if not data['lines']:
        pytest.skip('Local exported-save fixture is unavailable')
    data.update(history=[], save_path='望春市.save', save_key='regression')
    widget.on_completed(data)
    wait_home(widget)
    yield widget
    widget.close()
    widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    application.processEvents()


def test_overview_filters_metrics_and_navigation(window):
    assert metric(window, 'line-count') == len(window.data['lines'])
    for index in range(window.mode_combo.count()):
        window.mode_combo.setCurrentIndex(index)
        wait_home(window)
        mode = window.mode_combo.currentData()
        rows = [row for row in window.data['lines'] if mode == '综合' or row['运输制式'] == mode]
        assert metric(window, 'line-count') == len(rows)
        assert float(metric(window, 'weekly-income')) == pytest.approx(sum(row['每周收入'] for row in rows), abs=.005)
    for index in (2, 1, 0):
        window.navigate(index)
        assert window.pages.currentIndex() == index
        assert window.export_line_button.isHidden() == (index == 2)


def test_navigation_click_signal_opens_its_own_page(window):
    for index in (2, 0, 1):
        window.nav_buttons[index].clicked.emit(True)
        assert window.pages.currentIndex() == index
        assert window.page_title.text() == ('最新信息', '线路查询', '统计数据')[index]


def test_rebuilt_legacy_rows_are_hidden_before_deferred_deletion(window):
    window.show()
    QApplication.processEvents()
    panel = window.latest_info_page.passengers
    previous = panel._mode_widgets[0]
    assert not previous.isHidden()
    panel.clear()
    assert previous.isHidden()
    window.line_clicked(0, 0)
    cards = list(window.fact_cards)
    assert cards
    window.clear_fact_cards()
    assert all(card.isHidden() for card in cards)


def test_overview_ranking_pie_and_extreme_interaction(window):
    page = window.latest_info_page
    page.passengers.mode_combo.setCurrentIndex(1)
    assert page.passengers.mode_combo.currentData() != '综合'
    page.passengers.return_button.click()
    assert page.passengers.mode_combo.currentData() == '综合'
    page.departures.ranking_button.click()
    assert not page.departures.mode_combo.isHidden()
    page.departures.return_button.click()
    assert page.departures.capture_state()['view'] == 'structure'
    assert page.departures.mode_combo.currentData() == '综合'
    assert not page.departures.mode_combo.isEnabled()
    for card, highlight in zip(page.highlights, window.latest_info_controller.snapshot.highlights):
        assert len(card.values) == 4
        if highlight.line is not None:
            card.line_requested.emit(highlight.line.key)
            assert window.pages.currentIndex() == 1
            assert window._selected_line['key'] == highlight.line.key


def test_line_search_sort_details_and_fields(window):
    window.navigate(1)
    row = window.data['lines'][0]
    window.query.setText(row['线路名称'])
    assert window.line_table.rowCount() == len(window.filtered_lines())
    assert window.line_table.rowCount() > 0
    window.line_table.sortItems(7, Qt.SortOrder.DescendingOrder)
    key = window.line_table.item(0, 0).data(Qt.ItemDataRole.UserRole)
    window.line_clicked(0, 0)
    assert window._selected_line['key'] == key
    selected = window._selected_line
    assert window.schedule_panel.day_groups == tuple(selected['显示日组'])
    for day in window.schedule_panel.day_groups:
        entries = selected['班次'][day]
        window.schedule_panel.set_current_group(day)
        assert window.schedule_panel.summary['count'] == len(entries)
        assert len(window.schedule_panel.matrix.entries) == len(entries)
    assert '站点客流' not in window.schedule_panel.day_groups
    window.line_columns_menu.actions()[3].setChecked(False)
    assert window.line_table.isColumnHidden(3)
    window.line_columns_menu.actions()[3].setChecked(True)
    assert not window.line_table.isColumnHidden(3)
    window.fact_menu.actions()[0].setChecked(False)
    assert window.fact_cards[0].isHidden()
    window.fact_menu.actions()[0].setChecked(True)
    assert window.fact_cards[0].label.text() == '线路车库'
    window.query.setText('no-such-route')
    assert window.line_table.rowCount() == 0
    window.query.clear()
    assert window.line_table.rowCount() == len(window.data['lines'])


def test_existing_xlsx_exports_keep_identical_bytes(window, tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    for kind in ('line_workbook', 'company_workbook'):
        source, target = tmp_path / (kind + '.xlsx'), tmp_path / ('copy-' + kind + '.xlsx')
        book = Workbook()
        book.active.append(['原有工作簿', 123])
        book.save(source)
        window.data.setdefault('outputs', {})[kind] = source
        monkeypatch.setattr(desktop.QFileDialog, 'getSaveFileName', lambda *a: (str(target), ''))
        window.export_file(kind)
        assert target.read_bytes() == source.read_bytes()
        assert load_workbook(target).active['B1'].value == 123


def test_same_named_companies_filter_by_identity_in_both_existing_tabs(window):
    replacement = copy.deepcopy(window.data)
    template_company, template_line = replacement['companies'][0], replacement['lines'][0]
    replacement['companies'], replacement['lines'] = [], []
    for owner, passengers in (('a', 10), ('b', 20)):
        company, line = copy.deepcopy(template_company), copy.deepcopy(template_line)
        company.update(公司名称='同名公司', 公司标识=owner)
        line.update(公司名称='同名公司', 公司标识=owner, 今日客流=passengers,
                    key=owner + '|bus|1')
        replacement['companies'].append(company)
        replacement['lines'].append(line)
    window.on_completed(replacement)
    wait_home(window)
    for owner in ('a', 'b'):
        index = window.company_combo.findData(owner)
        assert index >= 0
        window.company_combo.setCurrentIndex(index)
        wait_home(window)
        assert metric(window, 'line-count') == 1
        window.line_company.setCurrentIndex(window.line_company.findData(owner))
        assert window.line_table.rowCount() == 1
        window.line_clicked(0, 0)
        assert window._selected_line['公司标识'] == owner


def test_replacement_import_clears_old_tabs_and_can_cancel(window, monkeypatch, tmp_path):
    class IdleWorker(QObject):
        progress = Signal(int, str)
        log = Signal(str)
        completed = Signal(object)
        failed = Signal(str)
        finished = Signal()
        def __init__(self, *args):
            super().__init__()
            self.cancelled = False
        def start(self): pass
        def isRunning(self): return True
        def cancel(self): self.cancelled = True
        def wait(self, _timeout): return True
    monkeypatch.setattr(desktop, 'locate_managed', lambda *a: tmp_path)
    monkeypatch.setattr(desktop, 'ParseWorker', IdleWorker)
    replacement = copy.deepcopy(window.data)
    from openpyxl import Workbook
    source = tmp_path / 'company.xlsx'
    Workbook().save(source)
    replacement.setdefault('outputs', {})['company_workbook'] = source
    window.navigate(1)
    window.line_clicked(0, 0)
    old_cards = list(window.fact_cards)
    window.start_parse(tmp_path / '中文存档.save')
    assert len(window.data) == 0
    assert window.line_table.rowCount() == 0
    assert not window.schedule_panel.day_groups
    assert window.schedule_panel.summary['count'] == 0
    assert all(card.isHidden() for card in old_cards)
    assert not window.export_line_button.isEnabled()
    assert not window.export_company_button.isEnabled()
    assert window.statistics_page.snapshot is None
    window.cancel_parse()
    assert window.worker.cancelled
    window.worker_finished()
    window.on_completed(replacement)
    assert window.line_table.rowCount() == len(replacement['lines'])
    assert window.export_company_button.isEnabled()


def test_close_cancels_real_parser_thread_and_backend(window, monkeypatch, tmp_path):
    (tmp_path / 'test.save').write_bytes(b'fixture')
    (tmp_path / 'Assembly-CSharp.dll').write_bytes(b'fixture')
    (tmp_path / 'extract_runtime_data.py').write_text(
        'import time\nprint("started", flush=True)\ntime.sleep(20)\n', encoding='utf8')
    monkeypatch.setattr(desktop, 'SRC', tmp_path)
    worker = desktop.ParseWorker(tmp_path / 'test.save', tmp_path, window)
    worker.job_dir = tmp_path
    window.worker = worker
    worker.start()
    try:
        deadline = time.monotonic() + 5
        while worker.process is None and time.monotonic() < deadline:
            QApplication.processEvents()
            time.sleep(.01)
        assert worker.isRunning()
        process = worker.process
        assert process is not None
        window.close()
        assert worker.cancel_requested
        assert not worker.isRunning()
        assert process.poll() is not None
    finally:
        worker.cancel()
        worker.wait(5000)
        window.worker = None


def test_finished_old_parser_cannot_replace_or_delete_new_parser(window, monkeypatch, tmp_path):
    class Parser(QObject):
        progress = Signal(int, str)
        log = Signal(str)
        completed = Signal(dict)
        failed = Signal(str)
        finished = Signal()
        def __init__(self, *args):
            super().__init__()
            self.running = True
        def start(self): pass
        def isRunning(self): return self.running
        def cancel(self): self.running = False
        def wait(self, *args): return not self.running
    stale_data = copy.deepcopy(window.data)
    monkeypatch.setattr(desktop, 'ParseWorker', Parser)
    monkeypatch.setattr(desktop, 'locate_managed', lambda *a: tmp_path)
    window.start_parse(tmp_path / 'old.save')
    old = window.worker
    old.running = False
    window.start_parse(tmp_path / 'new.save')
    current = window.worker
    old.completed.emit(stale_data)
    old.finished.emit()
    assert len(window.data) == 0
    assert window.worker is current
    current.running = False
    window.worker_finished()


@pytest.mark.parametrize('active,expected',[(True,15.0),(False,None)])
def test_overview_interval_excludes_preserved_non_operating_groups(window,active,expected):
    line=copy.deepcopy(window.data['lines'][0])
    line['班次']={
        '未启用':[{'time':'06:00'},{'time':'06:30'}],
        '运行日未知':[{'time':'08:00'},{'time':'09:00'}],
        '周一至周四':[{'time':'07:00'},{'time':'07:15'}] if active else []}
    window.data['lines']=[line]
    window.refresh_company()
    wait_home(window)
    assert metric(window, 'interval') == expected
    assert len(line['班次']['未启用'])==len(line['班次']['运行日未知'])==2
