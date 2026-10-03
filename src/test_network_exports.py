from openpyxl import load_workbook

from network_model import NetworkOptions, build_network_snapshot
from stats_text import label
from test_network_model import NAMES, dashboard, history


def test_network_summary_export_keeps_detail_sheets_and_unavailable_reason(tmp_path):
    from stats_exports import export_xlsx

    rows = [history('transport-by-type', company, 0, amount)
            for company, amount in [('a', 10), ('b', 30)]]
    source = dashboard(rows, grain='day')
    network = build_network_snapshot(source,
        NetworkOptions(mode='overall', facility='stopcount', vehicle='maximum'), NAMES)
    path = tmp_path / 'network.xlsx'
    export_xlsx(source, path, NAMES, network)
    book = load_workbook(path, data_only=True)
    assert '网络摘要' in book.sheetnames
    assert label('company') in book.sheetnames
    sheet = book['网络摘要']
    rows = list(sheet.values)
    assert rows[0][0] == '分析模式'
    assert '指标口径' not in rows[0]
    assert any(row[7] == '最大车辆需求数' and row[11] and row[16] for row in rows[1:])
    assert all(row[0] == '总体' for row in rows[1:])
    assert all(row[1] is not None and row[2] is not None for row in rows[1:])
