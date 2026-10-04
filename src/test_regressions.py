import sys
import zlib
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'frontend'))
from report_model import format_line_name, display_mode, _line_metrics
from save_container import locate_payload


@pytest.mark.parametrize('offset,padding', [(0x26000, 5), (0x27000, 0), (0x26000, 36), (0x2A000, 200)])
def test_payload_without_long_zero_padding(tmp_path, offset, padding):
    payload = b'\xfd\x77\xfd\xc9' + b'valid serialized content' * 100
    compressed = zlib.compress(payload, wbits=-15)
    path = tmp_path / '中文路径.save'
    path.write_bytes(b'x' * (offset - padding) + b'\0' * padding + compressed + b'trailer')
    location, actual = locate_payload(path)
    assert location.offset == offset
    assert location.compressed_size == len(compressed)
    assert actual == payload


@pytest.mark.parametrize('name,expected', [('', '12路'), ('A', '12A'), ('5-C', '5C'), ('101快线', '101快线'), ('机场', '12路机场'), ('旅游专线', '12路·旅游专线')])
def test_route_display(name, expected):
    assert format_line_name(12, name) == expected


@pytest.mark.parametrize('mode,expected', [('bus','公交'), ('trolley','无轨电车'), ('tram','有轨电车'), ('waterbus','水上巴士'), ('metro','地铁'), ('monorail','单轨'), ('ship','水上巴士')])
def test_modes(mode, expected):
    assert display_mode(mode) == expected


def test_reject_invalid_container(tmp_path):
    path = tmp_path / 'bad.save'
    path.write_bytes(b'not a save' * 20000)
    with pytest.raises(ValueError):
        locate_payload(path)


def test_metrics_zero_denominator():
    result = _line_metrics({'客流_今日': '100'}, [], 2, 0)
    assert result['今日平均单班人次'] == 0
    assert result['今日平均车公里人次'] == 0


def test_metrics_use_active_day_and_vehicle_km():
    result = _line_metrics({'客流_今日': '120'}, [
        {'时刻表_运行日掩码': '2'}, {'时刻表_运行日掩码': '2'},
        {'时刻表_运行日掩码': '4'},
    ], 2, 3)
    assert result['今日平均单班人次'] == 60
    assert result['今日平均车公里人次'] == 20


def test_equals_prefixed_save_text_is_not_written_as_formula(tmp_path):
    from openpyxl import Workbook, load_workbook
    from display_rules import store_text_literally
    workbook = Workbook()
    workbook.active.append(['=Express', 5])
    store_text_literally(workbook)
    workbook.save(tmp_path / 'names.xlsx')
    cell = load_workbook(tmp_path / 'names.xlsx').active['A1']
    assert (cell.value, cell.data_type) == ('=Express', 's')


def test_company_workbook_reads_only_its_own_save(tmp_path):
    import csv
    import subprocess
    from openpyxl import load_workbook
    def write(stem, tag, row):
        with (tmp_path / f'{stem}_{tag}.csv').open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.DictWriter(handle, fieldnames=list(row)); writer.writeheader(); writer.writerow(row)
    # '望' sorts before '运', so a CIM2_*_运行时.csv glob picks the other save first.
    for tag, name in (('运行时', 'Own'), ('望春市_运行时', 'Other')):
        write('CIM2_公司信息', tag, {'公司序号': 0, '公司名称': name})
        write('CIM2_公司车型数据', tag, {'公司序号': 0, '公司名称': name, '车型序号': 1, '车型类型': 'bus'})
        write('CIM2_城市历史元数据', tag, {'模拟开始': '2013-04-01 08:00:00', '模拟当前时间': '2013-04-03 10:00:00'})
    script = Path(__file__).with_name('build_company_workbook.py')
    subprocess.run([sys.executable, str(script), '运行时'], check=True, capture_output=True,
                   env={**__import__('os').environ, 'CIM2_EXPORT_DIR': str(tmp_path), 'PYTHONIOENCODING': 'utf-8'})
    workbook = load_workbook(tmp_path / 'CIM2_公司信息整理_运行时.xlsx', read_only=True)
    assert workbook['表1_公司信息']['A2'].value == 'Own'
    workbook.close()
