"""Synthetic event examples verify the native city and company scopes separately."""
from dataclasses import replace
from datetime import datetime as D, timedelta
from decimal import Decimal

import pytest

from dashboard_model import FilterState, build_dashboard
from journey_model import city_public_journeys, overall_transfer
from network_model import NetworkOptions, build_network_snapshot
from statistics_model import HistoryStore, Query, QueryCancelled, summarize_buckets

START = D(2024, 1, 2)
END = START + timedelta(hours=2)
NAMES = {'a': 'A', 'b': 'B'}


def row(metric, owner, value, hour=0, group='Student', divider=0):
    return dict(指标=metric, 公司标识=owner, 分组=group,
                模拟时间=START + timedelta(hours=hour), 值=value, 分母=divider, 当前槽位=False)


def source(rows, ids=('a', 'b'), start=START, end=END):
    return build_dashboard(HistoryStore(rows, END), FilterState(ids, start, end, 'hour'))


def counters(city=1, hour=0):
    return ([row('trip-types', c, 1, hour, group='single-line') for c in ('a', 'b')]
            + [row('transport-by-type', c, 1, hour, group='bus') for c in ('a', 'b')]
            + [row('public-transport', '', city, hour, divider=10)])


def network(snapshot, mode='overall', names=NAMES):
    return build_network_snapshot(snapshot, NetworkOptions(mode=mode, passenger='trip-types'), names)


def amount(result, key):
    return next(v for v in result.summaries[0].values if v.metric_id == key)


@pytest.mark.parametrize('ids', [(), ('a',), ('b',), ('a', 'b')])
def test_overall_one_journey_across_two_companies_ignores_company_selection(ids):
    result = network(source(counters(), ids))
    assert amount(result, 'trip-types').value == 1  # Native percentage would be 10%.
    assert amount(result, 'transfer-coefficient').value == 2
    assert '全市公共交通' in amount(result, 'trip-types').context
    coefficient = next(c.result for c in result.charts if c.key == 'transfer-coefficient')
    assert coefficient.query.companies == () and coefficient.comparison == {}
    assert coefficient.metric.scope == 'city'
    assert coefficient.series[('', '总计')][0].denominator == 1


def test_same_resident_two_separate_journeys_still_count_two():
    result = network(source(counters(city=2)))
    assert amount(result, 'trip-types').value == 2
    assert amount(result, 'transfer-coefficient').value == 1


def test_explicit_company_mode_with_one_selection_keeps_native_participation():
    snapshot = source(counters(city=2), ids=('a',))
    assert amount(network(snapshot), 'trip-types').value == 2
    company = network(snapshot, 'companies')
    assert company.options.mode == 'companies'
    assert amount(company, 'trip-types').value == 1
    assert amount(company, 'transfer-coefficient').value == 1
    compared = network(source(counters()), 'companies')
    assert [next(v.value for v in s.values if v.metric_id == 'trip-types')
            for s in compared.summaries] == [1, 1]


def test_half_open_interval_missing_city_and_recorded_zero():
    result = network(source(counters() + counters(city=2, hour=1) + counters(city=99, hour=2)))
    assert amount(result, 'trip-types').value == 3
    missing = network(source([r for r in counters() if r['指标'] != 'public-transport']))
    assert amount(missing, 'trip-types').value is None
    assert amount(missing, 'transfer-coefficient').value is None
    assert amount(missing, 'trip-types').reason == '公共交通行程数据不完整'
    zero = network(source(counters(city=0)))
    assert amount(zero, 'trip-types').value == 0
    assert amount(zero, 'transfer-coefficient').value is None


def test_zero_denominator_hour_keeps_boarding_numerator_for_network_period():
    snapshot = source(counters(city=0) + counters(city=1, hour=1))
    assert amount(network(snapshot), 'transfer-coefficient').value == 4
    # Latest-information policy still excludes individual zero-denominator hours.
    paired, _ = overall_transfer(snapshot, positive_hours_only=True)
    assert summarize_buckets(paired.series[('', '总计')], paired.metric) == 2


def test_missing_company_or_serialized_category_never_silently_shortens_city_scope():
    rows = [r for r in counters() if not (r['指标'] == 'transport-by-type' and r['公司标识'] == 'b')]
    result = network(source(rows, ('a',)))
    assert amount(result, 'trip-types').value == 1
    assert amount(result, 'transfer-coefficient').value is None
    rows = counters() + counters(hour=1)
    rows.append(row('transport-by-type', 'b', 7, hour=1, group='tram'))
    result = network(source(rows))
    assert amount(result, 'transfer-coefficient').value == 9  # Only complete 01:00 pairs.
    assert not amount(result, 'transfer-coefficient').complete


def test_missing_city_category_preserves_gap_instead_of_zero_or_percentage():
    rows = counters() + counters(hour=1) + [row('public-transport', '', 3, hour=1, group='Tourist', divider=10)]
    result = network(source(rows))
    assert amount(result, 'trip-types').value == 4
    assert not amount(result, 'trip-types').complete
    assert amount(result, 'transfer-coefficient').value == Decimal(2) / 4


def test_duplicate_category_is_invalid_not_two_journey_events():
    snapshot = source(counters())
    total = snapshot.results['public-transport']
    bucket = total.series[('', '总计')][0]
    corrupted = replace(total, series={('', '总计'): [replace(bucket, raw=bucket.raw * 2)]})
    result = city_public_journeys(corrupted, ['Student'])
    assert result.series[('', '总计')][0].value is None


def test_changed_players_or_obsolete_provenance_marker_cannot_change_city_count():
    rows = counters(city=7)
    rows[0]['journey_company_scope'] = {'company_ids': ['former-player'], 'complete': False}
    result = network(source(rows, ('a',)), names={'a': 'Current player'})
    assert amount(result, 'trip-types').value == 7
    store = HistoryStore([row('public-transport', '', 7, divider=10)], END)
    total = store.query(Query('public-transport', (), '__total__', START, END, 'hour'))
    city = city_public_journeys(total, ['Student'])
    assert city.metric.scope == 'city'
    assert summarize_buckets(city.series[('', '总计')], city.metric) == 7


def test_identical_saved_counters_cannot_determine_subset_union():
    from itertools import chain
    left = [('a', 'b'), ('c',)]
    right = [('a', 'c'), ('b',)]
    def saved(events):
        return ({c: sum(c in e for e in events) for c in set(chain.from_iterable(events))}, len(events))
    assert saved(left) == saved(right)
    assert sum(bool(set(e) & {'a', 'b'}) for e in left) == 1
    assert sum(bool(set(e) & {'a', 'b'}) for e in right) == 2


def test_only_city_journey_and_its_ratio_omit_baseline_other_metrics_keep_comparison():
    rows = counters() + [row('linecount', 'a', 5, group='bus'),
                          row('linecount', 'a', 2, hour=-168, group='bus')]
    result = network(source(rows, ('a',)))
    assert amount(result, 'trip-types').comparison.text == ''
    assert amount(result, 'transfer-coefficient').comparison.text == ''
    assert amount(result, 'linecount').comparison.amount == 150
    period = network(source(rows, ('a',)), 'period')
    assert next(c for c in period.charts if c.key == 'trip-types').result.metric.scope == 'company'


def test_filtered_boardings_or_different_period_cannot_use_city_denominator():
    snapshot = source(counters())
    board = snapshot.city_boardings
    bad = replace(snapshot, city_boardings=replace(board, query=replace(board.query, companies=('a',))))
    result, reason = overall_transfer(bad)
    assert result is None and '完整同期观测' in reason
    bad = replace(snapshot, city_boardings=replace(board, current_window=(START, END + timedelta(hours=1))))
    assert overall_transfer(bad)[0] is None


def test_query_cancellation_is_preserved():
    snapshot = source(counters())
    with pytest.raises(QueryCancelled):
        overall_transfer(snapshot, lambda: True)


def test_export_keeps_city_scope_and_truthful_company_detail_name(tmp_path):
    from stats_exports import export_xlsx
    from openpyxl import load_workbook
    snapshot = source(counters(), ('a',))
    path = tmp_path / 'journeys.xlsx'
    export_xlsx(snapshot, path, NAMES, network(snapshot))
    book = load_workbook(path, data_only=True)
    assert any(r[7] == '出行量' and r[8] == 1 and '全市公共交通' in (r[16] or '')
               and r[3] is None and r[4] is None and not r[14]
               for r in list(book['网络摘要'].values)[1:])
    assert any('公司参与行程' in str(r[0]) for sheet in book for r in sheet.values)
    missing = source([r for r in counters() if r['指标'] != 'public-transport'])
    export_xlsx(missing, path, NAMES, network(missing))
    assert any(r[7] == '出行量' and r[8] is None and r[11] == '公共交通行程数据不完整'
               for r in load_workbook(path, data_only=True)['网络摘要'].values)
