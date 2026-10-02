import os
from datetime import datetime as D
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtCore import QSettings, Signal
from PySide6.QtGui import QImage
from PySide6.QtWidgets import QApplication, QFrame, QLabel, QPushButton, QVBoxLayout, QWidget
from qfluentwidgets import CardWidget, PushButton, RoundMenu
from openpyxl import load_workbook

from dashboard_model import DashboardResult, FilterState
from stats_integration import configure_dashboard


class Page(QWidget):
    snapshot_changed = Signal(object)
    export_requested = Signal()

    def __init__(self):
        super().__init__()
        self.snapshot = SimpleNamespace(alerts=[])
        self.session_key = 'save-a'
        self.alerts_host = QFrame(self)
        host_layout = QVBoxLayout(self.alerts_host)
        host_layout.addWidget(QLabel('关键变化提醒'))
        self.board_host = QWidget(self)
        board_layout = QVBoxLayout(self.board_host)
        for _ in range(12):
            label = QLabel('large board')
            label.setFixedHeight(60)
            board_layout.addWidget(label)
        self.board_host.resize(160, 100)
        self.alert_toggle = QPushButton(self)
        self.export_button = QPushButton(self)

    def _names(self):
        return {'p1': 'Company'}

    def export_target(self):
        return self.board_host


def test_configure_keeps_export_without_duplicate_home_alerts(tmp_path):
    from stats_alerts import AlertsPanel
    QApplication.instance() or QApplication([])
    page = Page()
    controller = configure_dashboard(page, QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    assert page.findChildren(AlertsPanel) == []
    assert configure_dashboard(page) is controller
    page.snapshot_changed.emit(page.snapshot)
    assert page.findChildren(AlertsPanel) == []
    page.snapshot = None
    page.snapshot_changed.emit(None)
    assert page.findChildren(AlertsPanel) == []


def test_png_menu_exports_full_board_and_restores_size(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    page = Page()
    controller = configure_dashboard(page, QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    old_size = page.board_host.size()
    target = tmp_path / 'complete.png'
    monkeypatch.setattr('stats_integration.QFileDialog.getSaveFileName',
                        lambda *_a, **_kw: (str(target), ''))
    monkeypatch.setattr(RoundMenu, 'exec', lambda menu, *_: menu.actions()[0].trigger())
    page.export_requested.emit()
    image = QImage(str(target))
    assert not image.isNull()
    assert image.height() > 600
    assert page.board_host.size() == old_size
    assert controller.last_error is None


def test_changed_snapshot_during_dialog_cancels_export(tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    page = Page()
    configure_dashboard(page, QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    target = tmp_path / 'stale.png'
    def dialog(*_a, **_kw):
        page.snapshot = SimpleNamespace(alerts=[])
        return str(target), ''
    monkeypatch.setattr('stats_integration.QFileDialog.getSaveFileName', dialog)
    monkeypatch.setattr(RoundMenu, 'exec', lambda menu, *_: menu.actions()[0].trigger())
    page.export_requested.emit()
    assert not target.exists()


def test_xlsx_menu_uses_current_snapshot(tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    page = Page()
    page.snapshot = DashboardResult(FilterState(('p1',), D(2024, 1, 1), D(2024, 1, 2)),
                                    {}, {}, [])
    configure_dashboard(page, QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    target = tmp_path / 'board.xlsx'
    monkeypatch.setattr('stats_integration.QFileDialog.getSaveFileName',
                        lambda *_a, **_kw: (str(target), ''))
    monkeypatch.setattr(RoundMenu, 'exec', lambda menu, *_: menu.actions()[1].trigger())
    page.export_requested.emit()
    assert len(load_workbook(target).sheetnames) == 5


def test_export_error_is_signalled_without_crash(tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    page = Page()
    controller = configure_dashboard(page, QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    failures = []
    controller.export_failed.connect(failures.append)
    monkeypatch.setattr('stats_integration.QFileDialog.getSaveFileName',
                        lambda *_a, **_kw: (str(tmp_path / 'bad.png'), ''))
    monkeypatch.setattr('stats_integration.export_png', lambda *_: (_ for _ in ()).throw(OSError('disk')))
    monkeypatch.setattr(RoundMenu, 'exec', lambda menu, *_: menu.actions()[0].trigger())
    page.export_requested.emit()
    assert isinstance(controller.last_error, OSError)
    assert failures == [controller.last_error]
