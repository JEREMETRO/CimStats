"""Bounded offscreen regression for the homepage controller and shell wiring.

Imports are local so --noconftest --collect-only never initializes Qt.
The parser is the only external boundary replaced; real pages/models run.
"""
from datetime import datetime as D
import ast
import importlib.util
from pathlib import Path
import time

import pytest
from openpyxl import load_workbook


def test_controller_imports_when_frontend_precedes_src():
    import os
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    code = """import sys
from pathlib import Path
root = Path.cwd()
sys.path.insert(0, str(root / 'src'))
sys.path.insert(0, str(root / 'frontend'))
import latest_info_controller as controller
assert controller.build_latest_alerts.__module__ == 'src.latest_info_alerts'
assert controller.LatestInfoAlertsPanel.__module__ == 'frontend.latest_info_alerts'
"""
    result = subprocess.run([sys.executable, '-c', code], cwd=root,
                            env=dict(os.environ, QT_QPA_PLATFORM='offscreen'),
                            capture_output=True, text=True, timeout=20)
    assert result.returncode == 0, result.stderr


def _start_from_foreign_cwd(tmp_path, entry):
    import os
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    code = """import os, sys
from pathlib import Path
frontend = Path(sys.argv[1])
assert os.environ['QT_QPA_PLATFORM'] == 'offscreen'
assert not os.environ.get('PYTHONPATH')
assert str(frontend.parent) not in sys.path
if sys.argv[2] == 'official':
    # Python sets sys.path[0] to the root when executing CIM2_SaveStats.py.
    sys.path.insert(0, str(frontend.parent))
    import CIM2_SaveStats
else:
    sys.path.insert(0, str(frontend))
import desktop_app
from PySide6.QtCore import QCoreApplication, QEvent, QSettings
from PySide6.QtWidgets import QApplication
app = QApplication([])
assert app.platformName() == 'offscreen'
settings = QSettings(str(Path.cwd() / 'startup.ini'), QSettings.Format.IniFormat)
desktop_app.QSettings = lambda *_: settings
desktop_app.MainWindow.check_install = lambda _self: None
window = desktop_app.MainWindow()
assert window.latest_info_page is window.pages.widget(0)
assert len(window.latest_info_page.metric_cards) == 13
assert sys.path.index(str(desktop_app.SRC)) < sys.path.index(str(desktop_app.PROJECT))
assert window.close()
window.deleteLater()
QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
app.processEvents()
"""
    environment = dict(os.environ, QT_QPA_PLATFORM='offscreen')
    environment.pop('PYTHONPATH', None)
    result = subprocess.run([sys.executable, '-c', code, str(root / 'frontend'), entry],
                            cwd=tmp_path, env=environment, capture_output=True,
                            text=True, timeout=25)
    assert result.returncode == 0, result.stderr


def test_main_window_starts_from_foreign_cwd_without_pythonpath(tmp_path):
    _start_from_foreign_cwd(tmp_path, 'frontend')


def test_official_launcher_starts_from_foreign_cwd_without_pythonpath(tmp_path):
    _start_from_foreign_cwd(tmp_path, 'official')


def session(key='session-a', save='秋山市.save'):
    def line(key, owner, name, passengers, departures):
        return {'key': key, '公司标识': owner, '公司名称': '同名公司', '运输制式': '公交',
                '线路名称': name, '今日客流': passengers, '当日发班数': departures,
                '行车总时间': 80, '地图里程': 10, '核定速度': 30,
                '每周收入': 10, '每周支出': 4, '班次': {}, '班次数据完整': True}
    return {'save_key': key, 'save_path': save, 'simulation_time': '2024-03-03 12:30:00',
            'metadata': {'地图名称': '秋山市', '当前人口数': 12345},
            'companies': [{'公司标识': 'p1', '公司名称': '同名公司', '车辆总数': 4, '车队': {'公交': 4}},
                          {'公司标识': 'p2', '公司名称': '同名公司', '车辆总数': 6, '车队': {'公交': 6}}],
            'lines': [line('p1|bus|1', 'p1', '一号线路', 100, 4),
                      line('p2|bus|1', 'p2', '二号线路', 200, 8)], 'history': [], 'outputs': {}}


@pytest.fixture
def window(qt_application, tmp_path, monkeypatch):
    assert qt_application.platformName() == 'offscreen'
    from PySide6.QtCore import QSettings
    import desktop_app
    settings = QSettings(str(tmp_path / 'shell.ini'), QSettings.Format.IniFormat)
    monkeypatch.setattr(desktop_app, 'QSettings', lambda *_: settings)
    # Installation discovery is unrelated to the new home and can show a modal.
    monkeypatch.setattr(desktop_app.MainWindow, 'check_install', lambda _self: None)
    shell = desktop_app.MainWindow()
    shell.resize(1440, 960)
    shell.show()
    qt_application.processEvents()
    yield shell
    shell.close()
    shell.deleteLater()
    qt_application.processEvents()


def wait_for(qt_application, predicate, timeout=8):
    from PySide6.QtTest import QTest
    deadline = time.monotonic() + timeout
    while not predicate() and time.monotonic() < deadline:
        qt_application.processEvents()
        QTest.qWait(10)
    assert predicate(), '最新信息异步查询未在期限内完成'


def ready(window, qt_application, data=None):
    window.on_completed(data or session())
    wait_for(qt_application, lambda: window.latest_info_controller.snapshot is not None)
    return window.latest_info_controller


def test_navigation_renames_only_home_and_keeps_other_pages(window):
    from latest_info_page import LatestInfoPage
    assert [button.text() for button in window.nav_buttons] == ['最新信息', '线路查询', '统计数据']
    assert window.pages.count() == 3
    assert isinstance(window.pages.widget(0), LatestInfoPage)
    # One shared header: only its title and the statistics sub-tabs change.
    window.navigate(0)
    assert window.header.title.text() == '最新信息' and window.stats_tabs.isHidden()
    window.navigate(1)
    assert window.pages.currentWidget() is window.lines_page
    assert window.header.title.text() == '线路查询' and window.stats_tabs.isHidden()
    window.navigate(2)
    assert window.pages.currentWidget() is window.statistics_page
    assert window.header.title.text() == '统计数据' and window.stats_tabs.parentWidget() is window.header


def test_reminders_have_one_home_entry_and_keep_shared_thresholds(window):
    from PySide6.QtWidgets import QAbstractButton
    from frontend.latest_info_alerts import LatestInfoAlertsPanel
    from stats_alerts import AlertsPanel
    panel = window.latest_info_controller.alerts_panel
    assert isinstance(panel, LatestInfoAlertsPanel)
    assert panel in window.latest_info_page.findChildren(LatestInfoAlertsPanel)
    assert window.statistics_page.findChildren(AlertsPanel) == []
    controls = window.statistics_page.findChildren(QAbstractButton)
    assert not any(button.isVisibleTo(window.statistics_page) and button.text() in
                   ('关键变化提醒', '提醒阈值') for button in controls)
    panel.thresholds_changed.emit((7, 30, 150))
    assert window.statistics_page.thresholds == (7, 30, 150)
    assert float(window.settings.value('statistics/percentage_points')) == 7
    assert float(window.settings.value('statistics/relative_percent')) == 30
    assert float(window.settings.value('statistics/passenger_absolute')) == 150


def test_home_alerts_are_queried_without_visiting_statistics_tab(window, qt_application):
    controller = ready(window, qt_application)
    window.navigate(0)
    alerts = controller.alerts_snapshot
    assert alerts is not None and alerts.alerts == []
    assert alerts.filters.start == D(2024, 3, 2)
    assert alerts.filters.end == D(2024, 3, 3)
    assert alerts.filters.comparison == (D(2024, 3, 1), D(2024, 3, 2))
    assert alerts.filters.grain == 'hour'
    assert set(alerts.filters.companies) == {'p1', 'p2'}


def test_share_and_report_use_current_scope(window, qt_application, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    controller = ready(window, qt_application)
    page = window.latest_info_page
    page.company_combo.setCurrentIndex(page.company_combo.findData('p2'))
    wait_for(qt_application, lambda: controller.snapshot is not None and controller.snapshot.company_id == 'p2')
    controller.copy_summary()
    summary = QApplication.clipboard().text()
    assert '同名公司 [p2]' in summary and '二号线路' in summary
    assert '一号线路' not in summary
    target = tmp_path / '当前范围.xlsx'
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName',
                        lambda *_a, **_kw: (str(target), ''))
    controller._export('xlsx')
    wb = load_workbook(target, data_only=True)
    assert list(wb['客流前十'].values)[1][2] == '二号线路'
    assert len(list(wb['客流前十'].values)) == 2
    assert list(wb['班次分类'].values)[-1][1] == 8


@pytest.mark.parametrize('kind', ['xlsx', 'png'])
def test_report_save_cancel_has_no_file(window, qt_application, tmp_path, monkeypatch, kind):
    controller = ready(window, qt_application)
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', lambda *_a, **_kw: ('', ''))
    controller._export(kind)
    assert list(tmp_path.glob('*.xlsx')) == [] and list(tmp_path.glob('*.png')) == []


@pytest.mark.parametrize('kind', ['xlsx', 'png'])
def test_new_session_during_save_dialog_cannot_export_old_snapshot(window, qt_application, tmp_path, monkeypatch, kind):
    controller = ready(window, qt_application)
    target = tmp_path / ('过期.' + kind)
    def dialog(*_args, **_kwargs):
        controller.set_session(session('session-b', '新存档.save'))
        return str(target), ''
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', dialog)
    controller._export(kind)
    assert not target.exists()


def test_superseded_query_cannot_restore_previous_scope(window, qt_application):
    controller = ready(window, qt_application)
    old_token, old_snapshot, old_alerts = controller.token, controller.snapshot, controller.alerts_snapshot
    controller.set_session(session('session-b', '新存档.save'))
    controller._receive(old_token, (old_snapshot, old_alerts))
    assert controller.snapshot is None or controller.snapshot.session_key == 'session-b'
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.snapshot.session_key == 'session-b'


@pytest.mark.parametrize('message', ['已取消解析', '读取失败'])
def test_import_cancel_and_failure_clear_home(window, qt_application, monkeypatch, message):
    controller = ready(window, qt_application)
    old_token, old_snapshot, old_alerts = controller.token, controller.snapshot, controller.alerts_snapshot
    monkeypatch.setattr('desktop_app.QMessageBox.critical', lambda *_a, **_kw: None)
    window.on_failed(message)
    assert controller.snapshot is None and controller.alerts_snapshot is None
    controller._receive(old_token, (old_snapshot, old_alerts))
    assert controller.snapshot is None


def test_reimport_clears_before_parser_runs_and_preserves_open_action(window, qt_application, tmp_path, monkeypatch):
    from PySide6.QtCore import QObject, Signal
    controller = ready(window, qt_application)
    import desktop_app
    class PausedParser(QObject):
        progress = Signal(int, str)
        log = Signal(str)
        completed = Signal(object)
        failed = Signal(str)
        finished = Signal()
        cancel_requested = False
        process = None
        def __init__(self, _path, _managed, parent):
            super().__init__(parent)
        def start(self):
            pass
        def isRunning(self):
            return False
        def cancel(self):
            self.cancel_requested = True
        def wait(self, _timeout):
            return True
    monkeypatch.setattr(desktop_app, 'ParseWorker', PausedParser)
    monkeypatch.setattr(desktop_app, 'locate_managed', lambda *_: tmp_path)
    target = tmp_path / '替换.save'
    target.write_bytes(b'fixture')
    window.start_parse(target)
    assert controller.snapshot is None and controller.alerts_snapshot is None
    assert not window.header.export_actions['line_workbook'].isEnabled()
    assert not window.header.export_actions['company_workbook'].isEnabled()
    calls = []
    def open_dialog(*_args, **_kwargs):
        calls.append('open')
        return '', ''
    monkeypatch.setattr('desktop_app.QFileDialog.getOpenFileName', open_dialog)
    window.latest_info_page.open_save_requested.emit()
    assert calls == ['open']


def test_original_workbooks_remain_available_from_home(window, qt_application, tmp_path, monkeypatch):
    ready(window, qt_application)
    source = tmp_path / '原线路.xlsx'
    source.write_bytes(b'original workbook fixture')
    target = tmp_path / '复制.xlsx'
    window.data['outputs']['line_workbook'] = str(source)
    monkeypatch.setattr('desktop_app.QFileDialog.getSaveFileName', lambda *_a, **_kw: (str(target), ''))
    window.latest_info_page.line_export_requested.emit()
    assert target.read_bytes() == source.read_bytes()


def test_close_interrupts_latest_info_workers(window, qt_application):
    controller = ready(window, qt_application)
    window.close()
    assert controller.snapshot is None and controller.alerts_snapshot is None
    assert all(not worker.isRunning() for worker in controller.workers)


# The parent may release this component subset before releasing MainWindow.
# These cases use real B/C widgets and models; no application shell is imported.
@pytest.fixture
def component(qt_application, tmp_path):
    assert qt_application.platformName() == 'offscreen'
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QVBoxLayout, QWidget
    from latest_info_page import LatestInfoPage
    host = QWidget()
    settings = QSettings(str(tmp_path / 'component.ini'), QSettings.Format.IniFormat)
    page = LatestInfoPage(settings, host)
    layout = QVBoxLayout(host)
    layout.addWidget(page)
    controller = None
    if importlib.util.find_spec('latest_info_controller') is not None:
        from latest_info_controller import LatestInfoController
        controller = LatestInfoController(page, settings, host)
    host.resize(1300, 900)
    host.show()
    qt_application.processEvents()
    yield page, controller
    if controller is not None:
        controller.stop_workers()
    host.close()
    host.deleteLater()
    qt_application.processEvents()


def ready_component(component, qt_application):
    page, controller = component
    assert controller is not None, '首页控制器尚未实现'
    controller.set_session(session())
    wait_for(qt_application, lambda: controller.snapshot is not None)
    return page, controller


def test_component_builds_real_home_and_alerts_in_one_current_session(component, qt_application):
    from frontend.latest_info_alerts import LatestInfoAlertsPanel
    from latest_info_model import LatestInfoSnapshot
    page, controller = ready_component(component, qt_application)
    assert isinstance(controller.snapshot, LatestInfoSnapshot)
    assert isinstance(controller.alerts_panel, LatestInfoAlertsPanel)
    assert controller.alerts_panel in page.findChildren(LatestInfoAlertsPanel)
    assert controller.snapshot.session_key == 'session-a'
    assert {line.key for line in controller.snapshot.lines} == {'p1|bus|1', 'p2|bus|1'}
    assert controller.snapshot.total_departures == 12
    assert controller.alerts_snapshot is not None and controller.alerts_snapshot.alerts == []
    assert set(controller.alerts_snapshot.filters.companies) == {'p1', 'p2'}
    assert controller.alerts_snapshot.filters.start == D(2024, 3, 2)
    assert controller.alerts_snapshot.filters.comparison == (D(2024, 3, 1), D(2024, 3, 2))


def test_component_scope_change_clears_before_query_and_keeps_stable_company_id(component, qt_application):
    page, controller = ready_component(component, qt_application)
    page.company_combo.setCurrentIndex(page.company_combo.findData('p2'))
    assert controller.snapshot is None and controller.alerts_snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.snapshot.company_id == 'p2'
    assert [line.name for line in controller.snapshot.lines] == ['二号线路']
    assert controller.alerts_snapshot.filters.companies == ('p2',)


def test_component_clear_rejects_old_query_result(component, qt_application):
    _page, controller = ready_component(component, qt_application)
    old_token, snapshot, alerts = controller.token, controller.snapshot, controller.alerts_snapshot
    controller.clear_session()
    controller._receive(old_token, (snapshot, alerts))
    assert controller.snapshot is None and controller.alerts_snapshot is None


def test_component_reimport_rejects_old_query_result(component, qt_application):
    _page, controller = ready_component(component, qt_application)
    old_token, snapshot, alerts = controller.token, controller.snapshot, controller.alerts_snapshot
    controller.set_session(session('session-b', '新存档.save'))
    controller._receive(old_token, (snapshot, alerts))
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.snapshot.session_key == 'session-b'
    assert controller.snapshot.save_name == '新存档.save'


@pytest.mark.parametrize('kind', ['xlsx', 'png'])
def test_component_report_save_cancel_does_not_write(component, qt_application, tmp_path, monkeypatch, kind):
    _page, controller = ready_component(component, qt_application)
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', lambda *_a, **_kw: ('', ''))
    controller._export(kind)
    assert not list(tmp_path.glob('*.xlsx')) and not list(tmp_path.glob('*.png'))


@pytest.mark.parametrize('kind', ['xlsx', 'png'])
def test_component_dialog_session_change_does_not_export_old_snapshot(component, qt_application, tmp_path, monkeypatch, kind):
    _page, controller = ready_component(component, qt_application)
    target = tmp_path / ('过期.' + kind)
    def dialog(*_a, **_kw):
        controller.set_session(session('session-b', '新存档.save'))
        return str(target), ''
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', dialog)
    controller._export(kind)
    assert not target.exists()


def test_component_report_and_share_use_same_selected_snapshot(component, qt_application, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    page, controller = ready_component(component, qt_application)
    page.company_combo.setCurrentIndex(page.company_combo.findData('p2'))
    wait_for(qt_application, lambda: controller.snapshot is not None and controller.snapshot.company_id == 'p2')
    target = tmp_path / '当前公司.xlsx'
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', lambda *_a, **_kw: (str(target), ''))
    controller._export('xlsx')
    controller.copy_summary()
    summary = QApplication.clipboard().text()
    wb = load_workbook(target, data_only=True)
    assert '同名公司 [p2]' in summary and '二号线路' in summary
    assert '一号线路' not in summary
    assert list(wb['客流前十'].values)[1][2] == '二号线路'
    assert len(list(wb['客流前十'].values)) == 2
    assert list(wb['班次分类'].values)[-1][1] == 8


def test_component_missing_simulation_clock_keeps_calculable_home(component, qt_application, monkeypatch):
    _page, controller = component
    data = session()
    data.pop('simulation_time')
    def forbidden(*_a, **_kw):
        raise AssertionError('缺少模拟时钟不能主动查询提醒')
    monkeypatch.setattr('latest_info_controller.build_latest_alerts', forbidden)
    errors = []
    controller.query_failed.connect(errors.append)
    controller.set_session(data)
    wait_for(qt_application, lambda: controller.snapshot is not None or bool(errors))
    assert errors == []
    assert controller.snapshot.simulation_time is None
    assert controller.snapshot.total_departures == 12
    assert len(controller.snapshot.lines) == 2
    assert controller.alerts_snapshot is None and controller.alerts_panel.alerts == []


def test_component_disabled_alerts_stop_active_query_and_reenable_requeries(component, qt_application, monkeypatch):
    _page, controller = ready_component(component, qt_application)
    from latest_info_alerts import build_latest_alerts
    def forbidden(*_a, **_kw):
        raise AssertionError('提醒关闭后不能主动查询')
    monkeypatch.setattr('latest_info_controller.build_latest_alerts', forbidden)
    controller.alerts_panel.set_alerts_enabled(False)
    assert controller.snapshot is None and controller.alerts_snapshot is None
    errors = []
    controller.query_failed.connect(errors.append)
    wait_for(qt_application, lambda: controller.snapshot is not None or bool(errors))
    assert errors == [] and controller.alerts_snapshot is None
    monkeypatch.setattr('latest_info_controller.build_latest_alerts', build_latest_alerts)
    controller.alerts_panel.set_alerts_enabled(True)
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.alerts_snapshot is not None


def test_component_threshold_change_invalidates_and_updates_real_query(component, qt_application):
    _page, controller = component
    data = session()
    data['history'] = [dict(指标='cashflow', 分组='总计', 公司标识='p2',
                            模拟时间=f'2024-03-0{day} {hour:02}:00:00', 值=str(value), 分母='0', 当前槽位='False')
                       for day, value in ((1, 100), (2, 150)) for hour in range(24)]
    controller.set_session(data)
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert len(controller.alerts_snapshot.alerts) == 1
    old_token = controller.token
    controller.alerts_panel.set_thresholds((5, 60, 100))
    assert controller.token > old_token and controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.alerts_snapshot.alerts == []
    assert controller.thresholds == (5, 60, 100)


@pytest.mark.parametrize('kind', ['xlsx', 'png'])
def test_component_dialog_session_identity_change_blocks_report(component, qt_application, tmp_path, monkeypatch, kind):
    _page, controller = ready_component(component, qt_application)
    target = tmp_path / ('替换会话.' + kind)
    def dialog(*_a, **_kw):
        controller.data['save_key'] = 'changed-session-identity'
        return str(target), ''
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', dialog)
    controller._export(kind)
    assert not target.exists()


def test_component_png_exports_actual_page_and_restores_geometry(component, qt_application, tmp_path, monkeypatch):
    from PySide6.QtGui import QImage
    page, controller = ready_component(component, qt_application)
    target = tmp_path / '实际页面.png'
    old_size = page.export_target().size()
    monkeypatch.setattr('latest_info_controller.QFileDialog.getSaveFileName', lambda *_a, **_kw: (str(target), ''))
    controller._export('png')
    image = QImage(str(target))
    assert not image.isNull()
    assert image.height() >= old_size.height() and image.width() >= old_size.width()
    assert page.export_target().size() == old_size
    assert controller.last_error is None


def test_component_current_failure_clears_and_old_failure_is_ignored(component, qt_application, monkeypatch):
    _page, controller = ready_component(component, qt_application)
    errors = []
    controller.query_failed.connect(errors.append)
    old_token = controller.token
    controller._failed(old_token - 1, RuntimeError('old'))
    assert errors == [] and controller.snapshot is not None
    def broken(*_a, **_kw):
        raise RuntimeError('当前查询失败')
    monkeypatch.setattr('latest_info_controller.build_latest_info', broken)
    controller.schedule_query()
    wait_for(qt_application, lambda: bool(errors))
    assert str(errors[0]) == '当前查询失败'
    assert controller.snapshot is None and controller.alerts_snapshot is None


def test_component_stop_clears_and_rejects_queued_result(component, qt_application):
    _page, controller = ready_component(component, qt_application)
    token, result = controller.token, (controller.snapshot, controller.alerts_snapshot)
    assert controller.stop_workers()
    controller._receive(token, result)
    assert controller.snapshot is None and controller.alerts_snapshot is None
    assert all(not worker.isRunning() for worker in controller.workers)


def test_component_statistics_has_no_second_reminder_ui(component, qt_application, tmp_path):
    from PySide6.QtCore import QSettings
    from stats_alerts import AlertsPanel
    from statistics_page import StatisticsPage
    from PySide6.QtWidgets import QAbstractButton
    _page, controller = component
    statistics = StatisticsPage(QSettings(str(tmp_path / 'statistics.ini'), QSettings.Format.IniFormat))
    try:
        statistics.show()
        qt_application.processEvents()
        assert statistics.findChildren(AlertsPanel) == []
        assert not any(button.isVisibleTo(statistics) and button.text() in ('关键变化提醒', '提醒阈值')
                       for button in statistics.findChildren(QAbstractButton))
        controller.thresholds_changed.connect(statistics.set_thresholds)
        controller.alerts_panel.set_thresholds((7, 30, 150))
        assert statistics.thresholds == (7, 30, 150)
        assert len(statistics._controls) == len(statistics._fields) == 7
        for width in (1400, 920):
            statistics.resize(width, 850)
            statistics.reflow(width)
            statistics.set_filters_collapsed(True, animate=False)
            statistics.set_filters_collapsed(False, animate=False)
    finally:
        statistics.stop_workers()
        statistics.close()
        statistics.deleteLater()


@pytest.mark.parametrize('line_available, company_available', [(False, False), (True, False), (False, True), (True, True)])
def test_component_original_workbook_actions_follow_real_files(component, qt_application, tmp_path, line_available, company_available):
    from openpyxl import Workbook
    from PySide6.QtWidgets import QPushButton
    from PySide6.QtGui import QAction
    page, controller = component
    data = session()
    for kind, available in (('line_workbook', line_available), ('company_workbook', company_available)):
        path = tmp_path / (kind + '.xlsx')
        if available:
            Workbook().save(path)
        data['outputs'][kind] = str(path)
    controller.set_session(data)
    wait_for(qt_application, lambda: controller.snapshot is not None)
    for name, text, enabled in (('lineExportAction', '导出线路 XLSX', line_available),
                                ('companyExportAction', '导出公司 XLSX', company_available)):
        assert page.findChild(QPushButton, name).isEnabled() == enabled
        actions = [action for action in page.findChildren(QAction) if action.text() == text]
        assert actions and all(action.isEnabled() == enabled for action in actions)
    assert page.findChild(QPushButton, 'shareAction').isEnabled()
    assert page.findChild(QPushButton, 'reportAction').isEnabled()
    page.company_combo.setCurrentIndex(page.company_combo.findData('p2'))
    assert not page.findChild(QPushButton, 'lineExportAction').isEnabled()
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert page.findChild(QPushButton, 'lineExportAction').isEnabled() == line_available
    controller.set_session(session('session-b'))
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert not page.findChild(QPushButton, 'lineExportAction').isEnabled()
    assert not page.findChild(QPushButton, 'companyExportAction').isEnabled()


def test_component_original_workbook_directory_is_unavailable(component, qt_application, tmp_path):
    from PySide6.QtWidgets import QPushButton
    page, controller = component
    data = session()
    directory = tmp_path / '目录.xlsx'
    directory.mkdir()
    data['outputs']['line_workbook'] = str(directory)
    controller.set_session(data)
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert not page.findChild(QPushButton, 'lineExportAction').isEnabled()


def shell_method(name, **namespace):
    """Exercise a bounded shell method without importing/constructing MainWindow."""
    source = Path(__file__).resolve().parents[1] / 'frontend' / 'desktop_app.py'
    tree = ast.parse(source.read_text(encoding='utf-8'))
    shell = next(node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == 'MainWindow')
    method = next(node for node in shell.body if isinstance(node, ast.FunctionDef) and node.name == name)
    exec(compile(ast.Module(body=[method], type_ignores=[]), str(source), 'exec'), namespace)
    return namespace[name]


def test_component_shell_readiness_waits_for_home_snapshot(component):
    from types import SimpleNamespace as NS
    _page, controller = component
    finished = []
    data = session()
    stats = NS(snapshot=object(), network_snapshot=object(), query_timer=NS(isActive=lambda: False), workers=[], network_workers=[])
    shell = NS(_awaiting_dashboards=True, data=data, _ready_data=data, _ready_timer=NS(stop=lambda: None),
               _progress_timer=NS(stop=lambda: None), latest_info_controller=controller, statistics_page=stats,
               cancel_action=NS(setEnabled=lambda _value: None), _progress_predictor=None,
               loading_overlay=NS(finish=lambda: finished.append(True)), _ready_started=time.monotonic())
    shell_method('_check_dashboard_ready', time=time)(shell)
    assert finished == [] and shell._awaiting_dashboards


def test_component_shell_failure_invalidates_home(component, qt_application):
    from types import SimpleNamespace as NS
    _page, controller = ready_component(component, qt_application)
    widget = NS(setEnabled=lambda _v: None, setText=lambda _v: None, setToolTip=lambda _v: None, hide=lambda: None)
    shell = NS(_awaiting_dashboards=False, _ready_timer=NS(stop=lambda: None), loading_overlay=NS(finish=lambda: None),
               _progress_timer=NS(stop=lambda: None), cancel_action=widget, progress=widget,
               latest_info_controller=controller, statistics_page=NS(clear_session=lambda: None),
               data=controller.data, header=NS(save_chip=NS(set_context=lambda *_: None)),
               empty_state=NS(set_note=lambda _v: None), _refresh_body=lambda: None,
               _refresh_exports=lambda: None)
    shell_method('on_failed')(shell, '已取消解析')
    assert controller.snapshot is None and controller.alerts_snapshot is None


def test_component_shell_original_workbook_dialog_rejects_reimport(component, qt_application, tmp_path):
    from types import SimpleNamespace as NS
    from openpyxl import Workbook
    import shutil
    _page, controller = ready_component(component, qt_application)
    source, target = tmp_path/'原始.xlsx', tmp_path/'过期复制.xlsx'
    Workbook().save(source)
    controller.data['outputs']['line_workbook'] = str(source)
    shell = NS(data=controller.data, latest_info_controller=controller,
               status_label=NS(setText=lambda _text: None))
    def dialog(*_a, **_kw):
        shell.data = session('session-b')
        controller.set_session(shell.data)
        return str(target), ''
    shell_method('export_file', Path=Path, shutil=shutil,
                 QFileDialog=NS(getSaveFileName=dialog), QMessageBox=NS(warning=lambda *_a: None, critical=lambda *_a: None))(shell, 'line_workbook')
    assert not target.exists()


# Revision 2: public chart-state/controller contracts only. Prepared while B/C
# own the Qt slot; do not execute until the parent explicitly releases D.
def revision_views(page, passengers=('ranking', '公交'), departures=('line_share', None)):
    from copy import deepcopy
    state = page.capture_chart_state()
    state['passengers'] = dict(zip(('view', 'mode'), passengers))
    state['departures'] = dict(zip(('view', 'mode'), departures))
    assert page.restore_chart_state(state)
    actual = page.capture_chart_state()
    assert actual == state
    return deepcopy(actual)


def revision_default_views(page):
    state = page.capture_chart_state()
    assert state['passengers'] == state['departures'] == {'view': 'structure', 'mode': None}


def revision_restore_calls(page, monkeypatch):
    """Observe the public contract without replacing its rendering behavior."""
    from copy import deepcopy
    restore = page.restore_chart_state
    calls = []
    def observed(state, *, allow_scope_change=False):
        calls.append((deepcopy(state), allow_scope_change))
        return restore(state, allow_scope_change=allow_scope_change)
    monkeypatch.setattr(page, 'restore_chart_state', observed)
    return calls


@pytest.mark.parametrize('passengers,departures', [
    (('structure', None), ('line_share', None)),
    (('ranking', None), ('ranking', '公交')),
    (('ranking', '公交'), ('line_share', None)),
    (('line_share', None), ('ranking', None)),
    (('line_share', '公交'), ('ranking', '公交')),
    (('ranking', '公交'), ('line_share', '公交')),
])
def test_component_revision_refresh_preserves_independent_views(component, qt_application, passengers, departures):
    from PySide6.QtWidgets import QWidget
    page, controller = ready_component(component, qt_application)
    expected = revision_views(page, passengers, departures)
    old_snapshot, token = controller.snapshot, controller.token
    published = []
    controller.snapshot_changed.connect(lambda value: published.append(
        (value, page.capture_chart_state()) if value is not None else (None, None)))
    controller.schedule_query()
    assert controller.token > token and controller.snapshot is None
    assert not page.findChild(QWidget, 'reportAction').isEnabled()
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert controller.snapshot is not old_snapshot
    assert page.capture_chart_state() == expected
    assert published[0] == (None, None)
    assert published[-1][0] is controller.snapshot and published[-1][1] == expected


@pytest.mark.parametrize('change', ['threshold', 'disable', 'enable'])
def test_component_revision_reminder_requery_preserves_views(component, qt_application, change):
    page, controller = ready_component(component, qt_application)
    if change == 'enable':
        controller.alerts_panel.set_alerts_enabled(False)
        wait_for(qt_application, lambda: controller.snapshot is not None)
    expected = revision_views(page, ('line_share', None), ('ranking', '公交'))
    token = controller.token
    if change == 'threshold':
        controller.alerts_panel.set_thresholds((7, 30, 150))
    else:
        controller.alerts_panel.set_alerts_enabled(change == 'enable')
    assert controller.token > token and controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert page.capture_chart_state() == expected
    if change == 'disable':
        assert controller.alerts_snapshot is None
    else:
        assert controller.alerts_snapshot is not None


def test_component_revision_repeated_pending_refresh_keeps_original_intent(component, qt_application):
    page, controller = ready_component(component, qt_application)
    expected = revision_views(page)
    token = controller.token
    for _ in range(3):
        controller.schedule_query()
        assert controller.snapshot is None
    assert controller.token >= token + 3
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert page.capture_chart_state() == expected


@pytest.mark.parametrize('passengers', [
    ('ranking', '公交'), ('ranking', None), ('line_share', None), ('line_share', '公交')])
def test_component_revision_company_scope_preserves_pure_intent_with_explicit_gate(
        component, qt_application, monkeypatch, passengers):
    from PySide6.QtWidgets import QWidget
    page, controller = ready_component(component, qt_application)
    expected = revision_views(page, passengers, ('line_share', None))
    calls = revision_restore_calls(page, monkeypatch)
    page.company_combo.setCurrentIndex(page.company_combo.findData('p2'))
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    actual = page.capture_chart_state()
    assert actual['scope'] == ('p2', '综合')
    assert actual['passengers'] == expected['passengers']
    assert actual['departures'] == expected['departures']
    assert (expected, True) in calls, '来源 scope/state 不变，跨范围恢复必须走公开 True 门禁'
    assert {line.key for line in controller.snapshot.lines} == {'p2|bus|1'}
    assert controller.snapshot.total_departures == 8
    departure_chart = page.findChild(QWidget, 'departureStructure')
    rows = departure_chart.visible_share_rows()
    assert len(rows) == 1 and rows[0].entry.line_key == 'p2|bus|1'
    assert rows[0].number.text() == '8 班' and rows[0].share.text() == '100.0%'
    passenger_chart = page.findChild(QWidget, 'passengerRanking')
    if passengers[0] == 'line_share':
        rows = passenger_chart.visible_share_rows()
        assert len(rows) == 1 and rows[0].entry.line_key == 'p2|bus|1'
        assert rows[0].number.text() == '200 人次' and rows[0].share.text() == '100.0%'
    else:
        rows = passenger_chart.visible_ranking_rows()
        assert len(rows) == 1 and rows[0].line.key == 'p2|bus|1'
        assert rows[0].number.text() == '200 人次'


@pytest.mark.parametrize('passenger_view', ['ranking', 'line_share'])
def test_component_revision_invalid_drill_mode_resets_only_one_chart(component, qt_application, passenger_view):
    page, controller = component
    data = session()
    data['lines'][1]['运输制式'] = '有轨电车'
    controller.set_session(data)
    wait_for(qt_application, lambda: controller.snapshot is not None)
    revision_views(page, (passenger_view, '公交'), ('line_share', None))
    page.mode_combo.setCurrentIndex(page.mode_combo.findData('有轨电车'))
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    actual = page.capture_chart_state()
    assert actual['passengers'] == {'view': 'structure', 'mode': None}
    assert actual['departures'] == {'view': 'line_share', 'mode': None}
    assert {line.mode for line in controller.snapshot.lines} == {'有轨电车'}
    assert controller.snapshot.total_departures == 8


def test_component_revision_pending_scope_changes_keep_intent_and_latest_target(component, qt_application):
    page, controller = ready_component(component, qt_application)
    expected = revision_views(page, ('line_share', None), ('ranking', None))
    old_token, old_result = controller.token, (controller.snapshot, controller.alerts_snapshot)
    for owner in ('p1', 'p2'):
        page.company_combo.setCurrentIndex(page.company_combo.findData(owner))
        assert controller.snapshot is None
    controller._receive(old_token, old_result)
    controller._failed(old_token, RuntimeError('过期范围失败'))
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    actual = page.capture_chart_state()
    assert actual['scope'] == ('p2', '综合')
    assert actual['passengers'] == expected['passengers']
    assert actual['departures'] == expected['departures']
    assert {line.key for line in controller.snapshot.lines} == {'p2|bus|1'}
    assert controller.last_error is None


def test_component_revision_wrong_scope_ready_cannot_consume_pending(component, qt_application):
    from latest_info_model import build_latest_info
    page, controller = ready_component(component, qt_application)
    expected = revision_views(page)
    controller.schedule_query()
    controller._receive(controller.token, (build_latest_info(controller.data, 'p2'), None))
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    assert page.capture_chart_state() == expected


@pytest.mark.parametrize('replacement', ['new-key', 'same-key-new-object', 'same-object-new-key'])
def test_component_revision_session_replacement_drops_pending(component, qt_application, replacement):
    page, controller = ready_component(component, qt_application)
    revision_views(page)
    old_token, old_result = controller.token, (controller.snapshot, controller.alerts_snapshot)
    controller.schedule_query()
    if replacement == 'same-object-new-key':
        controller.data['save_key'] = 'session-b'
        controller.schedule_query()
    else:
        controller.set_session(session('session-b' if replacement == 'new-key' else 'session-a', '新存档.save'))
    controller._receive(old_token, old_result)
    assert controller.snapshot is None
    wait_for(qt_application, lambda: controller.snapshot is not None)
    revision_default_views(page)
    assert controller.snapshot.session_key == ('session-a' if replacement == 'same-key-new-object' else 'session-b')


@pytest.mark.parametrize('clear_kind', ['clear', 'failure', 'close'])
def test_component_revision_clear_failure_close_drop_pending(component, qt_application, clear_kind):
    page, controller = ready_component(component, qt_application)
    revision_views(page)
    controller.schedule_query()
    token = controller.token
    if clear_kind == 'failure':
        controller._failed(token, RuntimeError('当前刷新失败'))
        controller.schedule_query()
        wait_for(qt_application, lambda: controller.snapshot is not None)
        revision_default_views(page)
    elif clear_kind == 'clear':
        controller.clear_session()
        controller.set_session(session())
        wait_for(qt_application, lambda: controller.snapshot is not None)
        revision_default_views(page)
    else:
        assert controller.stop_workers()
        assert controller.snapshot is None
        revision_default_views(page)
