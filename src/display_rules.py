"""Shared presentation rules for parser, desktop, and workbook output."""
import re
from decimal import Decimal, ROUND_DOWN, localcontext


def number_places(metric=None):
    return 1 if metric == 'vehicles-running' else 2


def truncated_number(value, places=2):
    """Truncate presentation values toward zero without changing source data."""
    if value is None:
        return None
    number = value if isinstance(value, Decimal) else Decimal(str(value))
    if not number.is_finite():
        return None
    with localcontext() as context:
        context.prec = max(context.prec, len(number.as_tuple().digits), number.adjusted() + places + 1)
        result = number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_DOWN)
    return result.copy_abs() if result.is_zero() else result


def format_number(value, places=2, *, grouped=True, fixed=False):
    number = truncated_number(value, places)
    if number is None:
        return '—'
    text = format(number, (',' if grouped else '') + f'.{places}f')
    if not fixed and '.' in text:
        text = text.rstrip('0').rstrip('.')
    return text


def workbook_number(value, places=2):
    number = truncated_number(value, places)
    if number is None:
        return None
    return int(number) if number == number.to_integral_value() else float(number)


def export_precision_workbook(source, target):
    """Export a cached workbook with presentation precision, preserving its source."""
    from pathlib import Path
    from openpyxl import load_workbook
    source, target = Path(source), Path(target)
    if source.resolve() == target.resolve() or (target.exists() and source.samefile(target)):
        raise ValueError('请选择缓存工作簿以外的导出位置')
    workbook = load_workbook(source)
    try:
        for sheet in workbook:
            vehicle_columns = set()
            for row in sheet:
                for cell in row:
                    # Sections can repeat headers; names, dates, IDs and formulas
                    # retain their original types and contents.
                    if isinstance(cell.value, str) and '平均运行车辆' in cell.value:
                        vehicle_columns.add(cell.column)
                    if isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool):
                        places = 1 if cell.column in vehicle_columns else 2
                        if '%' in cell.number_format:
                            cell.value = float(truncated_number(Decimal(str(cell.value)) * 100, places) / 100)
                            cell.number_format = '0.##%'
                        else:
                            cell.value = workbook_number(cell.value, places)
        workbook.save(target)
    finally:
        workbook.close()

MODE_NAMES = {
    "bus": "公交", "trolley": "无轨电车", "trolleybus": "无轨电车",
    "tram": "有轨电车", "metro": "地铁", "subway": "地铁",
    "monorail": "单轨", "waterbus": "水上巴士", "ferry": "水上巴士",
    "ship": "水上巴士", "anyvehicletype": "综合",
}


def display_mode(value):
    raw = str(value or "").strip()
    if raw == '单轨列车':
        return '单轨'
    return MODE_NAMES.get(raw.lower(), raw)


def format_line_name(number_value, raw_name):
    number_text = str(int(float(number_value or 0)))
    name = str(raw_name or "").strip()
    if not name:
        return f"{number_text}路"
    if re.fullmatch(r"[A-Za-z]", name):
        return f'{number_text}{name}'
    suffix = re.fullmatch(r"(\d+)(?:路-?|-)?([A-Za-z])", name)
    if suffix:
        return ''.join(suffix.groups())
    if name.isdigit():
        return f'{name}路'
    if re.match(r"\d", name):
        return name
    if len(name) < 3:
        return f"{number_text}路{name}"
    return f"{number_text}路·{name}"


def store_text_literally(workbook):
    """Keep save-supplied text such as a line named "=Express" out of formulas.

    openpyxl turns every string starting with "=" into a formula; these
    workbooks never contain real formulas, so store all such cells as text.
    """
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def format_garage(company, mode, value):
    """Keep serialized depot names; only bare positive indices need a label."""
    raw = str(value or '').strip()
    if not raw or raw.lower() in ('unknown', 'none', 'null', '未知', '—'):
        return ''
    if re.fullmatch(r'[-+]?\d+(?:\.0+)?', raw):
        index = int(float(raw))
        if index <= 0:
            return ''
        if not company or not mode:
            return ''
        return f'{company}{mode}车库{index}'
    return raw if not company or raw.startswith(company) else f'{company}{raw}'
