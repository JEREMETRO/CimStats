from __future__ import annotations
import csv, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
PROJECT=Path(__file__).resolve().parents[1]
SAVE=Path(sys.argv[1]) if len(sys.argv)>1 else PROJECT/'data'/'quicksave.76561198362520556-76561198845688243.save'
sys.argv=[sys.argv[0],str(SAVE)]
import extract_runtime_data as e
def old_rows():
 roots = [PROJECT / 'restored_reanalysis5', PROJECT / 'archive' / 'recycle_bin' / 'restored_reanalysis5']
 candidates = [p for root in roots if root.exists() for p in root.rglob('*线路发班整理*.xlsx')]
 if not candidates:
  return {}
 p=candidates[0]; ws=load_workbook(p,data_only=True,read_only=True)['线路信息']; vals=list(ws.values); hdr=next(r for r in vals if r and '线路号' in r); ix={v:i for i,v in enumerate(hdr)}; length_key='折算里程' if '折算里程' in ix else '线路长度'; out={}
 for r in vals:
  try:n=int(r[ix['线路号']])
  except:continue
  out[(str(r[ix['公司名称']]),str(r[ix['运输制式']]),n)]={'里程':r[ix[length_key]],'预计时间':r[ix['单程时间']],'最大需求':r[ix['理论最大车辆需求数']]}
 return out
root,_=e.load_root(); players,companies,_=e.company_rows(root); names={r['公司序号']:r['公司名称'] for r in companies}; display={"826272703's Company":'六进公交',"jeremylin2005's Company":'八连交通集团'}; old=old_rows();
simulation_start=e.field(root,'m_simulationStart'); simulation_time=e.field(root,'m_simulationTime'); simulation_frame=int(e.field(root,'m_simulationFrame') or 0); SPAN=350000
if simulation_start is not None and simulation_time is not None and simulation_frame:
 try: SPAN=max(1,int(round((simulation_time-simulation_start).Ticks/simulation_frame)))
 except Exception: pass
t=e.field(root,'m_transportManagerData'); line=e.field(t,'m_firstLine'); seen=set(); rows=[]
while line is not None:
 oid=int(e.field(line,'m_objectID') or 0)
 if oid in seen:break
 seen.add(oid); mode=e.line_mode(line); n=int(e.field(line,'m_number') or 0); depot=e.field(line,'m_depot'); owner=e.field(depot,'m_owner') if depot is not None else None; idx=next((i for i,p in enumerate(players) if owner is not None and e.System.Object.ReferenceEquals(owner,e.field(p,'m_companyData'))),None); company=display.get(names.get(idx,''),names.get(idx,'')); key=(company,mode,n)
 length,duration,top=e.recompute_runtime_line_values(line,mode,SPAN); base=old.get(key,{})
 def status(field,value,tol=1e-6):
  try:return '一致' if abs(float(base.get(field))-float(value))<=tol else '不一致'
  except:return '旧值无效'
 rows.append({'公司名称':company,'运输制式':mode,'线路号':n,'旧表里程_km':base.get('里程'),'重算里程_km':length/1024000,'里程比对':status('里程',length/1024000),'旧表预计时间_分钟':base.get('预计时间'),'重算预计时间_分钟':duration/600_000_000,'预计时间比对':status('预计时间',duration/600_000_000,.01),'旧表最大需求':base.get('最大需求'),'重算最大需求':top,'最大需求比对':status('最大需求',top,.01),'原始缓存里程':int(e.field(line,'m_lineLength') or 0),'原始缓存预计时间':int(e.field(line,'m_estimatedDuration') or 0),'原始缓存最大需求':int(e.field(line,'m_vehiclesNeededTop') or 0)})
 line=e.field(line,'m_nextLine')
out=PROJECT/'exports'/'CIM2_三项指标_强制重算全量校验.csv'; xlsx=out.with_suffix('.xlsx')
with out.open('w',encoding='utf-8-sig',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='逐线路校验';ws.append(list(rows[0]));[ws.append([r[k] for k in rows[0]]) for r in rows];sm=wb.create_sheet('汇总');sm.append(['指标','一致','不一致','旧值无效']);
for f in ('里程比对','预计时间比对','最大需求比对'):sm.append([f,sum(r[f]=='一致' for r in rows),sum(r[f]=='不一致' for r in rows),sum(r[f]=='旧值无效' for r in rows)])
wb.save(xlsx)
print('线路',len(rows));
for f in ('里程比对','预计时间比对','最大需求比对'):print(f,{s:sum(r[f]==s for r in rows) for s in ('一致','不一致','旧值无效')})
for r in rows:
 if r['线路号'] in {107,201,301,327,901}:print(r)
print(out,xlsx)
